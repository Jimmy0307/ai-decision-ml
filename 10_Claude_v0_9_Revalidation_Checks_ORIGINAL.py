"""
Claude re-validation checks for AI x Decision x MLP prototype v0.9 (audit date 2026-10-09).
Usage:  python3 claude_v0_9_revalidation_checks.py <path-to-CLAUDE_REVALIDATION_PACK_v0_9_20261009>
All coefficients are SIMULATED. Nothing here is an empirical result.

K0  SHA-256 manifest
K1  04 rerun vs 05 (all keys except runtime_seconds)
K2  06 rerun vs 07 (consistency check only, not independent validation)
K3  Generator-vs-canonical with an oracle written in this file (thresholds re-typed, not imported)
K4  Mutation: built-in flags (both / alpha-only / A0-only) and an EXTERNAL source-level mutant
K5  Option alpha binds only at A=0; L1 allocation-domain sizes
K6  G09 permission channel exercised from outside the solver (full solve)
K7  G10 tau sensitivity and HumanOnly reporting read from 05
K8  Regression for v0.7 R1: v0.7-literal P008 EA_req rule (no floor of 1) vs code (floor of 1 for A>=1)
K9  psi range: internal vs reportable (non-HumanOnly)
K10 Argmax tie-set sizes on the reported path (do T1-T3 bind?)
K11 06 re-implementation vs 05 on per-process fields that 06 itself does not compare
K12 16 depletion O/E vs manuscript Table 3 degrees and Table 5 observed counts (M = 2,180)
K13 Workbook Configuration_Rows / Governance_Scenarios / Process_Data vs solver constants
K14 Documented component formulas (Objective_Components) re-implemented from the workbook text vs solver f2/l3 on every generated option
K15 Threshold text in 01 and 08 P007/P008 vs code (adj inside max)
K16 Tie-break order: 01, 02, 08 Tie_Breaking and 05 vs solver CONFIG_ORDER
K17 Priority_Gaps frozen CSV vs workbook Priority_Gaps_Freeze vs Depletion_Quantification
K18 Priority_Gaps internal consistency: O_E_Ratio vs Observed/Null_Mean vs Observed/E(analytic); Null_Mean vs E within Monte Carlo error
K19 Priority_Gaps vs current conference manuscript Table 5 (observed, all four q <= .05, direction)
K20 Stale version labels in 08
"""
import sys, os, io, csv, json, hashlib, itertools, tempfile, importlib.util, contextlib, subprocess
import openpyxl

PACK = sys.argv[1] if len(sys.argv) > 1 else "."
P = lambda f: os.path.join(PACK, f)
L = range(4)

def load(path, name):
    s = importlib.util.spec_from_file_location(name, path); m = importlib.util.module_from_spec(s)
    with contextlib.redirect_stdout(io.StringIO()):
        s.loader.exec_module(m)
    return m

# K0
rows = list(csv.DictReader(open(P("14_SHA256_MANIFEST.csv"), encoding="utf-8-sig")))
ok = [hashlib.sha256(open(P(r["file"]), "rb").read()).hexdigest() == r["sha256"] for r in rows]
print("K0 manifest:", f"{sum(ok)}/{len(ok)} files re-hash correctly")

# K1
out = subprocess.run([sys.executable, P("04_solver_v0_9.py")], capture_output=True, text=True, check=True).stdout
a, b = json.loads(out), json.load(open(P("05_solver_results_v0_9.json"), encoding="utf-8"))
diff = [k for k in set(a) | set(b) if k != "runtime_seconds" and a.get(k) != b.get(k)]
print("K1 04 vs 05: differing keys (excluding runtime):", diff or "none")

# K2
out = subprocess.run([sys.executable, P("06_reimplementation_consistency_check_v0_9.py")], capture_output=True, text=True, check=True).stdout
a2, b2 = json.loads(out), json.load(open(P("07_reimplementation_consistency_results_v0_9.json"), encoding="utf-8"))
print("K2 06 vs 07: identical =", a2 == b2, "| official_all_pass =", a2["official_all_pass"])

S = load(P("04_solver_v0_9.py"), "s7")

# K3 oracle with thresholds re-typed from code semantics
def Hreq(A, k): return 0 if A == 0 else min(3, max(0, A + k - 2))
def Ereq_code(A, k): return 0 if A == 0 else min(3, max(1, A + k - 2))
def Wreq(A, ex): return 0 if A == 0 else (3 if ex and A == 3 else A)
def canon(p, caps, Ef=Ereq_code):
    Dc, Ac, Rc, Gc = caps; out = set()
    for A, H, WI, EA in itertools.product(L, L, L, L):
        if A > Ac or H > Gc or EA > Gc or WI > Dc: continue
        if Dc < A or Rc < A: continue
        if H < Hreq(A, p["risk"]) or EA < Ef(A, p["risk"]) or WI < Wreq(A, p["execution"]): continue
        if H > 3*A or EA > 3*A or WI > 3*A: continue
        out.add((A, H, WI, EA))
    return out
def conf(mod, **kw):
    c = s = m = 0
    for i, p in enumerate(mod.PROCESSES):
        for caps in itertools.product(L, L, L, L):
            gen = {(o["A"], o["H"], o["WI"], o["EA"]) for o in mod.l2_options_for_proc(i, caps, mod.GOVERNANCE["F"], **kw)}
            can = canon(p, caps); c += 1; s += len(can); m += gen != can
    return c, s, m
print("K3 conformance (cases, canonical states, mismatches):", conf(S))

# K4
print("K4 flag mutant, both guards off :", conf(S, enforce_option_alpha=False, mutant_relax_A0_generator=True))
print("K4 flag mutant, alpha line only :", conf(S, enforce_option_alpha=False))
print("K4 flag mutant, A0 branch only  :", conf(S, mutant_relax_A0_generator=True))
src = open(P("04_solver_v0_9.py"), encoding="utf-8").read()
mut = src.replace("Hvals = Evals = Wvals = [0]", "Hvals = range(0,Gcap+1); Evals = range(0,Gcap+1); Wvals = range(0,Dcap+1)", 1)
mut = mut.replace("if enforce_option_alpha and (H > 3*A or EA > 3*A or WI > 3*A):", "if False:", 1)
assert mut != src and mut.count("if False:") == 1, "mutation anchors not found"
t = os.path.join(tempfile.mkdtemp(), "mutant.py"); open(t, "w", encoding="utf-8").write(mut)
print("K4 external source-level mutant :", conf(load(t, "mutant")))

# K5
viol = [x for x in itertools.product(L, L, L, L) if not (x[1] <= 3*x[0] and x[2] <= 3*x[0] and x[3] <= 3*x[0])]
print("K5 alpha-violating tuples:", len(viol), "| all at A=0:", all(x[0] == 0 for x in viol))
dp = [1] + [0]*7
for _ in range(12):
    nd = [0]*8
    for s_, c_ in enumerate(dp):
        for v in range(4):
            if s_ + v <= 7: nd[s_ + v] += c_
    dp = nd
print("K5 allocations B=3/5/7:", [len(S.generate_allocations(B)) for B in (3, 5, 7)], "| combinatorial:", [sum(dp[:B+1]) for B in (3, 5, 7)])

# K6
fmt = lambda r: [(o["A"], o["H"], o["WI"], o["EA"], o["l3"]["q"]) for o in r["l2"]["opts"]]
S.best_l2_cached.cache_clear(); r0 = S.solve("C", 5)
S.ALLOW_CONFIG["C"]["AIAssist"] = False; S.best_l2_cached.cache_clear(); r1 = S.solve("C", 5)
S.ALLOW_CONFIG["C"]["AIAssist"] = True; S.best_l2_cached.cache_clear()
print("K6 C,B=5 baseline  :", round(r0["F1"], 6), fmt(r0))
print("K6 C,B=5 no AIAssist:", round(r1["F1"], 6), fmt(r1))

# K7
for row in b["g10_full_model_sensitivity"]:
    print("K7 tau", row["tau_P2"], "F1", row["F1"], "P2", row["P2_configuration"],
          "AI d", row["P2_AI_recommended_decision"], "adopted d", row["P2_adopted_decision"], "psi", row["P2_reported_psi"])
ho = [(r["budget"], r["governance"], p["process"], p["AI_recommended_decision"], p["adopted_decision"], p["reported_psi"])
      for r in b["results"] for p in r["processes"] if p["configuration"] == "HumanOnly"]
print("K7 HumanOnly rows in 05 (B, G, proc, AI d, adopted d, psi):", ho)

# K8
V = load(P("04_solver_v0_9.py"), "variant")
V.EA_req = lambda A, risk, task_class: 0 if A == 0 else min(3, max(0, A + risk - 2))
nstates = lambda mod: sum(len(mod.canonical_state_set(i, c)) for i in range(3) for c in itertools.product(L, L, L, L))
base = {(r["budget"], r["governance"]): r for r in b["results"]}
changed = 0
for B in (3, 5, 7):
    for g in "CFD":
        r = V.public_result(V.solve(g, B)); o = base[(B, g)]
        key = lambda x: (x["F1_best"], x["L1_allocation_vector"], [(p["A"], p["H"], p["WI"], p["EA"], p["configuration"]) for p in x["processes"]])
        changed += key(r) != key(o)
print("K8 canonical states: code rule", nstates(S), "| v0.7-literal P008 rule (no floor of 1; now superseded)", nstates(V), "| baseline optima changed:", changed, "/ 9")

# K9
vi, vr = set(), set()
for i, p in enumerate(S.PROCESSES):
    for caps in itertools.product(L, L, L, L):
        for (A, H, WI, EA) in S.canonical_state_set(i, caps):
            for fac in ([1.0] if not p["selection"] else [1.0, 0.35]):
                for qn, qd in S.Q_DEFS.items():
                    if A < qd["A"] or H < qd["H"] or WI < qd["WI"] or EA < qd["EA"]: continue
                    v = round((1 + .05*EA + .03*WI)*fac, 6); vi.add(v)
                    if qn != "HumanOnly": vr.add(v)
print("K9 psi internal range", (min(vi), max(vi)), "| reportable (non-HumanOnly)", (min(vr), max(vr)),
      "| reported in nine runs", sorted({p["reported_psi"] for r in b["results"] for p in r["processes"] if p["reported_psi"] is not None}))

# K10
sizes = []
for B in (3, 5, 7):
    for g in "CFD":
        vals = []
        for al in S.generate_allocations(B):
            l2 = S.best_l2_cached(g, S.caps_from_alloc(al))
            if l2: vals.append(round(S.f1_for(al, l2)[0], 12))
        sizes.append(vals.count(max(vals)))
print("K10 L1 argmax-set sizes (nine scenarios):", sizes)

# K11
R = load(P("06_reimplementation_consistency_check_v0_9.py"), "reimpl")
bad = 0
for B in (3, 5, 7):
    for g in "CFD":
        r = R.solve(g, B); o = base[(B, g)]
        for x, op in zip(r["opts"], o["processes"]):
            q = x["l3"]["q"]
            mine = (x["A"], x["H"], x["WI"], x["EA"], q, None if q == "HumanOnly" else round(x["l3"]["psi"], 6),
                    None if q == "HumanOnly" else x["l3"]["decision"], round(x["l3"]["utility"], 6))
            theirs = (op["A"], op["H"], op["WI"], op["EA"], op["configuration"], op["reported_psi"], op["adopted_decision"], op["local_utility"])
            bad += mine != theirs
print("K11 per-process field mismatches (06 functions vs 05):", bad)

# K12 degrees from manuscript Table 3; observed counts from manuscript Table 5
M = 2180
deg = {"HUMAN_AUGMENTATION": 166, "PREDICTION_RISK_ESTIMATION": 462, "COORDINATION_EXECUTION": 219, "GENERATION_SYNTHESIS": 245,
       "EXPLANATION_ASSURANCE": 279, "EVALUATION_ASSESSMENT": 294, "RECOMMENDATION_OPTIMIZATION": 373,
       "PREDICTION_DIAGNOSIS": 340, "HUMAN_REVIEW_ACCOUNTABILITY": 212, "RISK_ASSESSMENT": 167, "ALLOCATION_PLANNING": 294,
       "EVALUATION_SELECTION": 128, "GOVERNANCE_OVERSIGHT": 420, "INFORMATION_ACQUISITION": 322, "GENERAL_DECISION_PROCESS": 97}
obs = dict(zip([f"G{i:02d}" for i in range(1, 17)], [3, 6, 3, 7, 7, 9, 2, 8, 26, 8, 12, 21, 15, 8, 7, 10]))
ws = openpyxl.load_workbook(P("08_16Depletion_Quantification_v0_9.xlsx"))["Depletion_Quantification"]
match = 0
for row in list(ws.iter_rows(values_only=True))[1:]:
    gid, a_, d_, _, oe = row[:5]
    match += abs(round(obs[gid] / (deg[a_]*deg[d_]/M), 4) - float(oe)) < 1e-9
print("K12 workbook O/E equal to manuscript-derived O/E:", f"{match}/16")

# K13
wb = openpyxl.load_workbook(P("08_16Depletion_Quantification_v0_9.xlsx"))
bad = []
for r in list(wb["Configuration_Rows"].iter_rows(values_only=True))[1:]:
    q = S.Q_DEFS[r[0]]
    if [r[1], r[2], r[3], r[4]] != [q["A"], q["H"], q["WI"], q["EA"]] or [float(x) for x in r[5:8]] != q["gain"] \
       or float(r[8]) != q["burden"] or float(r[9]) != q["autonomy"]:
        bad.append(r[0])
for r in list(wb["Governance_Scenarios"].iter_rows(values_only=True))[1:]:
    g = S.GOVERNANCE[r[0]]
    if [float(x) for x in r[1:5]] != [g["K"], g["risk_mult"], g["cost_mult"], g["flex_mult"]] or any(int(x) != 1 for x in r[5:10]):
        bad.append(r[0])
for r, p in zip(list(wb["Process_Data"].iter_rows(values_only=True))[1:], S.PROCESSES):
    doc = (r[0], r[2], r[3], r[4], r[5], r[6], float(r[7]), str(r[8]), str(r[9]), float(str(r[10]).split()[0]), r[11], r[12], r[13], r[14])
    code = (p["id"], p["risk"], p["D0"], p["R0"], p["A0"], p["G0"], p["value"], str(p["selection"]), str(p["execution"]), p["tau"],
            p["score"], p["truth"], p["use_mode"], p["task_class"])
    if doc != code: bad.append((doc, code))
print("K13 workbook tables vs solver constants, mismatches:", bad or "none")

# K14 documented formulas, typed from 01 / Objective_Components, evaluated independently
def clip(x): return max(0.0, min(1.0, x))
mism = n = 0
for gname, g in S.GOVERNANCE.items():
    for i, p in enumerate(S.PROCESSES):
        k = p["risk"]
        for caps in itertools.product(L, L, L, L):
            Dc, Ac, Rc, Gc = caps
            for o in S.l2_options_for_proc(i, caps, g):
                A, H, WI, EA = o["A"], o["H"], o["WI"], o["EA"]
                hr, er = Hreq(A, k), Ereq_code(A, k)
                feas = clip(.55 + .08*(Dc - A) + .08*(Rc - A) + .04*(Gc - max(hr, er)))
                rel = clip(.45 + .10*EA + .06*H + .04*WI - .030*A*k)
                sc = clip(.30 + .08*Dc + .08*Rc + .04*WI)
                ic = g["cost_mult"]*(.10*A + .06*WI + .05*H + .05*EA)
                gr = 0 if A == 0 else g["risk_mult"]*.08*k*A/(1 + .4*H + .4*EA)
                if p["selection"]:
                    d = int(p["score"] >= p["tau"]); V = 1.0 if d == p["truth"] else .35
                else:
                    V = 1.0
                psi = (1 + .05*EA + .03*WI)*V
                best = None
                for qn, qd in S.Q_DEFS.items():
                    if A < qd["A"] or H < qd["H"] or WI < qd["WI"] or EA < qd["EA"]: continue
                    U = qd["gain"][i]*p["value"]*psi*g["flex_mult"] - .45*(qd["burden"] + .03*H + .02*EA + .02*WI) \
                        - .75*(k*qd["autonomy"]*g["risk_mult"]/(1 + .45*H + .45*EA))
                    best = U if best is None else max(best, U)
                f2 = .90*feas + .90*rel + .45*sc - .55*ic - .75*gr + .65*max(0, best)
                cu = .42*A + .24*H + .24*EA + .20*WI
                n += 1; mism += abs(f2 - o["f2"]) > 1e-12 or abs(cu - o["capuse"]) > 1e-12
print("K14 documented-formula F2/CapUse vs solver, options checked:", n, "mismatches:", mism)

# K15
t01 = open(P("01_CANONICAL_MODEL_v0_9.md"), encoding="utf-8").read()
pt = {r[0]: r[7] for r in list(wb["Parameter_Table"].iter_rows(values_only=True))[1:]}
src4 = open(P("04_solver_v0_9.py"), encoding="utf-8").read()
print("K15 01 H_req adj inside:", "min(3,max(0,A+kappa-2+adj_u))" in t01, "| 01 EA_req adj inside:", "min(3,max(1,A+kappa-2+adj_t))" in t01,
      "| P007:", "max(0,A+κ−2+adj_u)" in pt["P007"], "| P008:", "max(1,A+κ−2+adj_t)" in pt["P008"],
      "| code:", "max(0, A + risk - 2 + mode_adjust)" in src4 and "max(1, A + risk - 2 + task_adjust)" in src4)

# K16
order = ["HumanOnly", "AIAssist", "HumanFirstAICheck", "AIFirstReview", "BoundedAutomation"]
t02 = open(P("02_FORMAL_PROOF_v0_9.md"), encoding="utf-8").read()
tb = {r[0]: r[4] for r in list(wb["Tie_Breaking"].iter_rows(values_only=True))[1:]}
print("K16 solver order == stated:", S.CONFIG_ORDER == order, "| 01:", ", ".join(order) in t01, "| 02:", ", ".join(order) in t02,
      "| 08 T3:", ", ".join(order) in tb["T3"], "| 08 T2 uses T3 rank:", "T3" in tb["T2"], "| 05:", b.get("tie_breaking_T3_order") == order)

# K17
csvrows = list(csv.DictReader(open(P("11_Priority_Gaps_Frozen_Export_G01_G16_20261009.csv"), encoding="utf-8-sig")))
fz = list(wb["Priority_Gaps_Freeze"].iter_rows(values_only=True)); hdr = fz[0]; fz = [dict(zip(hdr, r)) for r in fz[1:]]
dq = {r[0]: r for r in list(wb["Depletion_Quantification"].iter_rows(values_only=True))[1:]}
cols = ["Gap_ID", "AI_Family", "Decision_Family", "Observed", "Null_Mean", "O_E_Ratio", "q_BH", "Direction", "Official_Robust_Depletion", "Source_Spreadsheet", "Source_Tab", "Export_Date"]
def norm(v):
    if isinstance(v, bool): return str(v)
    try: return round(float(v), 6)
    except (TypeError, ValueError): return str(v)
m_fz = sum(all(norm(c[k]) == norm(f[k]) for k in cols) for c, f in zip(csvrows, fz))
m_dq = sum(dq[c["Gap_ID"]][1] == c["AI_Family"] and dq[c["Gap_ID"]][2] == c["Decision_Family"] and abs(float(dq[c["Gap_ID"]][4]) - float(c["O_E_Ratio"])) < 1e-9
           and abs(float(dq[c["Gap_ID"]][5]) - float(c["q_BH"])) < 1e-9 for c in csvrows)
print("K17 CSV rows:", len(csvrows), "| CSV == Priority_Gaps_Freeze on all 12 columns:", f"{m_fz}/16", "| CSV codes/O/E/q == Depletion_Quantification:", f"{m_dq}/16")

# K18
def hyp_sd(ka, kd): return (ka*(kd/M)*(1-kd/M)*(M-ka)/(M-1))**0.5
lab = {}
for c in csvrows:
    ka, kd, O, nm, oe = deg[c["AI_Family"]], deg[c["Decision_Family"]], float(c["Observed"]), float(c["Null_Mean"]), float(c["O_E_Ratio"])
    E = ka*kd/M; z = (nm - E)/(hyp_sd(ka, kd)/100)   # SE of the mean over 10,000 replicates
    lab[c["Gap_ID"]] = (round(O/E, 4) == round(oe, 4), round(O/nm, 4) == round(oe, 4), round(z, 2))
print("K18 O_E_Ratio == O/E_analytic:", sum(v[0] for v in lab.values()), "/16 | == O/Null_Mean:", sum(v[1] for v in lab.values()), "/16")
print("K18 (Null_Mean - E_analytic) in Monte Carlo SE units:", {k: v[2] for k, v in lab.items()})
print("K18 max |z|:", max(abs(v[2]) for v in lab.values()))

# K19 manuscript Table 5 (main, exact, document-preserving, raw-record q); "< .0001" entered as 0.0001
ms = {"G01": (3, .0009, .0001, .0024, .0001), "G02": (6, .0009, .0001, .0024, .0001), "G03": (3, .0009, .0002, .0024, .0001),
      "G04": (7, .0009, .0001, .0042, .0001), "G05": (7, .0009, .0001, .0024, .0001), "G06": (9, .0009, .0001, .0024, .0001),
      "G07": (2, .0197, .0147, .0304, .0290), "G08": (8, .0009, .0001, .0024, .0001), "G09": (26, .0009, .0001, .0024, .0001),
      "G10": (8, .0009, .0001, .0058, .0001), "G11": (12, .0009, .0001, .0024, .0001), "G12": (21, .0009, .0001, .0024, .0001),
      "G13": (15, .0009, .0001, .0024, .0001), "G14": (8, .0045, .0032, .0089, .0001), "G15": (7, .0206, .0232, .0387, .0147),
      "G16": (10, .0169, .0168, .0137, .0086)}
ok_obs = sum(float(c["Observed"]) == ms[c["Gap_ID"]][0] for c in csvrows)
ok_q = sum(abs(float(c["q_BH"]) - ms[c["Gap_ID"]][1]) < 5e-5 for c in csvrows)
ok_four = sum(all(q <= .05 for q in ms[g][1:]) for g in ms)
print("K19 Observed == manuscript:", ok_obs, "/16 | primary q == manuscript (4 dp):", ok_q, "/16 | all four manuscript q <= .05:", ok_four,
      "/16 | CSV Direction DEPLETED & robust flag True:", sum(c["Direction"] == "DEPLETED" and c["Official_Robust_Depletion"] == "True" for c in csvrows), "/16",
      "| other three q columns in CSV:", [h for h in csvrows[0] if "q" in h.lower()])

# K20
stale = []
for ws in wb.worksheets:
    if ws.title.startswith("Audit_Corrections_v0.8"): continue
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "v0.8" in cell.value and "Claude v0.8" not in cell.value:
                stale.append(f"{ws.title}!{cell.coordinate}")
print("K20 cells still labelled v0.8 (excluding historical sheet and references to the Claude v0.8 audit):", stale)
g10 = [r[10] for r in list(wb["G10_Sensitivity"].iter_rows(values_only=True))[1:]]
print("K20 G10_Sensitivity interpretation wording:", set(g10))
