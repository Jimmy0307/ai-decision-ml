"""Automated test of the S1b Chinese display-label -> canonical-code map (s1b_label_map.json / .py).

Checks (PASS/FAIL per check; overall PASS only if all pass):
  M1  each field's map is a bijection (labels unique after normalization, codes unique);
  M2  each field's code set equals the analyzer's canonical set (Q1/Q2/Q3/Q4) or {COMPLIES, DOES_NOT_COMPLY};
  M3  every drop-down option ACTUALLY PRESENT in each workbook's data validations maps to a code, the option list
      equals the map's list for that field (same labels, same order), and every expected field has a validation;
  M4  NFKC / whitespace / fullwidth / bar variants of every label normalize to exactly one code (its own);
  M5  unknown labels (canonical English codes, near-misses) map to nothing.
Reads the blank workbooks only; writes nothing except an optional result file.

Usage: python -E -P test_s1b_label_map.py [--out-dir OUT] [--result FILE]
"""
from __future__ import annotations

import argparse
import os
import re
import sys

from openpyxl import load_workbook

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import s1b_label_map as LM  # noqa: E402

CANON = {
    "Q1": {"YES", "NO", "UNCERTAIN"},
    "Q2": {"M1", "M2", "M3", "M4", "LONGITUDINAL", "DATA_GATED", "EVIDENCE_ONLY", "NONE", "AMBIGUOUS"},
    "Q3": {"HIGH", "MEDIUM", "LOW"},
    "Q4": {"M1", "M2", "M3", "M4", "LONGITUDINAL", "DATA_GATED"},
    "ATTESTATION": {"COMPLIES", "DOES_NOT_COMPLY"},
}
CODER_FILES = [os.path.join("coder_distribution", f"CODER_{k}_ZH_TW.xlsx") for k in (1, 2, 3)]
MASTER = "20_MASTER_PACKET_S1B_ZH_TW.xlsx"


def fullwidth(s):
    return "".join(chr(ord(c) + 0xFEE0) if "!" <= c <= "~" else c for c in s)


def variants(lab):
    v = {lab, " " + lab + " ", "　" + lab + "　", " " + lab, lab + "\n", fullwidth(lab),
         lab.replace("｜", "|"), lab.replace("｜", " | "), lab.replace("｜", "│"), lab.replace("｜", " ｜ "),
         " ".join(lab), fullwidth(lab.replace("｜", "|"))}
    return sorted(v)


def check_bijection():
    errs = []
    for f in LM.FIELDS:
        labs, cods = LM.options(f), LM.codes(f)
        if len({LM.normalize(x) for x in labs}) != len(labs):
            errs.append(f"{f}: labels not unique after normalization")
        if len(set(cods)) != len(cods):
            errs.append(f"{f}: codes not unique")
        if set(cods) != CANON[f]:
            errs.append(f"{f}: code set {sorted(cods)} != canonical {sorted(CANON[f])}")
    return errs


def dv_fields(path):
    """{(sheet, field): [options]} from the workbook's list data validations."""
    wb = load_workbook(path)
    found = {}
    for ws in wb.worksheets:
        head = [str(c.value) if c.value is not None else "" for c in ws[1]] if ws.max_row else []
        for dv in ws.data_validations.dataValidation:
            if dv.type != "list":
                continue
            opts = dv.formula1.strip('"').split(",")
            cols = {re.match(r"[A-Z]+", str(r).split(":")[0]).group(0) for r in dv.sqref.ranges}
            if ws.title == "作答聲明":
                field = "ATTESTATION"
            else:
                idx = [ws[c + "1"].column - 1 for c in cols]
                fields = {head[i][:2] for i in idx}
                field = fields.pop() if len(fields) == 1 else "?"
            found.setdefault((ws.title, field), []).append(opts)
    return found


def check_workbooks(out):
    errs, n_opts = [], 0
    for rel in CODER_FILES + [MASTER]:
        p = os.path.join(out, rel)
        found = dv_fields(p)
        fields_present = {f for _, f in found}
        if not {"Q1", "Q2", "Q3", "Q4", "ATTESTATION"} <= fields_present:
            errs.append(f"{rel}: missing validation for {sorted({'Q1','Q2','Q3','Q4','ATTESTATION'} - fields_present)}")
        for (sheet, f), lists in found.items():
            if f not in CANON:
                errs.append(f"{rel}!{sheet}: list validation on unknown field {f}")
                continue
            for opts in lists:
                n_opts += len(opts)
                if opts != LM.options(f):
                    errs.append(f"{rel}!{sheet} {f}: options differ from map")
                codes = [LM.to_code(f, o) for o in opts]
                if None in codes:
                    errs.append(f"{rel}!{sheet} {f}: unmapped option(s)")
                if len(set(codes)) != len(codes) or set(codes) != CANON[f]:
                    errs.append(f"{rel}!{sheet} {f}: options not a bijection onto canonical codes")
    return errs, n_opts


def check_variants():
    errs, n = [], 0
    for f in LM.FIELDS:
        tab = LM.table(f)
        for lab, code in zip(LM.options(f), LM.codes(f)):
            for v in variants(lab):
                n += 1
                hits = {c for k, c in tab.items() if k == LM.normalize(v)}
                if hits != {code} or LM.to_code(f, v) != code:
                    errs.append(f"{f}: variant {v!r} -> {hits}")
    return errs, n


def check_unknown():
    errs = []
    for f in LM.FIELDS:
        for bad in list(CANON[f]) + ["M5", "是的", "非常高", "符合。", "OTHER", "x"]:
            if bad in LM.options(f):
                continue
            if LM.to_code(f, bad) is not None:
                errs.append(f"{f}: unknown {bad!r} mapped")
    return errs


def run(out):
    res = {}
    res["M1_M2_bijection_and_canonical_sets"] = check_bijection()
    e3, n_opts = check_workbooks(out)
    res["M3_workbook_dropdowns_map"] = e3
    e4, n_var = check_variants()
    res["M4_variants_normalize_to_one_code"] = e4
    res["M5_unknown_labels_unmapped"] = check_unknown()
    ok = not any(res.values())
    line = ("PASS" if ok else "FAIL") + " " + " ".join(f"{k}={'PASS' if not v else 'FAIL'}" for k, v in res.items()) + \
           f" dropdown_options_checked={n_opts} variants_checked={n_var}"
    return ok, line, res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir", default=os.path.dirname(HERE))
    ap.add_argument("--result")
    a = ap.parse_args()
    ok, line, res = run(a.out_dir)
    print(line)
    for k, v in res.items():
        for e in v[:20]:
            print("  ", k, e)
    if a.result:
        with open(a.result, "w", encoding="utf-8", newline="\n") as f:
            f.write(line + "\n")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
