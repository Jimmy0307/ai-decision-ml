"""Compatibility / correctness test for v1_construct_analyzer_s1b.py. No answer-filled workbook is created.

(a) reliability core (unchanged reliability_v2_2 functions, hash-checked) vs sklearn cohen_kappa_score and
    krippendorff.alpha on random label vectors held in memory; classification rule boundaries; OA-3 NONE rule;
    in-memory parsing semantics: Chinese labels -> codes feed validate(); DOES_NOT_COMPLY -> EXPOSED.
(b) label-map test (test_s1b_label_map.run).
(c) analyzer run on the 3 blank S1b coder files -> NO_HUMAN_DATA (outputs to tools/no_human_data_run_s1b/).
(d) structural parse of the blank files: sheet names, headers, 16 items, empty responses/attestation, validations.
Prints PASS/FAIL only (and max abs differences for (a)).

Run: .venv_docx/Scripts/python.exe -E -P test_v1_construct_analyzer_s1b.py [--repo REPO] [--result FILE]
"""
from __future__ import annotations

import argparse
import io
import os
import random
import sys
from contextlib import redirect_stdout

import krippendorff
import numpy as np
from openpyxl import load_workbook
from sklearn.metrics import cohen_kappa_score

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import build_v1_packet_s1b as BLD  # noqa: E402  (constants only; build() is not called)
import s1b_label_map as LM  # noqa: E402
import test_s1b_label_map as TLM  # noqa: E402
import v1_construct_analyzer_s1b as A  # noqa: E402

TOL = 1e-9
FILES = [os.path.join(OUT, "coder_distribution", f"CODER_{k}_ZH_TW.xlsx") for k in (1, 2, 3)]


def stats_test(rel):
    rng = random.Random(1829374650)
    cats_all = sorted(A.Q2)
    dk = da = 0.0
    nk = na = 0
    for _ in range(300):
        n, k = rng.choice([8, 16, 30]), rng.choice([2, 3, 4])
        cats = cats_all[:rng.randint(2, len(cats_all))]
        data = [[rng.choice(cats) for _ in range(k)] for _ in range(n)]
        x, y = [u[0] for u in data], [u[1] for u in data]
        mk, _, _ = rel.cohen_kappa(x, y)
        rk = cohen_kappa_score(x, y)
        if mk == mk and rk == rk:
            dk, nk = max(dk, abs(mk - rk)), nk + 1
        if len({v for u in data for v in u}) > 1:
            enc = {c: j for j, c in enumerate(cats)}
            arr = np.array([[enc[v] for v in u] for u in data], dtype=float).T
            da = max(da, abs(rel.kripp_alpha_nominal(data)
                             - krippendorff.alpha(reliability_data=arr, level_of_measurement="nominal")))
            na += 1
    return dk, nk, da, na


def rule_test():
    c = A.classify_primary
    return (c(0.800, 0.667)["label"] == "RELIABLE" and not c(0.800, 0.667)["imprecise"]
            and c(0.7999, 0.70)["label"] == "TENTATIVE" and c(0.667, 0.667)["label"] == "TENTATIVE"
            and c(0.6669, 0.70)["label"] == "UNRESOLVED/NOT_ACCEPTABLE" and c(0.90, 0.6669)["imprecise"]
            and c(float("nan"), float("nan"))["label"] == "UNRESOLVED")


def oa3_test():
    def one(q1, q2, q5):
        row = {A.RESP[0]: q1, A.RESP[1]: q2, A.RESP[2]: "LOW", A.RESP[3]: "", A.RESP[4]: q5}
        return A.validate(dict(sheet="T", rows={"ITEM-A01": row}, attestation=[]))
    e1, e2, e3 = one("UNCERTAIN", "NONE", "x"), one("YES", "NONE", ""), one("YES", "NONE", "x")
    return (any("NONE requires Q1=YES" in e for e in e1) and any("NONE requires Q5" in e for e in e2)
            and not any("NONE" in e for e in e3))


def parse_semantics_test():
    """In memory: Chinese labels go through the analyzer's _code() and reach the unchanged validate()."""
    q = {A.RESP[0]: A._code("Q1", "是"), A.RESP[1]: A._code("Q2", " 僅屬證據層 "), A.RESP[2]: A._code("Q3", "高"),
         A.RESP[3]: A._code("Q4", "M1 | 權責與問責"), A.RESP[4]: "x"}
    ok = q[A.RESP[0]] == "YES" and q[A.RESP[1]] == "EVIDENCE_ONLY" and q[A.RESP[3]] == "M1"
    errs = A.validate(dict(sheet="T", rows={"ITEM-A01": q}, attestation=[]))
    ok = ok and any("Q1=YES with Q2=EVIDENCE_ONLY" in e for e in errs)
    bad = A._code("Q2", "YES")                                    # English code typed instead of label -> kept raw
    ok = ok and bad == "YES" and any("Q2 'YES' invalid" in e for e in A.validate(
        dict(sheet="T", rows={"ITEM-A01": {**q, A.RESP[0]: "UNCERTAIN", A.RESP[1]: bad}}, attestation=[])))
    s_dnc = dict(attestation=[("聲明一", "COMPLIES"), ("聲明二", "DOES_NOT_COMPLY"), ("姓名", "x")])
    s_ok = dict(attestation=[("聲明一", "COMPLIES"), ("姓名", "x")])
    return ok and A.attestation_exposed(s_dnc) == ["聲明二"] and A.attestation_exposed(s_ok) == []


def blank_run_test(repo):
    out = os.path.join(HERE, "no_human_data_run_s1b")
    buf = io.StringIO()
    with redirect_stdout(buf):
        code = A.run(FILES + ["--out", out] + (["--repo", repo] if repo else []))
    return code == 0 and buf.getvalue().strip() == "NO_HUMAN_DATA"


def structure_test():
    errs = []
    exp_head = BLD.CODER_HEAD + [BLD.Q_TEXT[k] for k in BLD.RESP_KEYS]
    for k, p in enumerate(FILES, 1):
        c = f"CODER_{k}"
        wb = load_workbook(p)
        want = [BLD.SHEET_README, BLD.SHEET_CATS, BLD.SHEET_INFO, BLD.SHEET_CODER[c], BLD.SHEET_ATTEST]
        if wb.sheetnames != want:
            errs.append(f"{c}: sheets {wb.sheetnames}")
        ws = wb[BLD.SHEET_CODER[c]]
        if [x.value for x in ws[1]] != exp_head:
            errs.append(f"{c}: header mismatch")
        ids = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
        if sorted(ids) != A.ITEMS or ws.max_row != 17:
            errs.append(f"{c}: item set/rows")
        if [ws.cell(r, 1).value for r in range(2, 18)] != list(range(1, 17)):
            errs.append(f"{c}: item order column")
        dvs = {str(d.sqref): d for d in ws.data_validations.dataValidation}
        exp = {"H2:H17": "Q1", "I2:I17": "Q2", "J2:J17": "Q3", "K2:K17": "Q4"}
        for rng, f in exp.items():
            d = dvs.get(rng)
            if d is None or d.type != "list" or d.formula1 != '"' + ",".join(LM.options(f)) + '"' or not d.showErrorMessage:
                errs.append(f"{c}: validation {f} {rng}")
        d5 = dvs.get("L2:L17")
        if d5 is None or d5.type != "textLength" or d5.formula1 != "600":
            errs.append(f"{c}: validation Q5")
        if len(dvs) != 5:
            errs.append(f"{c}: {len(dvs)} validations on coder sheet")
        wa = wb[BLD.SHEET_ATTEST]
        st = [wa.cell(r, 2).value for r in range(3, wa.max_row + 1)]
        if st != BLD.ATTEST_ITEMS + BLD.ATTEST_FREE or any(wa.cell(r, 3).value is not None for r in range(3, 13)):
            errs.append(f"{c}: attestation rows")
        adv = wa.data_validations.dataValidation
        if len(adv) != 1 or str(adv[0].sqref) != " ".join(f"C{r}" for r in range(3, 10)) or not adv[0].showErrorMessage:
            errs.append(f"{c}: attestation validation {[str(d.sqref) for d in adv]}")
        s = A.read_sheet(p)
        if s["sheet"] != c or sorted(s["rows"]) != A.ITEMS or not A.is_empty(s) or len(s["attestation"]) != 10:
            errs.append(f"{c}: analyzer parse")
        # blank sheet: validate() must report exactly "Q1/Q2/Q3 '' invalid" + "Q5 rationale empty" per item and
        # "attestation incomplete" once (proves every required response/attestation cell is parsed as empty)
        ve = A.validate(s)
        expect = {f"{c} {i}: {m}" for i in A.ITEMS
                  for m in ("Q1 '' invalid", "Q2 '' invalid", "Q3 '' invalid", "Q5 rationale empty")}
        expect.add(f"{c}: attestation incomplete")
        if len(ve) != len(expect) or set(ve) != expect:
            errs.append(f"{c}: unexpected validation errors on blank")
    return errs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", default=None)
    ap.add_argument("--result")
    a = ap.parse_args()
    rel = A.load_rel(a.repo or A.default_repo())
    dk, nk, da, na = stats_test(rel)
    r_ok, o_ok, p_ok = rule_test(), oa3_test(), parse_semantics_test()
    a_ok = dk <= TOL and da <= TOL and nk > 0 and na > 0 and r_ok and o_ok and p_ok
    b_ok, b_line, _ = TLM.run(OUT)
    c_ok = blank_run_test(a.repo)
    d_errs = structure_test()
    d_ok = not d_errs
    ok = a_ok and b_ok and c_ok and d_ok
    line = (f"{'PASS' if ok else 'FAIL'} "
            f"(a)reliability_core={'PASS' if a_ok else 'FAIL'} kappa_vs_sklearn_max_abs_diff={dk:.3e} (n={nk}) "
            f"alpha_vs_krippendorff_max_abs_diff={da:.3e} (n={na}) classification_rule={'PASS' if r_ok else 'FAIL'} "
            f"oa3_none_validation={'PASS' if o_ok else 'FAIL'} s1b_parse_semantics={'PASS' if p_ok else 'FAIL'} | "
            f"(b)label_map={'PASS' if b_ok else 'FAIL'} | (c)blank_run_NO_HUMAN_DATA={'PASS' if c_ok else 'FAIL'} | "
            f"(d)structure={'PASS' if d_ok else 'FAIL'}")
    print(line)
    for e in d_errs:
        print("  ", e)
    if a.result:
        with open(a.result, "w", encoding="utf-8", newline="\n") as f:
            f.write(line + "\n" + "(b) " + b_line + "\n")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
