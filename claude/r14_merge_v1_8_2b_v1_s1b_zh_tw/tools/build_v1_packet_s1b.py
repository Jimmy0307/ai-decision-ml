"""Build the V1 S1b Traditional Chinese (ZH_TW) coder instrument (pre-human; researcher-approved S1b).

Design (orchestrator + researcher decision, see S1B_PROTOCOL.md):
- SAME neutral item IDs ITEM-A01..A16 with the SAME item content per ID as S1. Item content is read from the S1
  master packet ITEM_INFO (read-only; S1 hashes verified first) and cross-checked against the three S1 coder
  sheets. No G ID is read or written. The sealed ITEM<->G map is NOT opened (its commitment stays valid unchanged).
- New per-coder item orders from new, recorded, deterministic seeds; the three orders must differ from each other,
  from the identity order and from every S1 coder order (read from the public S1 coder files).
- All coder-facing text in Traditional Chinese; drop-down labels come from s1b_label_map.json.
- Writes: OUT/coder_distribution/CODER_k_ZH_TW.xlsx, OUT/20_MASTER_PACKET_S1B_ZH_TW.xlsx,
  OUT/S1B_ITEM_CONTENT_CORRESPONDENCE.csv, OUT/V1_SEALED/S1B_ITEM_ORDER_LOG.COMMITMENT.md,
  STUDY_TEAM/S1B_ITEM_ORDER_LOG.json (plaintext order log; study team only, not in git).
Never writes responses; never simulates coders; never reads the researcher key, the item map or AI-pilot outputs.

Usage: python -E -P build_v1_packet_s1b.py --repo REPO --study-team DIR
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import random
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import s1b_label_map as LM  # noqa: E402

OUT_REL = os.path.join("claude", "r14_merge_v1_8_2b_v1_s1b_zh_tw")
S1_REL = os.path.join("claude", "r14_merge_v1_8_2_v1_prefreeze")
CODERS = ["CODER_1", "CODER_2", "CODER_3"]
NEUTRAL = [f"ITEM-A{i:02d}" for i in range(1, 17)]
SEEDS = {c: f"S1B_ZH_TW_20261014_{c}_ORDER" for c in CODERS}       # recorded, deterministic, non-secret
CODER_ZH = {c: f"編碼者{c[-1]}" for c in CODERS}
SHEET_README, SHEET_CATS, SHEET_INFO, SHEET_ATTEST = "填答說明", "類別定義", "題目資訊", "作答聲明"
SHEET_CODER = {c: f"作答_{CODER_ZH[c]}" for c in CODERS}

# ---------------------------------------------------------------- family translations (faithful; keyed by S1 English)
AI_FAM = {
    "Human Augmentation": ("HUMAN_AUGMENTATION", "人機增能",
        "AI functions whose main role is to assist, enhance, support or extend human judgment, cognition, expertise "
        "or work, rather than independently replacing the human decision role.",
        "主要角色是協助、強化、支援或延伸人的判斷、認知、專業或工作，而不是獨立取代人在決策中之角色的 AI 功能。"),
    "Prediction and Risk Estimation": ("PREDICTION_RISK_ESTIMATION", "預測與風險估計",
        "AI functions that estimate future states, outcomes, probabilities, uncertainty, hazards, failures, or "
        "quantified risk.",
        "估計未來狀態、結果、機率、不確定性、危害、故障或量化風險的 AI 功能。"),
    "Coordination and Execution": ("COORDINATION_EXECUTION", "協調與執行",
        "AI functions that coordinate agents, tools, processes, workflows, tasks or actions, or execute actions in an "
        "operating environment.",
        "協調代理者、工具、程序、工作流程、任務或行動，或在作業環境中執行行動的 AI 功能。"),
    "Generation and Synthesis": ("GENERATION_SYNTHESIS", "生成與合成",
        "AI functions that create, compose, synthesize, summarize, transform or generate new information, text, "
        "visuals, representations or structured content.",
        "創造、撰寫、綜合、摘要、轉換或生成新的資訊、文字、影像、表徵或結構化內容的 AI 功能。"),
    "Explanation and Assurance": ("EXPLANATION_ASSURANCE", "解釋與保證",
        "AI functions that make model outputs, reasoning, uncertainty, confidence, provenance, reliability or "
        "behaviour understandable, traceable, verifiable or assessable by humans or downstream systems.",
        "使模型的輸出、推理、不確定性、信心程度、來源、可靠性或行為，能被人或下游系統理解、追溯、驗證或評估的 AI 功能。"),
    "Evaluation and Audit": ("EVALUATION_ASSESSMENT", "評估與審核",
        "AI functions that assess, check, validate, audit or inspect the quality, performance, correctness, "
        "compliance, consistency or reliability of data, models, outputs or processes.",
        "評估、檢查、驗證、稽核或檢驗資料、模型、輸出或程序之品質、效能、正確性、合規性、一致性或可靠性的 AI 功能。"),
    "Recommendation and Optimization": ("RECOMMENDATION_OPTIMIZATION", "推薦與最佳化",
        "AI functions that recommend, rank, prescribe, search or optimize preferred options, actions, configurations, "
        "solutions or resource use.",
        "推薦、排序、提出處方式建議、搜尋或最佳化較佳的選項、行動、配置、解決方案或資源使用的 AI 功能。"),
}
DEC_FAM = {
    "Prediction and Diagnosis": ("PREDICTION_DIAGNOSIS", "預測與診斷",
        "Decision functions that use prediction, diagnosis, prognosis, state classification or state estimation to "
        "determine current or future decision-relevant conditions.",
        "運用預測、診斷、預後判斷、狀態分類或狀態估計，來判定目前或未來與決策相關之狀況的決策功能。"),
    "Human Review and Accountability": ("HUMAN_REVIEW_ACCOUNTABILITY", "人工複核與問責",
        "Decision functions requiring explicit human review, approval, escalation, assignment of responsibility, "
        "traceability, appeal, sign-off or accountability.",
        "需要明確的人工複核、核准、向上呈報、責任指派、可追溯性、申訴、簽核或問責的決策功能。"),
    "Risk Assessment": ("RISK_ASSESSMENT", "風險評估",
        "Decision functions that evaluate, score, compare, prioritize or characterize risks, hazards, vulnerability, "
        "exposure, uncertainty or safety risk.",
        "評估、評分、比較、排定優先順序或描述風險、危害、脆弱性、暴露程度、不確定性或安全風險的決策功能。"),
    "Allocation and Planning": ("ALLOCATION_PLANNING", "配置與規劃",
        "Decision functions that allocate resources or plan future actions, schedules, capacity, budgets, routes, "
        "staffing, inventory, production, portfolios or other operating resources.",
        "分配資源，或規劃未來的行動、排程、產能、預算、路線、人員配置、庫存、生產、投資組合或其他營運資源的決策功能。"),
    "Evaluation and Selection": ("EVALUATION_SELECTION", "評估與選擇",
        "Decision functions that evaluate, compare, rank, prioritize, screen or select among options.",
        "在多個選項之間進行評估、比較、排序、排定優先順序、篩選或選擇的決策功能。"),
    "Governance and Oversight": ("GOVERNANCE_OVERSIGHT", "治理與監督",
        "Decision functions that establish, monitor, enforce or review policies, authority, controls, governance "
        "mechanisms, compliance requirements or organizational oversight.",
        "建立、監控、執行或審查政策、權限、控制措施、治理機制、合規要求或組織監督的決策功能。"),
    "Information Acquisition": ("INFORMATION_ACQUISITION", "資訊取得",
        "Decision functions that collect, obtain, retrieve, gather, search for or request the information needed "
        "before or during a decision.",
        "在決策之前或決策過程中，蒐集、取得、檢索、彙集、搜尋或請求所需資訊的決策功能。"),
    "General Decision Process": ("GENERAL_DECISION_PROCESS", "一般決策流程",
        "Decision functions that explicitly concern the overall decision process, decision-support process or "
        "decision workflow and cannot be assigned to a more specific decision function (residual category).",
        "明確涉及整體決策過程、決策支援過程或決策工作流程，且無法歸入更具體之決策功能的決策功能（屬剩餘類型）。"),
}
PATTERN_TEXT = ("在本研究分析的文獻中，這兩類功能被一起討論的情形比依整體分布所預期的少。"
                "這只是研究線索，不代表這種組合不可能、無效、不重要，或在實務中不存在。")

# ---------------------------------------------------------------- coder-facing text (Traditional Chinese)
Q_TEXT = {
    "Q1": "Q1 這段內容是否呈現一個可辨識的企業決策介面？（是／否／無法確定）",
    "Q2": "Q2 如果有，這段內容最主要對應哪一種企業決策機制？（單選）",
    "Q3": "Q3 你對 Q2 的判斷有多大信心？（高／中／低）",
    "Q4": "Q4 如果你認為還有第二個合理的機制，可在此選填。（選填）",
    "Q5": "Q5 請簡短說明你做出上述判斷的主要理由。",
}
TIME_TEXT = "約 60–120 分鐘，視閱讀速度而異。"
NO_SEARCH_TEXT = "請只依據本問卷提供的內容作答，不需要也請不要另外搜尋外部資料。"

README_ROWS = [
    "編碼問卷　填答說明（每位編碼者收到的說明內容都相同）",
    "",
    "預計作答時間：" + TIME_TEXT,
    "可以中途休息，但在交回問卷之前，請勿與任何人討論題目。",
    NO_SEARCH_TEXT,
    "請獨立作答。作答期間請勿與其他編碼者討論，也請勿使用任何 AI 工具。若作答前你已接觸過與本研究有關的資料或討論，請在開始前告知研究協調人。",
    "",
    "本問卷共 16 題，每題只以中性的題目編號標示。每題呈現一種「AI 功能類型」與一種「決策功能類型」及其定義，並附上一段所有題目都相同的文獻樣態說明。每一題都回答相同的五個問題。",
    "「企業決策介面」是指：企業針對「這類 AI 功能如何進入這類決策功能」必須做出的一項決策，而且這項決策是這個組合特有的。",
    "",
    Q_TEXT["Q1"],
    "　　是：你能具體說出一項企業必須做出的決策，而且這項決策是這個組合特有的。",
    "　　否：這段內容沒有涉及任何企業決策；或你能說出的決策，幾乎任何 AI 功能進入這個決策功能時都會同樣出現。",
    "　　無法確定：你無法在「是」與「否」之間做出判斷。",
    Q_TEXT["Q2"],
    "　　每一題都必須填寫 Q2。Q1 與 Q2 的搭配規則如下：",
    "　　Q1＝是：Q2 可選「僅屬證據層」以外的任一類別。",
    "　　Q1＝否：Q2 請選「僅屬證據層」。",
    "　　Q1＝無法確定：Q2 可選「無適用機制」以外的任一類別（包括「僅屬證據層」與「無法明確判定」）。",
    "　　各類別的定義與判斷順序，請見「類別定義」工作表。",
    Q_TEXT["Q3"],
    "　　信心只針對 Q2 的判斷。",
    Q_TEXT["Q4"],
    "　　沒有明顯需要時請留空；若填寫，Q4 不可與 Q2 相同。",
    Q_TEXT["Q5"],
    "　　每一題都必須填寫，以 1–3 句話為原則（最多 600 字）。若 Q2 選「無適用機制」，請寫出沒有被任何類別涵蓋的那項企業決策；若選「資料條件限制」，請說明是哪一種資料及其條件；若選「無法明確判定」，請寫出相互競爭的類別，或說明哪裡證據不足。",
    "",
    "作答方式：請在淡黃色的作答欄中使用下拉選單作答（Q5 請直接輸入文字）。請勿修改題目列、題目編號或題目順序。",
    "完成後請填寫「作答聲明」工作表並存檔，再將檔案直接交回給研究協調人。請用電腦上的試算表軟體開啟與存檔，不要上傳到任何線上試算表服務。",
]

M1L, M2L, M3L, M4L, LGL, DGL, EOL, NOL, AML = LM.options("Q2")
CATEGORY_ROWS = [
    (M1L, "核心在誰有決策權、誰負責、誰可核准、否決或停止。",
     "與 M3 的差別：M1＝誰有權、誰負責；M3＝人與 AI 怎麼分工與互動。若核心是 AI 輸出如何轉交、升級或復原，請選 M2；若核心是驗證或監控，請選 M4。"),
    (M2L, "AI 輸出如何進入流程，以及失敗、異常或不確定時如何轉交、升級或復原。",
     "若核心是誰有權核准、否決或停止，請選 M1；若核心是人與 AI 如何分工與互動，請選 M3。"),
    (M3L, "人與 AI 如何分工、互動、採納、修正與接手。",
     "與 M1 的差別：M1＝誰有權、誰負責；M3＝人與 AI 怎麼分工與互動。若核心是驗證或監控，請選 M4。"),
    (M4L, "驗證、評測、監控、稽核、安全、風險與合規。",
     "若核心是誰有權決定或誰負責，請選 M1；若核心是人與 AI 如何分工與互動，請選 M3。"),
    (LGL, "核心在跨時間的學習與適應：前一段時間的結果或回饋會改變後續的決策或做法，或能力會隨時間演進與改變。只看單一時點的配置，不足以描述這項決策。",
     "若單一時點的選擇就足以描述，請改選適合的類別。選擇此類別時，可在 Q4 填入相關的 M1–M4。"),
    (DGL, "這項決策能否成立，主要取決於資料條件：資料是否可取得、品質是否足夠、數量是否充足，以及是否已準備好可用。",
     "請在 Q5 說明是哪一種資料及其條件。若不需要特別的資料條件就能套用某一類別，請選該類別。選擇此類別時，可在 Q4 填入相關的 M1–M4。"),
    (EOL, "只提供資訊或證據，但尚未形成可辨識的企業決策介面。",
     "Q1＝否時，Q2 請選此類別；Q1＝無法確定時也可選。Q1＝是時不可選。若有某一類別明顯適用，表示已形成決策介面，請改選該類別。"),
    (NOL, "只有在存在企業決策介面，但其他所有機制（M1–M4、跨時間學習與適應、資料條件限制）都不適用時才能選。",
     "必須 Q1＝是，且必須在 Q5 寫出沒有被涵蓋的那項企業決策。證據不足或兩類以上同樣合理時，請選「無法明確判定」；沒有可辨識的決策介面時，請選「僅屬證據層」。"),
    (AML, "用於證據不足、兩類以上同樣合理，或無法可靠指定單一主要機制。",
     "請在 Q5 寫出相互競爭的類別，或說明哪裡證據不足。不要把它當成預設選項；只要有一類明顯較適合，就選那一類。Q1＝否時不可選。"),
    ("判斷順序（Q2）", "先回答 Q1。若 Q1＝否，Q2 直接選「僅屬證據層」。若 Q1＝是或無法確定，請依下列順序檢查，遇到第一個適用的就停止。", ""),
    ("判斷順序（一）", "這項決策的核心是否在於跨時間的學習與適應（前一段時間的結果或回饋會改變後續的決策或做法，或能力會隨時間演進）？", "→ 跨時間學習與適應"),
    ("判斷順序（二）", "這項決策能否成立，是否主要取決於資料條件（資料是否可取得、品質是否足夠、數量是否充足、是否已準備好可用）？", "→ 資料條件限制（請在 Q5 說明是哪一種資料及其條件）"),
    ("判斷順序（三）", "是否有某一個 M1–M4 機制最符合這項決策的主要議題？",
     "→ 選最符合的 M1–M4。若存在可辨識的企業決策介面，但 M1–M4、跨時間學習與適應、資料條件限制都不適用 → 無適用機制（僅限 Q1＝是，並須在 Q5 寫出該項決策）"),
    ("判斷順序（四）", "若只提供資訊或證據，尚未形成可辨識的企業決策介面。", "→ 僅屬證據層（此時 Q1 應為「否」或「無法確定」）"),
    ("判斷順序（補充）", "在任何一步，若證據不足、兩類以上同樣合理，或無法可靠指定單一主要機制。",
     "→ 無法明確判定（請在 Q5 寫出相互競爭的類別，或說明哪裡證據不足）"),
    ("一致性規則", "Q1＝是：Q2 不可選「僅屬證據層」。Q1＝否：Q2＝「僅屬證據層」。Q1＝無法確定：Q2 不可選「無適用機制」。Q4 不可與 Q2 相同。選「無適用機制」時，Q1 必須為「是」，且 Q5 必須寫出沒有被涵蓋的企業決策。", ""),
    ("剩餘類型", "「一般決策流程」是剩餘的決策功能類型（不是 Q2 的類別）。若有類別適用，請照常判斷；若沒有，可選「僅屬證據層」或「無法明確判定」。", ""),
    ("判斷對象", "請判斷企業決策本身，而不是技術、文獻主題，或功能類型名稱與類別名稱之間碰巧相同的字詞。部分功能類型名稱與類別名稱有相同的字詞（例如：問責、流程、人機、配置、評估、保證、風險），相同字詞不是選擇該類別的理由。", ""),
    ("Q3 信心", "高／中／低，只針對 Q2 的判斷。", ""),
]

ATTEST_PREAMBLE = ("作答聲明：請逐條閱讀下列敘述，並從下拉選單選擇「符合」（這句敘述與你的實際情況相符）或「不符合」"
                   "（這句敘述與你的實際情況不符）。請如實回答；選擇「不符合」會被記錄並另行分析，你的作答仍會被使用。")
ATTEST_ITEMS = [
    "我對本研究的了解全部來自本問卷，作答前與作答期間都未查閱與本研究相關的線上資料。",
    "我未曾閱讀本研究的論文或任何研究報告。",
    "我不知道、也未曾看過研究者對這些題目所做的歸類。",
    "我未曾看過任何 AI 工具對這些題目所做的分類結果。",
    "在交回問卷之前，我未與其他編碼者討論或交流這些題目。",
    "我作答時完全依靠自己的判斷，未使用 ChatGPT、Claude、Gemini、Copilot 或其他 AI 工具。",
    "我獨立完成了全部題目的作答。",
]
ATTEST_FREE = ["姓名", "日期", "簽名"]

INFO_HEAD = ["題序", "題目編號", "AI 功能類型", "AI 功能類型定義", "決策功能類型", "決策功能類型定義", "文獻樣態說明"]
CODER_HEAD = ["題序", "題目編號", "AI 功能類型", "決策功能類型", "AI 功能類型定義", "決策功能類型定義", "文獻樣態說明"]
RESP_KEYS = ["Q1", "Q2", "Q3", "Q4", "Q5"]
RESP_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
ITEM_ROW_HEIGHT = 200
HEAD_ROW_HEIGHT = 75
DV_ERR_TITLE, DV_ERR = "輸入值無效", "請從下拉選單中選擇。"


def sha(p):
    with open(p, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def sha_text(t):
    return hashlib.sha256(t.encode("utf-8")).hexdigest()


# ---------------------------------------------------------------- S1 source (read-only)
def verify_s1(repo):
    s1 = os.path.join(repo, S1_REL)
    man = os.path.join(s1, "V1_SEALED", "FREEZE_MANIFEST_V1_8_2_S1.txt")
    want = {}
    for line in open(man, encoding="utf-8"):
        if line.startswith("#") or not line.strip() or "SEALED_LOCAL:" in line:
            continue
        h, rel = line.split(None, 1)
        want[rel.strip()] = h
    bad = [r for r, h in want.items() if sha(os.path.join(s1, r)) != h]
    if bad or len(want) != 21:
        raise SystemExit(f"S1 hash verification failed: {bad} (n={len(want)})")
    return len(want)


def load_s1_items(repo):
    s1 = os.path.join(repo, S1_REL)
    wb = load_workbook(os.path.join(s1, "20_MASTER_PACKET_v1_8_2.xlsx"), read_only=True)
    rows = list(wb["ITEM_INFO"].iter_rows(values_only=True))
    h = list(rows[0])
    ix = {k: h.index(k) for k in ("Item_ID", "AI_family", "AI_family_definition", "Decision_family",
                                  "Decision_family_definition")}
    items = {}
    for r in rows[1:]:
        items[r[ix["Item_ID"]]] = dict(AI_family=r[ix["AI_family"]], AI_def=r[ix["AI_family_definition"]],
                                       Dec_family=r[ix["Decision_family"]], Dec_def=r[ix["Decision_family_definition"]])
    if sorted(items) != NEUTRAL:
        raise SystemExit("S1 master ITEM_INFO item set mismatch")
    s1_orders = {}
    for c in CODERS:
        wbc = load_workbook(os.path.join(s1, "coder_distribution", f"{c}.xlsx"), read_only=True)
        cr = list(wbc[c].iter_rows(values_only=True))
        hh = list(cr[0])
        s1_orders[c] = [r[hh.index("Item_ID")] for r in cr[1:]]
        for r in cr[1:]:
            it = items[r[hh.index("Item_ID")]]
            if (r[hh.index("AI_family_definition")] != it["AI_def"]
                    or r[hh.index("Decision_family_definition")] != it["Dec_def"]
                    or not str(r[hh.index("AI_family")]).startswith(it["AI_family"] + " (")
                    or not str(r[hh.index("Decision_family")]).startswith(it["Dec_family"] + " (")):
                raise SystemExit(f"S1 {c} content differs from master for {r[hh.index('Item_ID')]}")
    return items, s1_orders


def s1_en_text(it):
    return "\n".join([it["AI_family"], it["AI_def"], it["Dec_family"], it["Dec_def"]])


def to_zh(it):
    a, d = AI_FAM[it["AI_family"]], DEC_FAM[it["Dec_family"]]
    if a[2] != it["AI_def"] or d[2] != it["Dec_def"]:
        raise SystemExit(f"English definition mismatch for {it['AI_family']} / {it['Dec_family']}")
    return dict(AI_code=a[0], AI_zh=a[1], AI_def_zh=a[3], Dec_code=d[0], Dec_zh=d[1], Dec_def_zh=d[3])


def zh_text(z):
    return "\n".join([z["AI_zh"], z["AI_def_zh"], z["Dec_zh"], z["Dec_def_zh"]])


def make_orders(s1_orders):
    orders = {}
    for c in CODERS:
        o = list(NEUTRAL)
        random.Random(SEEDS[c]).shuffle(o)
        orders[c] = o
    tup = [tuple(o) for o in orders.values()]
    if len(set(tup)) != 3 or tuple(NEUTRAL) in tup or set(tup) & {tuple(o) for o in s1_orders.values()}:
        raise SystemExit("S1b orders not distinct from each other / identity / S1 orders")
    return orders


# ---------------------------------------------------------------- workbook parts
def _w(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def neutral_props(wb, title):
    p = wb.properties
    p.creator = "研究協調人"
    p.lastModifiedBy = "研究協調人"
    p.title = title
    for k in ("subject", "keywords", "description", "category"):
        setattr(p, k, "")


def add_readme(wb):
    ws = wb.create_sheet(SHEET_README)
    for i, t in enumerate(README_ROWS, 1):
        ws.cell(i, 1, t if t else None).alignment = WRAP
    ws.cell(1, 1).font = HEAD_FONT
    _w(ws, {"A": 140})


def add_categories(wb):
    ws = wb.create_sheet(SHEET_CATS)
    for j, h in enumerate(["類別或規則", "定義或內容", "區分說明與使用規則"], 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i, row in enumerate(CATEGORY_ROWS, 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v if v else None).alignment = WRAP
    _w(ws, {"A": 26, "B": 80, "C": 80})


def add_item_info(wb, zh, order):
    ws = wb.create_sheet(SHEET_INFO)
    for j, h in enumerate(INFO_HEAD, 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i, nid in enumerate(order, 2):
        z = zh[nid]
        for j, v in enumerate([i - 1, nid, z["AI_zh"], z["AI_def_zh"], z["Dec_zh"], z["Dec_def_zh"], PATTERN_TEXT], 1):
            ws.cell(i, j, v).alignment = WRAP
        ws.row_dimensions[i].height = ITEM_ROW_HEIGHT * 0.6
    _w(ws, {"A": 6, "B": 11, "C": 18, "D": 50, "E": 18, "F": 50, "G": 50})


def add_coder_sheet(wb, coder, zh, order):
    ws = wb.create_sheet(SHEET_CODER[coder])
    head = CODER_HEAD + [Q_TEXT[k] for k in RESP_KEYS]
    for j, h in enumerate(head, 1):
        c = ws.cell(1, j, h)
        c.font = HEAD_FONT
        c.alignment = WRAP
    ws.row_dimensions[1].height = HEAD_ROW_HEIGHT
    for i, nid in enumerate(order, 2):
        z = zh[nid]
        vals = [i - 1, nid, z["AI_zh"], z["Dec_zh"], z["AI_def_zh"], z["Dec_def_zh"], PATTERN_TEXT]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v).alignment = WRAP
        for j in range(len(vals) + 1, len(head) + 1):
            c = ws.cell(i, j)                       # response cell: EMPTY
            c.fill = RESP_FILL
            c.alignment = WRAP
        ws.row_dimensions[i].height = ITEM_ROW_HEIGHT
    n, first = len(order) + 1, len(CODER_HEAD) + 1

    def col(idx):
        return ws.cell(1, idx).column_letter

    for k, idx in (("Q1", first), ("Q2", first + 1), ("Q3", first + 2), ("Q4", first + 3)):
        dv = DataValidation(type="list", formula1='"' + ",".join(LM.options(k)) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle=DV_ERR_TITLE, error=DV_ERR)
        dv.add(f"{col(idx)}2:{col(idx)}{n}")
        ws.add_data_validation(dv)
    dv5 = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="600", allow_blank=True,
                         showErrorMessage=True, errorTitle=DV_ERR_TITLE, error="請以 1–3 句話說明，最多 600 字。")
    dv5.add(f"{col(first + 4)}2:{col(first + 4)}{n}")
    ws.add_data_validation(dv5)
    ws.freeze_panes = "C2"
    _w(ws, {"A": 6, "B": 11, "C": 16, "D": 16, "E": 40, "F": 40, "G": 40,
            col(first): 16, col(first + 1): 24, col(first + 2): 14, col(first + 3): 24, col(first + 4): 60})


def add_attestation(wb, coders):
    ws = wb.create_sheet(SHEET_ATTEST)
    ws.cell(1, 1, ATTEST_PREAMBLE).font = HEAD_FONT
    ws.cell(1, 1).alignment = WRAP
    ws.row_dimensions[1].height = 48
    for j, h in enumerate(["編碼者", "聲明", "回答（符合／不符合）"], 1):
        ws.cell(2, j, h).font = HEAD_FONT
    r, dv_cells = 3, []
    for c in coders:
        for st in ATTEST_ITEMS:
            ws.cell(r, 1, CODER_ZH[c])
            ws.cell(r, 2, st).alignment = WRAP
            ws.cell(r, 3).fill = RESP_FILL
            dv_cells.append(f"C{r}")
            r += 1
        for lab in ATTEST_FREE:
            ws.cell(r, 1, CODER_ZH[c])
            ws.cell(r, 2, lab)
            ws.cell(r, 3).fill = RESP_FILL
            r += 1
    dv = DataValidation(type="list", formula1='"' + ",".join(LM.options("ATTESTATION")) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle=DV_ERR_TITLE, error=DV_ERR)
    for a in dv_cells:
        dv.add(a)
    ws.add_data_validation(dv)
    _w(ws, {"A": 12, "B": 90, "C": 24})


def add_adjudication(wb):                                          # study-team master only
    ws = wb.create_sheet("ADJUDICATION_EMPTY")
    head = ["Item_ID", "Coder_Q1_values (paste after S2 freeze)", "Coder_Q2_values (paste after S2 freeze)",
            "Disagreement_type", "Route (CODER_CONSENSUS / INDEPENDENT_ADJUDICATOR / AUTHOR_ADJUDICATED)",
            "Final_Q1", "Final_Q2", "Final_Q4", "Adjudicator_role", "Evidence_used", "Rationale"]
    for j, h in enumerate(head, 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i, nid in enumerate(NEUTRAL, 2):
        ws.cell(i, 1, nid)
    _w(ws, {c: 22 for c in "ABCDEFGHIJK"})


def add_code_map(wb):                                              # study-team master only
    ws = wb.create_sheet("STUDY_TEAM_CODE_MAP")
    for j, h in enumerate(["Field", "Display_label_ZH_TW", "Canonical_code", "English_label_S1"], 1):
        ws.cell(1, j, h).font = HEAD_FONT
    r = 2
    for f in LM.FIELDS:
        for lab, code in zip(LM.options(f), LM.codes(f)):
            for j, v in enumerate([f, lab, code, ""], 1):
                ws.cell(r, j, v)
            r += 1
    for kind, fam in (("AI_FAMILY", AI_FAM), ("DECISION_FAMILY", DEC_FAM)):
        for en, (code, zh, _e, _z) in fam.items():
            for j, v in enumerate([kind, zh, code, en], 1):
                ws.cell(r, j, v)
            r += 1
    _w(ws, {"A": 18, "B": 30, "C": 30, "D": 34})


# ---------------------------------------------------------------- build
def build(repo, team):
    n_s1 = verify_s1(repo)
    items, s1_orders = load_s1_items(repo)
    zh = {i: to_zh(it) for i, it in items.items()}
    orders = make_orders(s1_orders)
    out = os.path.join(repo, OUT_REL)
    os.makedirs(os.path.join(out, "coder_distribution"), exist_ok=True)
    os.makedirs(os.path.join(out, "V1_SEALED"), exist_ok=True)
    os.makedirs(team, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    neutral_props(wb, "編碼問卷（研究團隊總表）")
    add_readme(wb)
    add_categories(wb)
    add_item_info(wb, zh, NEUTRAL)
    for c in CODERS:
        add_coder_sheet(wb, c, zh, orders[c])
    add_attestation(wb, CODERS)
    add_adjudication(wb)
    add_code_map(wb)
    paths = [os.path.join(out, "20_MASTER_PACKET_S1B_ZH_TW.xlsx")]
    wb.save(paths[0])
    for c in CODERS:
        wb = Workbook()
        wb.remove(wb.active)
        neutral_props(wb, "編碼問卷")
        add_readme(wb)
        add_categories(wb)
        add_item_info(wb, zh, orders[c])
        add_coder_sheet(wb, c, zh, orders[c])
        add_attestation(wb, [c])
        p = os.path.join(out, "coder_distribution", f"{c}_ZH_TW.xlsx")
        wb.save(p)
        paths.append(p)
    # content correspondence (S1 English per ID <-> S1b Chinese per ID); no G ID
    with open(os.path.join(out, "S1B_ITEM_CONTENT_CORRESPONDENCE.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["Item_ID", "S1_AI_family_EN", "S1_Decision_family_EN", "S1_EN_text_sha256", "S1B_AI_family_ZH",
                    "S1B_Decision_family_ZH", "S1B_ZH_text_sha256", "AI_family_code", "Decision_family_code"])
        for i in NEUTRAL:
            it, z = items[i], zh[i]
            w.writerow([i, it["AI_family"], it["Dec_family"], sha_text(s1_en_text(it)), z["AI_zh"], z["Dec_zh"],
                        sha_text(zh_text(z)), z["AI_code"], z["Dec_code"]])
    # order log: study team only; commitment in OUT
    log_path = os.path.join(team, "S1B_ITEM_ORDER_LOG.json")
    with open(log_path, "w", encoding="utf-8", newline="\n") as f:
        json.dump(dict(label="STUDY_TEAM_ONLY - S1b per-coder item orders - do not commit (public repository)",
                       seeds=SEEDS, rule="random.Random(seed_string).shuffle(ITEM-A01..A16)",
                       coder_orders=orders, coder_files={c: f"{c}_ZH_TW.xlsx" for c in CODERS},
                       checks=dict(orders_pairwise_distinct=True, differ_from_identity=True,
                                   differ_from_every_S1_order=True)), f, indent=1, ensure_ascii=False)
    with open(os.path.join(out, "V1_SEALED", "S1B_ITEM_ORDER_LOG.COMMITMENT.md"), "w", encoding="utf-8",
              newline="\n") as f:
        f.write("# S1B_ITEM_ORDER_LOG — COMMITMENT\n\n"
                "| File | SHA-256 | Bytes |\n|---|---|---|\n"
                f"| S1B_ITEM_ORDER_LOG.json | `{sha(log_path)}` | {os.path.getsize(log_path)} |\n\n"
                "Plaintext is held in study-team-only storage (`outputs\\20_SEALED_PRIVATE\\V1_8_2B_S1B_STUDY_TEAM_ONLY\\`), "
                "not in git (public repository). It contains the three S1b per-coder item orders and their seeds; it "
                "contains no item-to-G assignment. The seeds are non-secret (orders are over neutral IDs only and "
                "are visible in the coder files); they are also listed in `S1B_PACKET_MANIFEST.csv`.\n\n"
                "Checks recorded at build time: the three orders are pairwise distinct, differ from the identity "
                "order, and differ from every S1 per-coder order.\n\n"
                "The sealed ITEM-ID to G-ID map is unchanged from S1 and was not opened: see "
                "`claude/r14_merge_v1_8_2_v1_prefreeze/V1_SEALED/ITEM_ID_TO_GID_MAP.COMMITMENT.md`, "
                "SHA-256 of the map `2a84d5de58ec3d13878281d684d82b959d8bc73be7068845bd4cd2f2448cbe81`.\n")
    print(f"S1 verified {n_s1}/21; built {len(paths)} workbooks")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True)
    ap.add_argument("--study-team", required=True)
    a = ap.parse_args()
    build(a.repo, a.study_team)
