"""Leakage check for the V1.8.2 masked packet (pre-human freeze). Required result: 0 findings.

Scans
  (1) every coder-facing workbook (coder_distribution/CODER_1..3.xlsx and the study-team master
      20_MASTER_PACKET_v1_8_2.xlsx) at two levels:
      - cell level (openpyxl): forbidden researcher / AI-pilot / network / sensitivity vocabulary, category codes in
        item rows, mechanism-hint words in neutral item text, AI-pilot text overlap, network numbers,
        non-empty response / attestation / adjudication cells, per-item row formatting uniformity
        (incl. the sealed sensitivity items, identified via the sealed item map without printing it);
      - raw XML level (unzip ALL parts, incl. docProps/core.xml, app.xml, custom XML, comments, defined names,
        data validations, formulas): G-ID tokens, researcher/pilot terms, hidden sheets/rows/cols, comments,
        custom XML, defined names, formulas.
  (2) every text file in OUT for ITEM-ID <-> G-ID pairings, per-coder order lists, the secret seed, and any
      plaintext map/log file.
Prints a summary; writes OUT/tools/leakage_check_result_v1_8_2.json (no mapping, no order is written).

Usage: python -E -P leakage_check_v182.py --repo REPO --sealed-local DIR
"""
from __future__ import annotations

import argparse
import csv
import glob
import hashlib
import json
import os
import re
import sys
import zipfile

from openpyxl import load_workbook

sys.dont_write_bytecode = True
OUT_REL = os.path.join("claude", "r14_merge_v1_8_2_v1_prefreeze")
CATS = ["M1", "M2", "M3", "M4", "M5", "LONGITUDINAL", "DATA_GATED", "EVIDENCE_ONLY", "NONE", "AMBIGUOUS",
        "OPTIONAL_DATA_GATED", "OTHER", "INSUFFICIENT_INFORMATION"]
CAT_RE = re.compile(r"\b(" + "|".join(CATS) + r")\b")
GID_RE = re.compile(r"(?<![A-Za-z0-9_])G(0[1-9]|1[0-6])(?![0-9])")
ITEM_RE = re.compile(r"ITEM-A(0[1-9]|1[0-6])")
FORBIDDEN = [r"\bI[1-5]\b", r"\banchor\b", r"\bMerge\b", r"\bKeep\b", r"static core", r"static bridge",
             r"\bbridge\b", r"researcher'?s? (key|assignment|mapping) (is|=)", r"\bsealed\b", r"\bkey\b", r"\b41_",
             r"\b53A\b", r"\b54[ABC]\b", r"\b132\b", r"convergence", r"\bpilot\b", r"AI coder", r"Optional module",
             r"contested", r"sensitivity", r"爭議", r"研究者", r"裁決", r"答案", r"O/E", r"q\s*=", r"p_BH",
             r"observed\s*=", r"\bdynamic\b", r"\bstatic\b", r"evidence-only", r"G_ID", r"Gap_ID", r"\bGap\b",
             r"Study 1", r"robust core", r"enrichment"]
# researcher-assignment phrasing; the attestation legitimately says "the researcher's assignments" generically
ATTEST_OK = re.compile(r"researcher's assignments of these items to categories")
HINT_EN = ["authority", "accountab", "approve", "sign-off", "sign off", "veto", "stop", "governance", "policy",
           "routing", "route", "recovery", "rollback", "exception", "workflow", "handoff", "configuration",
           "division of work", "override", "delegat", "assurance", "validation", "monitor", "audit", "drift",
           "longitudinal", "feedback", "learning", "next period", "data-gated", "threshold", "label", "score",
           "mechanism", "interface", "evidence layer", "frontier", "pilot"]
HINT_ZH = ["誰", "負責", "問責", "簽核", "停止", "回復", "例外", "流程", "門檻", "回饋", "下一期", "試點", "分工",
           "配置決策", "組織配置", "技術應用前沿", "嵌入", "審", "驗證", "最終權", "授權"]
ITEM_TEXT_COLS = {"Literature_pattern", "Evidence_description"}
RESP_PREFIX = ("Q1_", "Q2_", "Q3_", "Q4_", "Q5_")
NET_COLS = ("E_analytic", "Observed", "O_E_Ratio", "q_primary_BH", "q_exact", "q_document_preserving", "q_raw_record")
ATTR_STRIP = re.compile(r'\s(r|ref|sqref|topLeftCell|activeCell|spans|xr:uid|uid)="[^"]*"')


def sha(p):
    with open(p, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def shingles(text, n=6):
    w = re.findall(r"[a-z0-9']+", text.lower())
    return {" ".join(w[i:i + n]) for i in range(len(w) - n + 1)}


def xml_scan(path, F):
    name = os.path.basename(path)
    with zipfile.ZipFile(path) as z:
        parts = z.namelist()
        for p in parts:
            low = p.lower()
            if "comments" in low or low.startswith("customxml") or "threadedcomment" in low:
                F.append((name, "XML_FORBIDDEN_PART", p))
            raw = z.read(p)
            try:
                txt = raw.decode("utf-8")
            except UnicodeDecodeError:
                continue
            body = ATTR_STRIP.sub("", txt)
            for m in GID_RE.finditer(body):
                F.append((name, "XML_GID_TOKEN", f"{p}: ...{body[max(0, m.start() - 30):m.end() + 10]}..."))
            if re.search(r"<definedName[ >]", txt):
                F.append((name, "XML_DEFINED_NAME", p))
            if re.search(r"<f[ >]", txt):
                F.append((name, "XML_FORMULA", p))
            if re.search(r'hidden="1"|state="(hidden|veryHidden)"', txt):
                F.append((name, "XML_HIDDEN_ELEMENT", p))
            for pat in (r"researcher", r"\bpilot\b", r"sealed", r"sensitivity", r"G_ID", r"O/E", r"p_BH", r"爭議"):
                hits = [x for x in re.finditer(pat, body, flags=re.I)
                        if not ATTEST_OK.search(body[max(0, x.start() - 5):x.start() + 80])]
                if hits:
                    F.append((name, "XML_FORBIDDEN_TERM", f"{p}: /{pat}/"))
            if p == "docProps/core.xml":
                for tag in ("dc:title", "dc:subject", "cp:keywords", "dc:creator", "cp:lastModifiedBy",
                            "dc:description", "cp:category"):
                    m = re.search(rf"<{tag}>(.*?)</{tag}>", txt)
                    if m and (GID_RE.search(m.group(1)) or re.search(r"G01|mapping|key", m.group(1), re.I)):
                        F.append((name, "DOCPROPS_LEAK", f"{tag}={m.group(1)}"))
    return len(parts)


def cell_scan(path, ai_sh, net, fam_labels, sens_ids, F, is_coder):
    wb = load_workbook(path)
    name = os.path.basename(path)
    coder_sheets = [s for s in wb.sheetnames if s.startswith("CODER_")]
    if is_coder and len(coder_sheets) != 1:
        F.append((name, "PER_CODER_FILE_SHEETS", str(coder_sheets)))
    nonempty = dict(response=0, attestation=0, adjudication=0)
    sigs = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        hi = 1 if ws.title == "ATTESTATION" else 0
        head = [str(h) if h is not None else "" for h in rows[hi]] if len(rows) > hi else []
        for ri, row in enumerate(rows):
            for ci, v in enumerate(row):
                if v is None:
                    continue
                s, col = str(v), (head[ci] if ci < len(head) else "")
                where = f"{ws.title}!R{ri + 1}C{ci + 1}"
                if ri > hi and col.startswith(RESP_PREFIX):
                    nonempty["response"] += 1
                    F.append((name, "NON_EMPTY_RESPONSE_CELL", where))
                if ws.title == "ATTESTATION" and ri > hi and ci == 2:
                    nonempty["attestation"] += 1
                    F.append((name, "NON_EMPTY_ATTESTATION_CELL", where))
                if ws.title == "ADJUDICATION_EMPTY" and ri > 0 and ci > 0:
                    nonempty["adjudication"] += 1
                    F.append((name, "NON_EMPTY_ADJUDICATION_CELL", where))
                if GID_RE.search(s):
                    F.append((name, "GID_IN_CELL", where))
                if ws.title.startswith(("CODER_", "ITEM_INFO", "ADJUDICATION")) and ri > 0 and CAT_RE.search(s):
                    F.append((name, "CATEGORY_CODE_IN_ITEM_ROW", where))
                for pat in FORBIDDEN:
                    if re.search(pat, s, flags=re.I) and not ATTEST_OK.search(s):
                        F.append((name, "FORBIDDEN_TOKEN", f"{where}: /{pat}/"))
                if col in ITEM_TEXT_COLS and ri > 0:
                    st = s
                    for fam in fam_labels:
                        st = st.replace(fam, "<FAMILY>")
                    for h in HINT_EN:
                        if h in st.lower():
                            F.append((name, "MECHANISM_HINT_IN_ITEM_TEXT", f"{where}: {h}"))
                    for h in HINT_ZH:
                        if h in st:
                            F.append((name, "MECHANISM_HINT_IN_ITEM_TEXT", f"{where}: {h}"))
                if shingles(s) & ai_sh:
                    F.append((name, "AI_PILOT_TEXT_OVERLAP", where))
                for num in net:
                    if num in s:
                        F.append((name, "NETWORK_NUMBER", f"{where}: {num}"))
        if ws.title.startswith("CODER_"):
            for r in ws.iter_rows(min_row=2):
                if r[1].value is None:
                    continue
                sig = tuple((c.fill.fgColor.rgb, c.font.b, c.font.i, c.font.color.rgb if c.font.color else None,
                             c.alignment.wrap_text) for c in r)
                sigs[(ws.title, r[1].value)] = (sig, sum(c.value is not None for c in r),
                                                ws.row_dimensions[r[0].row].hidden,
                                                ws.row_dimensions[r[0].row].height)
    for sheet in {k[0] for k in sigs}:
        vals = {i: v for (sh, i), v in sigs.items() if sh == sheet}
        if len(set(vals.values())) != 1:
            F.append((name, "ITEM_ROWS_NOT_UNIFORM", sheet))
        for sid in sens_ids:
            if sid not in vals:
                F.append((name, "SENSITIVITY_ITEM_MISSING", sheet))
    return nonempty


def text_scan(out, sealed_files, seed, F):
    n = 0
    for p in glob.glob(os.path.join(out, "**", "*"), recursive=True):
        if os.path.isdir(p) or p.lower().endswith(".xlsx"):
            continue
        base = os.path.basename(p)
        if base in ("ITEM_ID_TO_GID_MAP.json", "ITEM_ORDER_LOG.json") or "SECRET_SEED" in base:
            F.append((base, "PLAINTEXT_SEALED_FILE_IN_OUT", p))
        try:
            txt = open(p, encoding="utf-8").read()
        except (UnicodeDecodeError, OSError):
            continue
        n += 1
        if str(seed) in txt:
            F.append((base, "SECRET_SEED_IN_OUT", ""))
        for i, line in enumerate(txt.splitlines(), 1):
            if ITEM_RE.search(line) and GID_RE.search(line):
                F.append((base, "ITEM_GID_PAIR_ON_LINE", f"line {i}"))
            if len(set(ITEM_RE.findall(line))) >= 5:
                F.append((base, "ITEM_ORDER_LIST_ON_LINE", f"line {i}"))
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True)
    ap.add_argument("--sealed-local", required=True)
    a = ap.parse_args()
    out = os.path.join(a.repo, OUT_REL)
    imap = json.load(open(os.path.join(a.sealed_local, "ITEM_ID_TO_GID_MAP.json"), encoding="utf-8"))["map"]
    seed = json.load(open(os.path.join(a.sealed_local, "ITEM_ORDER_LOG.json"), encoding="utf-8"))["secret_seed"]
    sens_ids = [i for i, gid in imap.items() if gid in ("G04", "G10", "G11")]
    ai_sh, pre = set(), set()
    for f in ("ai_coder_A_output.csv", "ai_coder_B_output.csv"):
        for r in csv.DictReader(open(os.path.join(a.repo, "v1_blind", f), encoding="utf-8")):
            ai_sh |= shingles(r["rationale"]) | shingles(r["enterprise_decision_interface"])
    for p in ("53B_V1_BLIND_CODING_PROTOCOL_v2_2.md", os.path.join("v1_human_packet", "CODER_CLARIFICATIONS_v2_2_1.md")):
        pre |= shingles(open(os.path.join(a.repo, p), encoding="utf-8").read())
    ai_sh -= pre
    net = set()
    for r in csv.DictReader(open(os.path.join(a.repo, "11_Priority_Gaps_Frozen_Export_G01_G16_v1_0_20261009.csv"),
                                 encoding="utf-8-sig")):
        for c in NET_COLS:
            v = r[c]
            if re.fullmatch(r"[0-9.]+", v) and len(v) >= 4:
                net.add(v)
                net.add(f"{float(v):.2f}")
    net = {x for x in net if len(x) >= 4 and x != "0.00"}
    fam = []
    for r in csv.DictReader(open(os.path.join(a.repo, "54A_V1_BLIND_PACKET_A_v2_2.csv"), encoding="utf-8")):
        for x in (r["ai_family"], r["decision_family"]):
            if x not in fam:
                fam.append(x)
    F = []
    dist = sorted(glob.glob(os.path.join(out, "coder_distribution", "*")))
    dist_names = [os.path.basename(p) for p in dist]
    if dist_names != ["CODER_1.xlsx", "CODER_2.xlsx", "CODER_3.xlsx"]:
        F.append(("coder_distribution", "UNEXPECTED_FILES", str(dist_names)))
    files = dist + [os.path.join(out, "20_MASTER_PACKET_v1_8_2.xlsx")]
    per = {}
    for p in files:
        is_coder = "coder_distribution" in p
        parts = xml_scan(p, F)
        ne = cell_scan(p, ai_sh, net, fam, sens_ids, F, is_coder)
        per[os.path.relpath(p, out).replace("\\", "/")] = dict(sha256=sha(p), xml_parts_scanned=parts,
                                                              distributed_to_coders=is_coder, nonempty=ne)
    n_text = text_scan(out, None, seed, F)
    res = dict(checker="leakage_check_v182.py", files=per, text_files_scanned=n_text,
               ai_pilot_shingles_checked=len(ai_sh), network_numbers_checked=len(net),
               sensitivity_rows_checked_for_uniformity=len(sens_ids),
               coder_distribution_files=dist_names,
               findings=[dict(file=f, check=c, where=w) for f, c, w in F])
    res["n_findings"] = len(res["findings"])
    res["verdict"] = "PASS" if not F else "FAIL"
    os.makedirs(os.path.join(out, "tools"), exist_ok=True)
    with open(os.path.join(out, "tools", "leakage_check_result_v1_8_2.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump(res, f, indent=1, ensure_ascii=False)
    print(res["verdict"], res["n_findings"])
    for x in res["findings"][:40]:
        print(x)
    return 0 if not F else 1


if __name__ == "__main__":
    sys.exit(main())
