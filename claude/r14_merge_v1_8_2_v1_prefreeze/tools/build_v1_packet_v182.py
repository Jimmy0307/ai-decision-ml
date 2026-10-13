"""Build the V1.8.2 human construct-validation packet with MASKED item IDs (pre-human freeze).

Item IDs ITEM-A01..ITEM-A16 are assigned by a seeded random permutation whose seed is SECRET: it is generated once
with `secrets` and stored only in SEALED_LOCAL (study-team-only, not in git, repository is public). The
plaintext ITEM_ID_TO_GID_MAP.json and ITEM_ORDER_LOG.json are written ONLY to SEALED_LOCAL. Nothing written to
OUT pairs an item ID with a G ID or lists a per-coder order in readable form (only the workbooks themselves, which
show neutral IDs).

Never writes responses; never simulates coders; never reads the researcher key or AI-pilot outputs.

Usage: python -E -P build_v1_packet_v182.py --repo REPO --sealed-local DIR
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import random
import secrets
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.dont_write_bytecode = True
OUT_REL = os.path.join("claude", "r14_merge_v1_8_2_v1_prefreeze")
CODERS = ["CODER_1", "CODER_2", "CODER_3"]
GIDS = [f"G{i:02d}" for i in range(1, 17)]
NEUTRAL = [f"ITEM-A{i:02d}" for i in range(1, 17)]

Q1_OPTS = ["YES", "NO", "UNCERTAIN"]
Q2_OPTS = ["M1", "M2", "M3", "M4", "LONGITUDINAL", "DATA_GATED", "EVIDENCE_ONLY", "NONE", "AMBIGUOUS"]
Q3_OPTS = ["HIGH", "MEDIUM", "LOW"]
Q4_OPTS = ["M1", "M2", "M3", "M4", "LONGITUDINAL", "DATA_GATED"]

AI_LABEL = {
    "Human Augmentation": "HUMAN_AUGMENTATION", "Prediction and Risk Estimation": "PREDICTION_RISK_ESTIMATION",
    "Coordination and Execution": "COORDINATION_EXECUTION", "Generation and Synthesis": "GENERATION_SYNTHESIS",
    "Explanation and Assurance": "EXPLANATION_ASSURANCE", "Evaluation and Audit": "EVALUATION_ASSESSMENT",
    "Recommendation and Optimization": "RECOMMENDATION_OPTIMIZATION",
}
DEC_LABEL = {
    "Prediction and Diagnosis": "PREDICTION_DIAGNOSIS", "Human Review and Accountability": "HUMAN_REVIEW_ACCOUNTABILITY",
    "Risk Assessment": "RISK_ASSESSMENT", "Allocation and Planning": "ALLOCATION_PLANNING",
    "Evaluation and Selection": "EVALUATION_SELECTION", "Governance and Oversight": "GOVERNANCE_OVERSIGHT",
    "Information Acquisition": "INFORMATION_ACQUISITION", "General Decision Process": "GENERAL_DECISION_PROCESS",
}
ZH = {
    "HUMAN_AUGMENTATION": "人機增能", "PREDICTION_RISK_ESTIMATION": "預測與風險估計",
    "COORDINATION_EXECUTION": "協調與執行", "GENERATION_SYNTHESIS": "生成與合成",
    "EXPLANATION_ASSURANCE": "解釋與保證", "EVALUATION_ASSESSMENT": "評估與審核",
    "RECOMMENDATION_OPTIMIZATION": "推薦與最佳化",
    "PREDICTION_DIAGNOSIS": "預測與診斷", "HUMAN_REVIEW_ACCOUNTABILITY": "人工複核與問責",
    "RISK_ASSESSMENT": "風險評估", "ALLOCATION_PLANNING": "配置與規劃", "EVALUATION_SELECTION": "評估與選擇",
    "GOVERNANCE_OVERSIGHT": "治理與監督", "INFORMATION_ACQUISITION": "資訊取得",
    "GENERAL_DECISION_PROCESS": "一般決策流程",
}
PATTERN_TEXT = ("Literature pattern: DEPLETED. In this study's literature network, the number of published "
                "documents that directly link AI functions of type '{ai}' to decision functions of type '{dec}' "
                "is below the number expected under a statistical null model, and this holds under all four "
                "pre-specified analysis methods. This is a research lead only. It does not mean that the "
                "combination is impossible, ineffective, unimportant, or absent from practice.")

README_ROWS = [
    "V1 HUMAN CONSTRUCT-VALIDATION CODING — INSTRUCTIONS (identical for every coder)",
    "",
    "Work alone. Expected time: 45-60 minutes. Do not discuss the items with anyone until you have submitted.",
    "Do not consult the project's code repository, any thesis, manuscript or report of this project, any other "
    "coder, or any AI assistant (ChatGPT, Claude, Gemini, Copilot or other) while coding. If you have already seen "
    "any of these, tell the coordinator before you start.",
    "",
    "There are 16 items, identified only by a neutral item ID. Each item shows one AI function family and one "
    "decision function family, with their definitions, and the same standard description of a literature pattern. "
    "Every item is answered with the same five questions.",
    "",
    "Q1  Does this pairing correspond to an identifiable enterprise decision interface?  YES / NO / UNCERTAIN",
    "     An enterprise decision interface is a decision that an enterprise must make about how this kind of AI "
    "function feeds this kind of decision function, and that is specific to this pairing.",
    "Q2  Best category (answer for every item; see sheet CATEGORY_DEFINITIONS and the consistency rules there).",
    "Q3  Your confidence in Q2:  HIGH / MEDIUM / LOW.",
    "Q4  Optional secondary category. Leave blank unless a second category is clearly necessary.",
    "Q5  Short rationale (1-3 sentences). Required for every item. If Q2 is NONE, state the enterprise decision that "
    "no listed category captures. If Q2 is DATA_GATED, name the data. If Q2 is AMBIGUOUS, name the competing "
    "categories.",
    "",
    "Use the drop-down lists in the coloured response columns. Do not change the item rows, IDs or order. "
    "Fill sheet ATTESTATION, save the file, and return it directly to the study coordinator (not via any shared repository).",
]

CATEGORY_ROWS = [
    ("Q1", "YES", "You can name a decision that an enterprise must make about how this kind of AI function feeds this "
     "kind of decision function, and that decision is specific to this pairing.", ""),
    ("Q1", "NO", "The only decisions you can name would apply equally to almost any AI function feeding this decision "
     "function, or no enterprise decision is involved.", "If Q1 = NO, set Q2 = EVIDENCE_ONLY."),
    ("Q1", "UNCERTAIN", "You cannot decide between YES and NO.", "Q2 may be any category except NONE; AMBIGUOUS is allowed."),
    ("Q2/Q4", "M1  Authority & Accountability (授權與問責; = Governance & Decision Rights / authority-to-act)",
     "Use when the main issue of the connection is who may approve, reject, stop or change what the AI output "
     "is allowed to decide; formal versus effective control; accountability; enterprise policy boundaries.",
     "Do not use when the main issue is how outputs are routed or technically validated."),
    ("Q2/Q4", "M2  Workflow Routing & Recovery",
     "Use when the main issue of the connection is how the AI output enters the next decision or action step; "
     "handoffs across systems or units; exception handling; rollback and recovery.",
     "Do not use when the main issue is who holds authority."),
    ("Q2/Q4", "M3  Human-AI Configuration & Response",
     "Use when the main issue of the connection is the sequence and division of work between people and AI "
     "(who proposes, who decides, aggregation, delegation, override) and how users actually respond (use, verify, "
     "modify, escalate, reject, bypass).",
     "Do not use when the main issue is formal authority or validation evidence."),
    ("Q2/Q4", "M4  Evaluation & Assurance",
     "Use when the main issue of the connection is pre-deployment validation, monitoring, revalidation, "
     "auditability, explanation usability, drift, or evidence that the AI output is fit for the decision.",
     "Do not use when the main issue is who decides or how work is sequenced."),
    ("Q2/Q4", "LONGITUDINAL  Longitudinal extension",
     "Use when the enterprise decision can only be studied with outcomes observed over more than one period: results "
     "of one cycle change later resource allocation, policy or information acquisition. A single-period choice is "
     "not enough to represent it.",
     "Do not use when a single-period configuration choice is enough. If you use it, you may put the M1-M4 "
     "category involved in Q4."),
    ("Q2/Q4", "DATA_GATED  Data-gated module",
     "Use when an enterprise decision exists only if a specific kind of data is available. Name the data in Q5.",
     "Do not use when a category applies without special data. If you use it, you may put the M1-M4 category "
     "involved in Q4."),
    ("Q2", "EVIDENCE_ONLY  Evidence-layer only",
     "Use when the pairing describes a research or technology pattern but implies no distinct enterprise decision "
     "beyond what any AI function feeding this decision function would raise.",
     "Use together with Q1 = NO (or UNCERTAIN). Do not use when a specific category clearly applies."),
    ("Q2", "NONE  None of the listed categories",
     "Use when an identifiable enterprise decision interface exists but none of M1, M2, M3, M4, LONGITUDINAL or "
     "DATA_GATED adequately expresses it. Requires Q1 = YES, and Q5 must state the enterprise decision that is not "
     "captured.",
     "Do not use for insufficient evidence or when two categories tie (use AMBIGUOUS, or Q1 = UNCERTAIN), or when "
     "there is no distinct interface (use EVIDENCE_ONLY)."),
    ("Q2", "AMBIGUOUS",
     "Use when two or more categories fit equally well and you cannot choose one, or the definitions do not allow a "
     "defensible single choice. Name the competing categories in Q5.",
     "Do not use as a default; choose a category whenever one fits better."),
    ("Rule", "Consistency", "Q1 = YES -> Q2 is one of M1, M2, M3, M4, LONGITUDINAL, DATA_GATED, NONE, AMBIGUOUS.  "
     "Q1 = NO -> Q2 = EVIDENCE_ONLY.  Q1 = UNCERTAIN -> any Q2 except NONE.  Q4 must differ from Q2.", ""),
    ("Rule", "Order of tests for Q2",
     "(a) Can the decision be represented only with outcomes over more than one period? -> LONGITUDINAL. "
     "(b) Does the decision exist only if a named dataset is available? -> DATA_GATED. "
     "(c) Is there a distinct single-period decision? -> the M1-M4 category whose 'use when' text fits the main issue "
     "(or NONE if none of them expresses it). (d) Otherwise -> EVIDENCE_ONLY with Q1 = NO. "
     "Use AMBIGUOUS only if you cannot choose.", ""),
    ("Rule", "Residual decision family",
     "'General Decision Process' is a residual category. Code it like any other family if a category fits; otherwise "
     "use EVIDENCE_ONLY or AMBIGUOUS.", ""),
    ("Rule", "Code the enterprise decision",
     "Code the enterprise decision, not the technology, not the topic of the literature, and not shared words "
     "between a family name and a category name.", ""),
    ("Rule", "Q3 confidence", "HIGH / MEDIUM / LOW refers to Q2 only. It is descriptive and never weights agreement.", ""),
]

ATTEST_ITEMS = [
    "I have not read the public code repository (e.g. on GitHub) for this project.",
    "I have not read the thesis or any manuscript or report of this project.",
    "I have not seen the researcher's assignments of these items to categories.",
    "I have not seen any AI-generated coding output for these items.",
    "I did not consult another coder.",
    "I did not use ChatGPT, Claude, Gemini, Copilot or any other AI tool to answer.",
    "I coded every item independently.",
]

ITEM_HEAD = ["Item_ID", "AI_family", "AI_family_ZH", "AI_family_definition", "Decision_family", "Decision_family_ZH",
             "Decision_family_definition", "Literature_pattern", "Evidence_description"]
RESP_HEAD = ["Q1_identifiable_interface", "Q2_best_category", "Q3_confidence", "Q4_secondary_category_optional",
             "Q5_rationale"]
RESP_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def load_items(repo):
    gaps = {}
    with open(os.path.join(repo, "11_Priority_Gaps_Frozen_Export_G01_G16_v1_0_20261009.csv"), encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            gaps[(r["AI_Family"], r["Decision_Family"])] = r["Gap_ID"]
    items = {}
    with open(os.path.join(repo, "54A_V1_BLIND_PACKET_A_v2_2.csv"), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            pair = (AI_LABEL[r["ai_family"]], DEC_LABEL[r["decision_family"]])
            gid = gaps[pair]
            if gid in items:
                raise SystemExit(f"duplicate mapping for {gid}")
            items[gid] = dict(
                AI_family=r["ai_family"], AI_family_ZH=ZH[pair[0]], AI_family_definition=r["ai_family_definition"],
                Decision_family=r["decision_family"], Decision_family_ZH=ZH[pair[1]],
                Decision_family_definition=r["decision_family_definition"],
                Literature_pattern=PATTERN_TEXT.format(ai=r["ai_family"], dec=r["decision_family"]),
                Evidence_description=r["neutral_pair_description"] + " " + r["network_evidence"])
    if sorted(items) != GIDS:
        raise SystemExit("item set mismatch")
    return items


def secret_seed(sealed):
    p = os.path.join(sealed, "SECRET_SEED_v1_8_2.txt")
    if not os.path.exists(p):
        with open(p, "w", encoding="utf-8") as f:
            f.write(str(secrets.randbits(62)))
    with open(p, encoding="utf-8") as f:
        return int(f.read().strip())


def assign(seed):
    perm = list(GIDS)
    random.Random(f"{seed}-ITEM_ID").shuffle(perm)
    id_map = dict(zip(NEUTRAL, perm))                      # neutral -> G
    if perm == GIDS:
        raise SystemExit("degenerate permutation (identity); regenerate seed")
    orders = {}
    for c in CODERS:
        o = list(NEUTRAL)
        random.Random(f"{seed}-{c}-ORDER").shuffle(o)
        orders[c] = o
    if len({tuple(o) for o in orders.values()}) != len(CODERS):
        raise SystemExit("coder orders not distinct")
    return id_map, orders


def _w(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def neutral_props(wb):
    p = wb.properties
    p.creator = "study coordinator"
    p.lastModifiedBy = "study coordinator"
    p.title = "V1 coding packet"
    p.subject = ""
    p.keywords = ""
    p.description = ""
    p.category = ""


def add_readme(wb):
    ws = wb.create_sheet("README_INSTRUCTIONS")
    for i, t in enumerate(README_ROWS, 1):
        ws.cell(i, 1, t).alignment = WRAP
    ws.cell(1, 1).font = HEAD_FONT
    _w(ws, {"A": 150})


def add_categories(wb):
    ws = wb.create_sheet("CATEGORY_DEFINITIONS")
    for j, h in enumerate(["Question", "Category", "Use when (operational definition)", "Do not use when / rule"], 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i, row in enumerate(CATEGORY_ROWS, 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v).alignment = WRAP
    _w(ws, {"A": 9, "B": 40, "C": 90, "D": 60})


def add_item_info(wb, items, id_map, order):
    ws = wb.create_sheet("ITEM_INFO")
    for j, h in enumerate(["Item_order"] + ITEM_HEAD, 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i, nid in enumerate(order, 2):
        it = items[id_map[nid]]
        ws.cell(i, 1, i - 1)
        ws.cell(i, 2, nid)
        for j, h in enumerate(ITEM_HEAD[1:], 3):
            ws.cell(i, j, it[h]).alignment = WRAP
    _w(ws, {"A": 6, "B": 11, "C": 26, "D": 14, "E": 50, "F": 28, "G": 14, "H": 50, "I": 60, "J": 60})


def add_coder_sheet(wb, name, items, id_map, order):
    ws = wb.create_sheet(name)
    head = ["Item_order", "Item_ID", "AI_family", "Decision_family", "AI_family_definition",
            "Decision_family_definition", "Literature_pattern"] + RESP_HEAD
    for j, h in enumerate(head, 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i, nid in enumerate(order, 2):
        it = items[id_map[nid]]
        vals = [i - 1, nid, f"{it['AI_family']} ({it['AI_family_ZH']})",
                f"{it['Decision_family']} ({it['Decision_family_ZH']})",
                it["AI_family_definition"], it["Decision_family_definition"], it["Literature_pattern"]]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v).alignment = WRAP
        for j in range(len(vals) + 1, len(head) + 1):
            c = ws.cell(i, j)                       # response cell: EMPTY
            c.fill = RESP_FILL
            c.alignment = WRAP
    n, first = len(order) + 1, len(head) - len(RESP_HEAD) + 1

    def col(idx):
        return ws.cell(1, idx).column_letter

    for opts, idx in ((Q1_OPTS, first), (Q2_OPTS, first + 1), (Q3_OPTS, first + 2), (Q4_OPTS, first + 3)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid value",
                            error="Choose a value from the drop-down list.")
        dv.add(f"{col(idx)}2:{col(idx)}{n}")
        ws.add_data_validation(dv)
    dv5 = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="600", allow_blank=True,
                         showErrorMessage=True, error="Keep the rationale to 1-3 sentences (max 600 characters).")
    dv5.add(f"{col(first + 4)}2:{col(first + 4)}{n}")
    ws.add_data_validation(dv5)
    ws.freeze_panes = "C2"
    _w(ws, {"A": 6, "B": 11, "C": 30, "D": 30, "E": 45, "F": 45, "G": 55,
            col(first): 14, col(first + 1): 18, col(first + 2): 12, col(first + 3): 18, col(first + 4): 60})


def add_attestation(wb, sheet_names):
    ws = wb.create_sheet("ATTESTATION")
    ws.cell(1, 1, "Attestation. Answer every statement YES or NO. Answer NO honestly if it is not true; an honest NO "
                  "is recorded and does not invalidate your work silently.").font = HEAD_FONT
    for j, h in enumerate(["Coder_sheet", "Statement", "Answer (YES/NO)"], 1):
        ws.cell(2, j, h).font = HEAD_FONT
    r = 3
    for s in sheet_names:
        for st in ATTEST_ITEMS:
            ws.cell(r, 1, s)
            ws.cell(r, 2, st).alignment = WRAP
            ws.cell(r, 3).fill = RESP_FILL
            r += 1
        for lab in ("Name", "Date", "Signature"):
            ws.cell(r, 1, s)
            ws.cell(r, 2, lab)
            ws.cell(r, 3).fill = RESP_FILL
            r += 1
    dv = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
    for i in range(3, r):
        if ws.cell(i, 2).value not in ("Name", "Date", "Signature"):
            dv.add(f"C{i}")
    ws.add_data_validation(dv)
    _w(ws, {"A": 12, "B": 90, "C": 22})


def add_adjudication(wb):
    ws = wb.create_sheet("ADJUDICATION_EMPTY")
    head = ["Item_ID", "Coder_Q1_values (paste after S2 freeze)", "Coder_Q2_values (paste after S2 freeze)",
            "Disagreement_type", "Route (CODER_CONSENSUS / INDEPENDENT_ADJUDICATOR / AUTHOR_ADJUDICATED)",
            "Final_Q1", "Final_Q2", "Final_Q4", "Adjudicator_role", "Evidence_used", "Rationale"]
    for j, h in enumerate(head, 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i, nid in enumerate(NEUTRAL, 2):
        ws.cell(i, 1, nid)
    _w(ws, {c: 22 for c in "ABCDEFGHIJK"})


def build(repo, sealed):
    os.makedirs(sealed, exist_ok=True)
    items = load_items(repo)
    seed = secret_seed(sealed)
    id_map, orders = assign(seed)
    out = os.path.join(repo, OUT_REL)
    os.makedirs(os.path.join(out, "coder_distribution"), exist_ok=True)
    # study-team master (no G IDs): neutral-ID order, all three coder sheets
    wb = Workbook()
    wb.remove(wb.active)
    neutral_props(wb)
    add_readme(wb)
    add_categories(wb)
    add_item_info(wb, items, id_map, NEUTRAL)
    for c in CODERS:
        add_coder_sheet(wb, c, items, id_map, orders[c])
    add_attestation(wb, CODERS)
    add_adjudication(wb)
    paths = [os.path.join(out, "20_MASTER_PACKET_v1_8_2.xlsx")]
    wb.save(paths[0])
    for c in CODERS:
        wb = Workbook()
        wb.remove(wb.active)
        neutral_props(wb)
        add_readme(wb)
        add_categories(wb)
        add_item_info(wb, items, id_map, orders[c])
        add_coder_sheet(wb, c, items, id_map, orders[c])
        add_attestation(wb, [c])
        p = os.path.join(out, "coder_distribution", f"{c}.xlsx")
        wb.save(p)
        paths.append(p)
    # SEALED_LOCAL only
    with open(os.path.join(sealed, "ITEM_ID_TO_GID_MAP.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump(dict(label="STUDY_TEAM_ONLY - do not commit (public repository)", map=id_map), f, indent=1)
    with open(os.path.join(sealed, "ITEM_ORDER_LOG.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump(dict(label="STUDY_TEAM_ONLY - do not commit (public repository)", secret_seed=seed,
                       item_id_rule="random.Random(f'{seed}-ITEM_ID').shuffle(input_g_order); neutral IDs assigned in that order",
                       coder_order_rule="random.Random(f'{seed}-CODER_k-ORDER').shuffle(ITEM-A01..A16)",
                       input_g_order=GIDS, item_id_to_gid=id_map, coder_orders=orders), f, indent=1)
    print("built", len(paths), "workbooks")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True)
    ap.add_argument("--sealed-local", required=True)
    a = ap.parse_args()
    build(a.repo, a.sealed_local)
