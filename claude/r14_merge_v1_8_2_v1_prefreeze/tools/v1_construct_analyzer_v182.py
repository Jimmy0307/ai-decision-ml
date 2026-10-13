"""V1 human construct-validation analyzer, v1.8.2 (masked item IDs, 3 human coders; analysis plan 29).

Primary question: do independent human coders agree (inter-rater reliability)? NOT agreement with the researcher key.

Phases (each gated by a hash-verified freeze manifest; see 25 / 29):
  (blank)  all response cells empty                          -> NO_HUMAN_DATA, nothing computed
  R        --s2-manifest: the 3 returned sheets hash-frozen   -> reliability on NEUTRAL item IDs only. No item map,
                                                                 no researcher key. Sections:
           A. CURRENT_INSTRUMENT_RELIABILITY (primary: Krippendorff alpha nominal on Q2, all 16 items, 3 coders;
              secondary: pairwise Cohen kappa; Q1 alpha + kappa, no pass/fail; Q3 descriptive; Q4 Jaccard)
           B. 53B_CONTINUITY_RELIABILITY (derived primary_53B / status_53B, 53B C1/C3 classes) — never collapsed with A
  S3-gated --s3-manifest (adjudication file hash-frozen) unlocks:
           S  --item-map (hash verified vs ITEM_ID_TO_GID_MAP.COMMITMENT.md): G04/G10/G11 sensitivity (16 vs 13 items)
           K  --unseal --key (hash verified vs RESEARCHER_ASSIGNMENT_COMMITMENT.md) + --item-map: key comparison
           C  --classes (hash in --c-manifest): 53B s11 classes -> V1_PASS_HUMAN_C6 continuity label
Exposure: --exposed CODER_k (repeatable). EXPOSED coders are never silently pooled: results are reported for
CLEAN coders (primary, if >= 2 clean) and for ALL coders (labelled INCLUDES_EXPOSED).

Statistics reuse UNCHANGED functions from REPO/v1_blind/reliability_v2_2.py (hash-checked at runtime).
Thresholds .800/.667 = RESEARCH_CONVENTION (53B C1/C2); no Landis & Koch labels.

Usage:
  python -E -P v1_construct_analyzer_v182.py CODER_1.xlsx CODER_2.xlsx CODER_3.xlsx --out DIR [--repo REPO]
         [--s2-manifest M2] [--exposed CODER_k ...] [--adjudication ADJ.csv --s3-manifest M3]
         [--item-map ITEM_ID_TO_GID_MAP.json] [--unseal --key sealed_key.json] [--classes C.csv --c-manifest MC]
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import os
import random
import re
import sys
from collections import Counter
from itertools import combinations

from openpyxl import load_workbook

sys.dont_write_bytecode = True
VERSION = "v1_construct_analyzer_v1_8_2"
PRIMARY_RELIABILITY = "KRIPPENDORFF_ALPHA_Q2_3_CODERS"
PAIRWISE_KAPPA = "SECONDARY"
THRESHOLD_PROVENANCE = "RESEARCH_CONVENTION"
N_HUMAN_CODERS = 3
ITEMS = [f"ITEM-A{i:02d}" for i in range(1, 17)]
SENS_G = ("G04", "G10", "G11")
Q1 = {"YES", "NO", "UNCERTAIN"}
Q2 = {"M1", "M2", "M3", "M4", "LONGITUDINAL", "DATA_GATED", "EVIDENCE_ONLY", "NONE", "AMBIGUOUS"}
Q3 = {"HIGH", "MEDIUM", "LOW"}
Q4 = {"M1", "M2", "M3", "M4", "LONGITUDINAL", "DATA_GATED"}
RESP = ["Q1_identifiable_interface", "Q2_best_category", "Q3_confidence", "Q4_secondary_category_optional",
        "Q5_rationale"]
REL_SHA256 = "866e9005da8a20e996c93b5c8f6bee916992f24d60c6a96e8f1dcda4556848e7"
SEED, B = 20261011, 5000
LO, HI = 0.667, 0.800
C6_SENTENCE = ("V1_PASS_HUMAN_C6 is a continuity label under the earlier 53B protocol, not evidence that the amended "
               "Q1–Q5 instrument is identical to 53B.")
OUT_REL = os.path.join("claude", "r14_merge_v1_8_2_v1_prefreeze")
KEY_COMMIT = os.path.join(OUT_REL, "V1_SEALED", "RESEARCHER_ASSIGNMENT_COMMITMENT.md")
MAP_COMMIT = os.path.join(OUT_REL, "V1_SEALED", "ITEM_ID_TO_GID_MAP.COMMITMENT.md")


class InputInvalid(ValueError):
    pass


class Refused(RuntimeError):
    pass


def sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def default_repo():
    here = os.path.dirname(os.path.abspath(__file__))
    for cand in (os.path.abspath(os.path.join(here, "..", "..", "..")),
                 r"C:\NTUST_PAPER\outputs\20261014_V1_8_1_STUDY1_FREEZE\repo"):
        if os.path.exists(os.path.join(cand, "v1_blind", "reliability_v2_2.py")):
            return cand
    raise SystemExit("REPO not found; pass --repo")


def load_rel(repo):
    p = os.path.join(repo, "v1_blind", "reliability_v2_2.py")
    if sha(p) != REL_SHA256:
        raise SystemExit(f"reliability_v2_2.py hash mismatch: {p}")
    spec = importlib.util.spec_from_file_location("reliability_v2_2", p)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def manifest_hashes(path):
    if not path or not os.path.exists(path):
        return set()
    with open(path, encoding="utf-8") as f:
        return {m.group(1) for line in f for m in [re.match(r"^([0-9a-f]{64})\s", line)] if m}


def committed_hash(repo, rel, label):
    txt = open(os.path.join(repo, rel), encoding="utf-8").read()
    m = re.search(re.escape(label) + r"[^`\n]*`([0-9a-f]{64})`", txt)
    if not m:
        raise Refused(f"no committed hash for {label} in {rel}")
    return m.group(1)


# ------------------------------------------------------------------------------------------------ input
def _cell(v):
    return "" if v is None else str(v).strip()


def read_sheet(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    names = [s for s in wb.sheetnames if s.startswith("CODER_")]
    if len(names) != 1:
        raise InputInvalid(f"{os.path.basename(path)}: expected exactly one CODER_* sheet, found {names}")
    rows = list(wb[names[0]].iter_rows(values_only=True))
    head = [_cell(h) for h in rows[0]]
    need = ["Item_ID"] + RESP
    if any(h not in head for h in need):
        raise InputInvalid(f"{os.path.basename(path)}: missing columns")
    ix = {h: head.index(h) for h in need}
    data = {}
    for r in rows[1:]:
        if r is None or all(v is None for v in r):
            continue
        i = _cell(r[ix["Item_ID"]])
        if i in data:
            raise InputInvalid(f"{os.path.basename(path)}: duplicate {i}")
        data[i] = {h: _cell(r[ix[h]]) for h in RESP}
    att = []
    if "ATTESTATION" in wb.sheetnames:
        for r in list(wb["ATTESTATION"].iter_rows(values_only=True))[2:]:
            if r and r[0]:
                att.append((_cell(r[1]), _cell(r[2])))
    return dict(path=path, sheet=names[0], sha256=sha(path), rows=data, attestation=att)


def is_empty(s):
    return all(not any(v for v in row.values()) for row in s["rows"].values()) and not any(a for _, a in s["attestation"])


def validate(s):
    errs, name = [], s["sheet"]
    if sorted(s["rows"]) != ITEMS:
        errs.append(f"{name}: item set is not ITEM-A01..A16")
    for i, r in s["rows"].items():
        q1, q2, q3, q4, q5 = (r[h] for h in RESP)
        if q1 not in Q1:
            errs.append(f"{name} {i}: Q1 '{q1}' invalid")
        if q2 not in Q2:
            errs.append(f"{name} {i}: Q2 '{q2}' invalid")
        if q3 not in Q3:
            errs.append(f"{name} {i}: Q3 '{q3}' invalid")
        if q4 and (q4 not in Q4 or q4 == q2):
            errs.append(f"{name} {i}: Q4 '{q4}' invalid or equal to Q2")
        if not q5:
            errs.append(f"{name} {i}: Q5 rationale empty")
        if q1 == "YES" and q2 == "EVIDENCE_ONLY":
            errs.append(f"{name} {i}: Q1=YES with Q2=EVIDENCE_ONLY")
        if q1 == "NO" and q2 != "EVIDENCE_ONLY":
            errs.append(f"{name} {i}: Q1=NO requires Q2=EVIDENCE_ONLY")
        if q2 == "NONE" and q1 != "YES":
            errs.append(f"{name} {i}: NONE requires Q1=YES (OA-3)")
        if q2 == "NONE" and not q5:
            errs.append(f"{name} {i}: NONE requires Q5 stating the uncaptured enterprise decision (OA-3)")
    if len(s["attestation"]) and any(a == "" for _, a in s["attestation"]):
        errs.append(f"{name}: attestation incomplete")
    return errs


def attestation_exposed(s):
    return [st for st, a in s["attestation"] if a == "NO" and st not in ("Name", "Date", "Signature")]


# ------------------------------------------------------------------------------------------------ metrics
def derived(q2, q4):
    """Codebook 26 s5.2 (OA-2: secondary continuity only)."""
    if q2 in ("M1", "M2", "M3", "M4"):
        return q2, "static"
    return {"LONGITUDINAL": ("M5", "dynamic"),
            "DATA_GATED": (q4 if q4 in ("M1", "M2", "M3", "M4") else "OPTIONAL_DATA_GATED", "optional"),
            "EVIDENCE_ONLY": ("EVIDENCE_ONLY", "evidence-only"),
            "NONE": ("OTHER", "static"),
            "AMBIGUOUS": ("INSUFFICIENT_INFORMATION", "insufficient")}[q2]


def classify_primary(alpha, lo):
    if alpha != alpha:
        return dict(label="UNRESOLVED", sub="UNDEFINED", imprecise=True)
    sub = "RELIABLE" if alpha >= HI else "TENTATIVE" if alpha >= LO else "UNRESOLVED/NOT_ACCEPTABLE"
    imprecise = lo != lo or lo < LO
    return dict(label=sub, sub=sub + ("+IMPRECISE" if imprecise else ""), imprecise=imprecise,
                provenance=THRESHOLD_PROVENANCE)


def kappa_ci(rel, a, b):
    rng = random.Random(SEED)
    n, vals = len(a), []
    for _ in range(B):
        idx = [rng.randrange(n) for _ in range(n)]
        k, _, _ = rel.cohen_kappa([a[i] for i in idx], [b[i] for i in idx])
        if k == k:
            vals.append(k)
    vals.sort()
    if not vals:
        return [float("nan"), float("nan")], 0
    return [vals[int(0.025 * len(vals))], vals[int(0.975 * len(vals)) - 1]], len(vals)


def confusion(a, b):
    labels = sorted(set(a) | set(b))
    m = {x: {y: 0 for y in labels} for x in labels}
    for x, y in zip(a, b):
        m[x][y] += 1
    return dict(labels=labels, matrix=m)


def field(rel, coders, items, get, mode):
    data = [[get(c, i) for c in coders] for i in items]
    alpha = rel.kripp_alpha_nominal(data)
    lo, hi, nb = rel.bootstrap_alpha(data)
    out = dict(n_items=len(items), n_coders=len(coders), alpha_nominal=alpha, alpha_ci95=[lo, hi],
               alpha_boot_valid=nb, all_coders_agree_items=sum(len(set(u)) == 1 for u in data), pairs=[])
    if mode == "primary":
        out["classification"] = classify_primary(alpha, lo)
    elif mode == "53B":
        out["c1_c3_class_53B"] = rel.classify(alpha, lo)
    else:
        out["classification"] = "NOT_CLASSIFIED (no pass/fail; NO_SOURCE_THRESHOLD)"
    for ci, cj in combinations(coders, 2):
        a, b = [get(ci, i) for i in items], [get(cj, i) for i in items]
        k, po, _ = rel.cohen_kappa(a, b)
        kci, kb = kappa_ci(rel, a, b)
        out["pairs"].append(dict(pair=[ci["sheet"], cj["sheet"]], agreements=sum(x == y for x, y in zip(a, b)),
                                 raw_agreement=po, cohen_kappa=k, kappa_ci95=kci, kappa_boot_valid=kb,
                                 role="SECONDARY", confusion=confusion(a, b)))
    out["mean_pairwise_raw_agreement"] = sum(p["raw_agreement"] for p in out["pairs"]) / len(out["pairs"])
    return out


def g(fieldname):
    return lambda c, i: c["rows"][i][fieldname]


def d_primary(c, i):
    return derived(c["rows"][i][RESP[1]], c["rows"][i][RESP[3]])[0]


def d_status(c, i):
    return derived(c["rows"][i][RESP[1]], c["rows"][i][RESP[3]])[1]


def descriptives(coders, items):
    per = {}
    for i in items:
        q2s = [c["rows"][i][RESP[1]] for c in coders]
        pairs = list(combinations(range(len(coders)), 2))
        per[i] = dict(Q2={c["sheet"]: c["rows"][i][RESP[1]] for c in coders},
                      Q1={c["sheet"]: c["rows"][i][RESP[0]] for c in coders},
                      Q2_modal=Counter(q2s).most_common(1)[0][0],
                      Q2_pair_agreement_share=sum(q2s[a] == q2s[b] for a, b in pairs) / len(pairs),
                      Q2_all_agree=len(set(q2s)) == 1,
                      status="AGREED" if len(set(q2s)) == 1 else "TO_ADJUDICATION")
    freq = {c["sheet"]: dict(Counter(c["rows"][i][RESP[1]] for i in items)) for c in coders}
    freq["POOLED"] = dict(Counter(c["rows"][i][RESP[1]] for c in coders for i in items))

    def rates(cs):
        n = len(cs) * len(items)
        return {k: sum(c["rows"][i][f] == v for c in cs for i in items) / n
                for k, f, v in (("Q1_UNCERTAIN", RESP[0], "UNCERTAIN"), ("Q2_AMBIGUOUS", RESP[1], "AMBIGUOUS"),
                                ("Q2_NONE", RESP[1], "NONE"))}
    jac = []
    for i in items:
        sets = [{c["rows"][i][RESP[1]]} | ({c["rows"][i][RESP[3]]} - {""}) for c in coders]
        jac += [len(a & b) / len(a | b) for a, b in combinations(sets, 2)]
    order = ["LOW", "MEDIUM", "HIGH"]
    conf = {}
    for c in coders:
        def med(its):
            v = sorted(order.index(c["rows"][i][RESP[2]]) for i in its)
            return None if not v else order[v[(len(v) - 1) // 2]]
        agreed = [i for i in items if per[i]["Q2_all_agree"]]
        conf[c["sheet"]] = dict(distribution=dict(Counter(c["rows"][i][RESP[2]] for i in items)),
                                median_lower=med(items), median_lower_agreed=med(agreed),
                                median_lower_disagreed=med([i for i in items if i not in agreed]))
    return dict(per_item=per, Q2_category_frequency=freq,
                uncertainty_ambiguity_rates={**{c["sheet"]: rates([c]) for c in coders}, "POOLED": rates(coders)},
                none_rows=[dict(coder=c["sheet"], item=i, Q5=c["rows"][i][RESP[4]]) for c in coders for i in items
                           if c["rows"][i][RESP[1]] == "NONE"],
                Q4_jaccard_Q2_union_Q4_mean=sum(jac) / len(jac),
                Q3_confidence_descriptive_only=conf)


def reliability_block(rel, coders, items):
    return dict(
        A_CURRENT_INSTRUMENT_RELIABILITY=dict(
            primary=PRIMARY_RELIABILITY,
            Q2_primary=field(rel, coders, items, g(RESP[1]), "primary"),
            Q1_secondary=field(rel, coders, items, g(RESP[0]), None),
            descriptives=descriptives(coders, items)),
        B_53B_CONTINUITY_RELIABILITY=dict(
            note="OA-2 APPROVED_WITH_SCOPE_LIMIT: secondary continuity metrics only. " + C6_SENTENCE,
            derived_primary_53B_C1=field(rel, coders, items, d_primary, "53B"),
            derived_status_53B_C3=field(rel, coders, items, d_status, "53B")))


# ------------------------------------------------------------------------------------------------ S3-gated
def load_adjudication(path, per_item):
    final = {i: dict(Q2=v["Q2_modal"], Q4="", route="AGREED") for i, v in per_item.items() if v["Q2_all_agree"]}
    with open(path, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            i = r["Item_ID"].strip()
            if i not in ITEMS or r["Final_Q2"].strip() not in Q2:
                raise InputInvalid(f"adjudication row invalid: {i}")
            final[i] = dict(Q2=r["Final_Q2"].strip(), Q4=r.get("Final_Q4", "").strip(), route=r["Route"].strip())
    missing = [i for i in ITEMS if i not in final]
    if missing:
        raise InputInvalid(f"adjudication missing items {missing}")
    return final


def load_item_map(repo, path):
    want = committed_hash(repo, MAP_COMMIT, "ITEM_ID_TO_GID_MAP.json")
    if sha(path) != want:
        raise Refused(f"item map sha256 {sha(path)} != committed {want}")
    m = json.load(open(path, encoding="utf-8"))["map"]
    if sorted(m) != ITEMS:
        raise Refused("item map does not cover ITEM-A01..A16")
    return m


def flatten_key(entry):
    primary, _s, status = entry
    return {"dynamic": "LONGITUDINAL", "optional": "DATA_GATED", "evidence-only": "EVIDENCE_ONLY"}.get(status, primary)


def unseal(repo, key_path, item_map, final, manifest_path, adj_sha):
    want = committed_hash(repo, KEY_COMMIT, "sealed_key.json")
    if sha(key_path) != want:
        raise Refused(f"STOP: sealed_key.json sha256 {sha(key_path)} != committed {want}")
    key = json.load(open(key_path, encoding="utf-8"))["key"]
    rows = []
    for i in ITEMS:
        gid = item_map[i]
        fk = flatten_key(key[gid])
        fp, fs = derived(final[i]["Q2"], final[i]["Q4"])
        rows.append(dict(item=i, gap=gid, final_Q2=final[i]["Q2"], route=final[i]["route"], key_flat=fk,
                         flat_match=final[i]["Q2"] == fk, primary_53B_match=fp == key[gid][0],
                         status_53B_match=fs == key[gid][2], sensitivity_target=gid in SENS_G,
                         needs_53B_s11_class=final[i]["Q2"] != fk))
    with open(manifest_path, "a", encoding="utf-8", newline="\n") as f:
        f.write(f"# KEY_OPENED adjudication={adj_sha} by {VERSION}\n")
    return dict(key_sha256=want, rows=rows, flat_agreement=sum(r["flat_match"] for r in rows),
                note="Agreement with the researcher key is NOT the primary V1 question; the key is a hypothesis.")


def c6_continuity(cont, classes, n_clean):
    c1 = cont["derived_primary_53B_C1"]["c1_c3_class_53B"]
    c3 = cont["derived_status_53B_C3"]["c1_c3_class_53B"]
    structural = [r for r in classes if r["model_change_class"] == "model-structural"]
    ok = (n_clean >= 2 and not c1.startswith(("NOT_ACCEPTABLE", "UNDEFINED")) and "IMPRECISE" not in c1
          and not c3.startswith(("NOT_ACCEPTABLE", "UNDEFINED")) and not structural)
    return dict(label="V1_PASS_HUMAN_C6" if ok else "V1_C6_NOT_MET", C1=c1, C3=c3, structural=structural,
                statement=C6_SENTENCE)


# ------------------------------------------------------------------------------------------------ output
def fmt(x):
    return "n/a" if x is None or (isinstance(x, float) and x != x) else f"{x:.3f}"


def render_md(res):
    L = [f"# V1 construct-validation analysis ({VERSION})", "", f"**Status: `{res['status']}`**", ""]
    if res["status"] == "NO_HUMAN_DATA":
        L += ["All response and attestation cells in every supplied coder sheet are empty. Nothing was computed.", "",
              "| Sheet | SHA-256 |", "|---|---|"] + [f"| {s['sheet']} | `{s['sha256']}` |" for s in res["inputs"]]
        return "\n".join(L) + "\n"
    if res["status"] in ("INPUT_INVALID", "REFUSED"):
        return "\n".join(L + [f"- {e}" for e in res["errors"]]) + "\n"
    for scope, blk in res["reliability"].items():
        a = blk["A_CURRENT_INSTRUMENT_RELIABILITY"]["Q2_primary"]
        L += [f"## {scope}", "", "### A. CURRENT_INSTRUMENT_RELIABILITY (primary)", "",
              f"Q2 Krippendorff alpha (nominal) = {fmt(a['alpha_nominal'])} [{fmt(a['alpha_ci95'][0])}, "
              f"{fmt(a['alpha_ci95'][1])}]; class {a['classification']['sub']} ({THRESHOLD_PROVENANCE}); "
              f"all-coder agreement {a['all_coders_agree_items']}/{a['n_items']}", "",
              "| pair (secondary) | agree | kappa [95% CI] |", "|---|---|---|"]
        L += [f"| {p['pair'][0]}–{p['pair'][1]} | {p['agreements']}/{a['n_items']} | {fmt(p['cohen_kappa'])} "
              f"[{fmt(p['kappa_ci95'][0])}, {fmt(p['kappa_ci95'][1])}] |" for p in a["pairs"]]
        b = blk["B_53B_CONTINUITY_RELIABILITY"]
        L += ["", "### B. 53B_CONTINUITY_RELIABILITY (secondary; never collapsed with A)", "",
              f"C1 derived primary_53B: alpha {fmt(b['derived_primary_53B_C1']['alpha_nominal'])} "
              f"({b['derived_primary_53B_C1']['c1_c3_class_53B']}); C3 derived status_53B: alpha "
              f"{fmt(b['derived_status_53B_C3']['alpha_nominal'])} ({b['derived_status_53B_C3']['c1_c3_class_53B']})",
              "", C6_SENTENCE, ""]
    L.append("Full detail (Q1, Q3, Q4, per-item, confusion matrices, frequencies, rates) in the JSON.")
    return "\n".join(L) + "\n"


def run(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("sheets", nargs="+")
    ap.add_argument("--out", required=True)
    ap.add_argument("--repo")
    ap.add_argument("--s2-manifest")
    ap.add_argument("--exposed", action="append", default=[])
    ap.add_argument("--adjudication")
    ap.add_argument("--s3-manifest")
    ap.add_argument("--item-map")
    ap.add_argument("--unseal", action="store_true")
    ap.add_argument("--key")
    ap.add_argument("--classes")
    ap.add_argument("--c-manifest")
    a = ap.parse_args(argv)
    repo = a.repo or default_repo()
    os.makedirs(a.out, exist_ok=True)
    res = dict(analyzer=VERSION, primary_reliability=PRIMARY_RELIABILITY, pairwise_kappa=PAIRWISE_KAPPA,
               threshold_provenance=THRESHOLD_PROVENANCE, n_human_coders_design=N_HUMAN_CODERS, seed=SEED,
               bootstrap_reps=B, thresholds=dict(reliable=HI, tentative=LO, imprecise_ci_lower=LO))
    code = 0
    try:
        sheets = [read_sheet(p) for p in a.sheets]
        res["inputs"] = [dict(sheet=s["sheet"], file=os.path.basename(s["path"]), sha256=s["sha256"]) for s in sheets]
        if all(is_empty(s) for s in sheets):
            res["status"] = "NO_HUMAN_DATA"
        else:
            errs = []
            if len(sheets) != N_HUMAN_CODERS:
                errs.append(f"design requires {N_HUMAN_CODERS} human coder sheets, got {len(sheets)}")
            if len({s["sheet"] for s in sheets}) != len(sheets):
                errs.append("duplicate coder sheet names")
            for s in sheets:
                errs += validate(s)
            for s1, s2 in combinations(sheets, 2):
                if s1["rows"] == s2["rows"]:
                    errs.append(f"identical content: {s1['sheet']} and {s2['sheet']}")
            m2 = manifest_hashes(a.s2_manifest)
            errs += [f"{s['sheet']}: not hash-frozen in S2 manifest" for s in sheets if s["sha256"] not in m2]
            if errs:
                raise InputInvalid("\n".join(errs))
            rel = load_rel(repo)
            for s in sheets:
                s["status"] = "EXPOSED" if (s["sheet"] in a.exposed or attestation_exposed(s)) else "CLEAN"
            res["coder_status"] = {s["sheet"]: dict(status=s["status"], attestation_no=attestation_exposed(s))
                                   for s in sheets}
            clean = [s for s in sheets if s["status"] == "CLEAN"]
            res["reliability"] = {}
            if len(clean) == len(sheets):
                res["reliability"]["ALL_CODERS_CLEAN"] = reliability_block(rel, sheets, ITEMS)
                res["primary_scope"] = "ALL_CODERS_CLEAN"
            else:
                res["reliability"]["ALL_CODERS_INCLUDES_EXPOSED"] = reliability_block(rel, sheets, ITEMS)
                if len(clean) >= 2:
                    res["reliability"]["CLEAN_CODERS_ONLY"] = reliability_block(rel, clean, ITEMS)
                res["primary_scope"] = ("CLEAN_CODERS_ONLY (fewer than 3: replace exposed coder per plan 29)"
                                        if len(clean) >= 2 else "NOT_COMPUTABLE: fewer than 2 clean coders")
            res["status"] = "RELIABILITY_COMPUTED"
            scope = "ALL_CODERS_CLEAN" if "ALL_CODERS_CLEAN" in res["reliability"] else "CLEAN_CODERS_ONLY"
            base = clean if scope == "CLEAN_CODERS_ONLY" else sheets
            # ---------------- S3-gated stages
            if a.adjudication or a.item_map or a.unseal or a.classes:
                if not (a.adjudication and a.s3_manifest):
                    raise Refused("item map / key / classes require --adjudication and --s3-manifest (phase S3)")
                adj_sha = sha(a.adjudication)
                if adj_sha not in manifest_hashes(a.s3_manifest):
                    raise Refused("adjudication file is not hash-frozen in the S3 manifest; refused")
                per_item = res["reliability"][scope]["A_CURRENT_INSTRUMENT_RELIABILITY"]["descriptives"]["per_item"] \
                    if scope in res["reliability"] else None
                if per_item is None:
                    raise Refused("no clean scope for adjudication")
                final = load_adjudication(a.adjudication, per_item)
                res["adjudication_sha256"] = adj_sha
                res["status"] = "ADJUDICATION_IMPORTED"
                if a.item_map:
                    imap = load_item_map(repo, a.item_map)
                    sub = [i for i in ITEMS if imap[i] not in SENS_G]
                    res["sensitivity_excluding_G04_G10_G11"] = reliability_block(rel, base, sub)
                    res["sensitivity_items"] = {i: imap[i] for i in ITEMS if imap[i] in SENS_G}
                    res["status"] = "SENSITIVITY_COMPUTED"
                    if a.unseal:
                        if not a.key:
                            raise Refused("--unseal requires --key")
                        res["key_comparison"] = unseal(repo, a.key, imap, final, a.s3_manifest, adj_sha)
                        res["status"] = "KEY_COMPARED"
                        if a.classes:
                            if sha(a.classes) not in manifest_hashes(a.c_manifest):
                                raise Refused("classes file not hash-frozen in --c-manifest")
                            classes = list(csv.DictReader(open(a.classes, encoding="utf-8-sig")))
                            res["c6_continuity"] = c6_continuity(
                                res["reliability"][scope]["B_53B_CONTINUITY_RELIABILITY"], classes, len(clean))
                            res["status"] = "CLASSES_APPLIED"
                elif a.unseal:
                    raise Refused("--unseal requires --item-map (verified) to index the key")
    except InputInvalid as e:
        res["status"], res["errors"], code = "INPUT_INVALID", str(e).split("\n"), 2
    except Refused as e:
        res["status"], res["errors"], code = "REFUSED", [str(e)], 3
    with open(os.path.join(a.out, "V1_RESULTS_v1_8_2.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump(res, f, indent=1, ensure_ascii=False, default=str)
    with open(os.path.join(a.out, "V1_RESULTS_v1_8_2.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write(render_md(res))
    print(res["status"])
    return code


if __name__ == "__main__":
    sys.exit(run(sys.argv[1:]))
