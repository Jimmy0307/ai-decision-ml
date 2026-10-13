"""Task A (V1.8.1, agent H): item-by-item source closure of the 2026-09-29 flags in artifact 132.

Usage: python -I task_a_132.py <W> <REPO_V18_DIR> <OUT_DIR>
Reads historical 132 .csv/.xlsx from W (read-only), verifies each flagged G row against
V9 ROBUST_CORE_LOCAL_SENSITIVITY.csv rows and V1.8 04 T2, writes current copies + JSON log.
"""
import sys
sys.path.append(r'C:/Users/user/AppData/Roaming/Python/Python314/site-packages')
sys.stdout.reconfigure(encoding='utf-8')
import csv, hashlib, json, os, re
import openpyxl

W, V18, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
HIST = os.path.join(W, r'01_ALL_OUTPUTS\AI_DECISION_THESIS_ALL_OUTPUTS\r14_merge_v1_7')
V9 = os.path.join(W, r'02_V9_HUMAN_VALIDATION\20260929_V9_HUMAN_VALIDATION')
SRC_CSV = os.path.join(HIST, '132_GAP_TO_MECHANISM_TRACEABILITY_MATRIX.csv')
SRC_XLSX = os.path.join(HIST, '132_GAP_TO_MECHANISM_TRACEABILITY_MATRIX.xlsx')
SENS = os.path.join(V9, 'ROBUST_CORE_LOCAL_SENSITIVITY.csv')
R04 = os.path.join(V18, '04_LOCAL_SENSITIVITY_REPRODUCTION.md')
CA = os.path.join(OUT, 'current_artifacts')
DST_CSV = os.path.join(CA, '132_GAP_TO_MECHANISM_TRACEABILITY_MATRIX_v1_8_1.csv')
DST_XLSX = os.path.join(CA, '132_GAP_TO_MECHANISM_TRACEABILITY_MATRIX_v1_8_1.xlsx')

FLAG = '〔2026-09-29 紀錄，待核〕'
EXPECTED_SENS_SHA = 'e3166ee183e18f340c539b983108e2c972d18b916672a62510a25638db8b5b59'
PAT = re.compile(r'2026-09-29|待核|20260929')


def sha(p):
    h = hashlib.sha256()
    with open(p, 'rb') as f:
        h.update(f.read())
    return h.hexdigest()


def load_sens():
    with open(SENS, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    return {(r['ai_family'], r['decision_family']): r for r in rows}


def load_t2():
    t2 = {}
    txt = open(R04, encoding='utf-8').read()
    sec = txt.split('### T2.')[1].split('### T3.')[0]
    for line in sec.splitlines():
        m = re.match(r'\| (G\d\d) \| (\S+) x (\S+) \| (\d+) \| (\d+) \| (\w+) \| (\w+) \| (\d+) \|', line)
        if m:
            t2[m.group(1)] = dict(ai=m.group(2), dec=m.group(3), o132=int(m.group(4)),
                                  off=m.group(6), sc=m.group(7))
    return t2


def verify_row(gid, ai, dec, o132, sens, t2):
    r = sens.get((ai, dec))
    checks = {
        'sens_row_found': r is not None,
        'official_core_member=YES': r is not None and r['official_core_member'] == 'YES',
        'official_core_direction=DEPLETED': r is not None and r['official_core_direction'] == 'DEPLETED',
        'sidecar_core_member=YES': r is not None and r['sidecar_core_member'] == 'YES',
        'sidecar_core_direction=DEPLETED': r is not None and r['sidecar_core_direction'] == 'DEPLETED',
        'core_membership_changed=NO': r is not None and r['core_membership_changed'] == 'NO',
        'O132==official_observed': r is not None and int(r['official_observed']) == o132,
        't2_row_match': gid in t2 and t2[gid]['ai'] == ai and t2[gid]['dec'] == dec
                         and t2[gid]['off'] == 'DEPLETED' and t2[gid]['sc'] == 'DEPLETED',
    }
    return all(checks.values()), checks, r


def main():
    os.makedirs(CA, exist_ok=True)
    sens_sha = sha(SENS)
    assert sens_sha == EXPECTED_SENS_SHA, sens_sha
    r04_sha = sha(R04)
    sens, t2 = load_sens(), load_t2()
    assert len(t2) == 16, len(t2)

    raw = open(SRC_CSV, 'rb').read()
    text = raw.decode('utf-8-sig')
    lines = text.split('\n')
    rows = list(csv.reader(lines[:1] + [l for l in lines[1:] if l]))
    hdr = rows[0]
    col = hdr.index('Network status')
    log = []
    new_lines = list(lines)
    for i, row in enumerate(rows[1:], start=1):
        flagged = [j for j, c in enumerate(row) if PAT.search(c)]
        if not flagged:
            continue
        assert flagged == [col], (i, flagged)
        gid, ai, dec, o132 = row[0], row[1], row[3], int(row[5])
        ok, checks, r = verify_row(gid, ai, dec, o132, sens, t2)
        old = row[col]
        assert old.count(FLAG) == 1
        if ok:
            frag = ('〔VERIFIED_SOURCE: V9 ROBUST_CORE_LOCAL_SENSITIVITY.csv (sha256 %s…) row %s×%s: '
                    'official_core_member=YES/DEPLETED; sidecar_core_member=YES/DEPLETED; '
                    'core_membership_changed=NO; cross-check V1.8 04_LOCAL_SENSITIVITY_REPRODUCTION.md T2 '
                    '(sha256 %s…); local label-replacement sidecar only (not a full reanalysis)〕'
                    % (sens_sha[:12], ai, dec, r04_sha[:12]))
            status = 'VERIFIED'
        else:
            frag = FLAG
            status = 'KEPT_PENDING'
        new = old.replace(FLAG, frag)
        assert ',' not in frag and '"' not in frag
        # line-level replacement keeps all other bytes identical
        li = i  # data line index (no embedded newlines verified below)
        assert lines[li].count(FLAG) == 1
        new_lines[li] = lines[li].replace(FLAG, frag)
        log.append(dict(file='132 .csv', sheet='-', row=i + 1, column='%d (%s)' % (col + 1, hdr[col]),
                        gid=gid, cell=f'{ai}×{dec}', old=old, new=new, status=status, checks=checks,
                        sidecar_main_O=(r or {}).get('sidecar_observed'),
                        sidecar_doc_q=(r or {}).get('sidecar_document_q')))
    out_text = '\n'.join(new_lines)
    with open(DST_CSV, 'wb') as f:
        f.write(b'\xef\xbb\xbf' + out_text.encode('utf-8'))

    # ---- XLSX ----
    wb = openpyxl.load_workbook(SRC_XLSX)  # formulas kept (data_only=False)
    ws = wb['G_to_M_matrix']
    xhdr = [c.value for c in ws[1]]
    xcol = xhdr.index('Network status') + 1
    planned = {}
    for rr in range(2, ws.max_row + 1):
        c = ws.cell(rr, xcol)
        if not (isinstance(c.value, str) and PAT.search(c.value)):
            continue
        gid = ws.cell(rr, 1).value
        ai, dec, o = ws.cell(rr, 2).value, ws.cell(rr, 4).value, int(ws.cell(rr, 6).value)
        ok, checks, r = verify_row(gid, ai, dec, o, sens, t2)
        csv_entry = next(e for e in log if e['gid'] == gid)
        assert csv_entry['old'] == c.value, gid  # identical old wording in both formats
        new = csv_entry['new'] if ok else c.value
        planned[('G_to_M_matrix', c.coordinate)] = new
        log.append(dict(file='132 .xlsx', sheet='G_to_M_matrix', row=rr, column=c.coordinate,
                        gid=gid, cell=f'{ai}×{dec}', old=c.value, new=new,
                        status='VERIFIED' if ok else 'KEPT_PENDING', checks=checks))
        c.value = new
    rd = wb['README']
    a8 = rd['A8'].value
    assert PAT.search(a8)
    all_ok = all(e['status'] == 'VERIFIED' for e in log)
    if all_ok:
        new_a8 = ("VERIFIED_SOURCE: label-substitution note ('未移除〔VERIFIED_SOURCE: …〕'): "
                  "the V9 human-validation package (folder 20260929_V9_HUMAN_VALIDATION; zip sha256 8f63e88e…) "
                  "was recovered and partially verified in V1.8 (V9_SOURCE_VERIFIED=PARTIAL). "
                  "Each G row was checked against the manifest-pinned ROBUST_CORE_LOCAL_SENSITIVITY.csv "
                  "(sha256 %s…; sidecar produced 2026-10-03, not 2026-09-29) and V1.8 "
                  "04_LOCAL_SENSITIVITY_REPRODUCTION.md T2 (sha256 %s…): no G01–G16 cell removed. "
                  "Scope: local label-replacement sidecar (20 reconciled + 30 first-coder labels), not a full reanalysis; "
                  "the V9.1 evidence re-segmentation sensitivity is separate (G07/G15 boundary-sensitive). "
                  "Manuscript Appendix E39." % (sens_sha[:12], r04_sha[:12]))
        rd['A8'].value = new_a8
        planned[('README', 'A8')] = new_a8
        log.append(dict(file='132 .xlsx', sheet='README', row=8, column='A8', gid='-', cell='legend',
                        old=a8, new=new_a8, status='VERIFIED',
                        checks={'all_16_G_rows_verified': True}))
    wb.save(DST_XLSX)

    json.dump(dict(src_csv_sha=sha(SRC_CSV), src_xlsx_sha=sha(SRC_XLSX), sens_sha=sens_sha, r04_sha=r04_sha,
                   dst_csv_sha=sha(DST_CSV), dst_xlsx_sha=sha(DST_XLSX),
                   planned={f'{k[0]}!{k[1]}': v for k, v in planned.items()}, log=log),
              open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'task_a_out.json'), 'w',
                   encoding='utf-8'), ensure_ascii=False, indent=1)
    print('csv items', sum(1 for e in log if e['file'] == '132 .csv'),
          'xlsx items', sum(1 for e in log if e['file'] == '132 .xlsx'),
          'verified', sum(1 for e in log if e['status'] == 'VERIFIED'),
          'pending', sum(1 for e in log if e['status'] != 'VERIFIED'))


if __name__ == '__main__':
    main()
