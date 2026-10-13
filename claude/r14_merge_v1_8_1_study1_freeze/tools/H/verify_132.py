"""Verify 132 current copies against the historical originals (agent H, V1.8.1).

Usage: python -I verify_132.py <src_csv> <dst_csv> <src_xlsx> <dst_xlsx> <task_a_out.json>
Checks: sheet names, dimensions, every formula, every non-flag cell identical; changed cells == planned.
"""
import sys
sys.path.append(r'C:/Users/user/AppData/Roaming/Python/Python314/site-packages')
sys.stdout.reconfigure(encoding='utf-8')
import csv, io, json
import openpyxl

src_csv, dst_csv, src_x, dst_x, plan_json = sys.argv[1:6]
plan = json.load(open(plan_json, encoding='utf-8'))
planned = plan['planned']
res = {}

# CSV
def rows(p):
    b = open(p, 'rb').read()
    return b[:3] == b'\xef\xbb\xbf', b.count(b'\r'), b.endswith(b'\n'), list(csv.reader(io.StringIO(b.decode('utf-8-sig'), newline='')))
bs, crs, ends, a = rows(src_csv)
bd, crd, endd, b = rows(dst_csv)
res['csv_bom_same'] = bs == bd
res['csv_cr_count_same'] = crs == crd
res['csv_trailing_newline_same'] = ends == endd
res['csv_shape_same'] = len(a) == len(b) and all(len(x) == len(y) for x, y in zip(a, b))
changed = [(i + 1, j + 1) for i, (x, y) in enumerate(zip(a, b)) for j, (u, v) in enumerate(zip(x, y)) if u != v]
csv_log = {(e['row'], int(e['column'].split()[0])): e['new'] for e in plan['log'] if e['file'] == '132 .csv' and e['old'] != e['new']}
res['csv_changed_cells'] = len(changed)
res['csv_changed_eq_planned'] = set(changed) == set(csv_log) and all(b[r - 1][c - 1] == csv_log[(r, c)] for r, c in changed)
res['csv_header_same'] = a[0] == b[0]

# XLSX (formulas kept: data_only=False)
wa, wb = openpyxl.load_workbook(src_x), openpyxl.load_workbook(dst_x)
res['sheetnames_same'] = wa.sheetnames == wb.sheetnames
res['sheetnames'] = wa.sheetnames
dims, nform, ndiff, bad = {}, 0, [], []
for n in wa.sheetnames:
    sa, sb = wa[n], wb[n]
    dims[n] = (sa.dimensions, sb.dimensions)
    for r in range(1, max(sa.max_row, sb.max_row) + 1):
        for c in range(1, max(sa.max_column, sb.max_column) + 1):
            ca, cb = sa.cell(r, c), sb.cell(r, c)
            if isinstance(ca.value, str) and ca.value.startswith('='):
                nform += 1
                if ca.value != cb.value:
                    bad.append(('FORMULA', n, ca.coordinate))
            key = f'{n}!{ca.coordinate}'
            if ca.value != cb.value:
                ndiff.append(key)
                if key not in planned or cb.value != planned[key]:
                    bad.append(('UNPLANNED', key))
            elif ca.number_format != cb.number_format or ca.data_type != cb.data_type:
                bad.append(('FORMAT', key))
res['dimensions'] = dims
res['dimensions_same'] = all(x == y for x, y in dims.values())
res['formulas_checked'] = nform
res['xlsx_changed_cells'] = sorted(ndiff)
res['xlsx_changed_eq_planned'] = set(ndiff) == set(planned)
res['problems'] = bad
res['PASS'] = all([res['csv_bom_same'], res['csv_cr_count_same'], res['csv_trailing_newline_same'],
                   res['csv_shape_same'], res['csv_changed_eq_planned'], res['csv_header_same'],
                   res['sheetnames_same'], res['dimensions_same'], res['xlsx_changed_eq_planned'], not bad])
print(json.dumps(res, ensure_ascii=False, indent=1))
