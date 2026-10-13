"""Leakage check for the V1 human construct-validation packet (v1.8.1, sub-agent I).

Confirms that no coder-facing file contains: researcher G->M assignments, AI-pilot outputs, mechanism-name
hints in item text, network numbers, special marking of G04/G10/G11, or non-empty response cells.
It also screens the 132 'Literature interpretation' texts (the alternative item-text source) for mechanism hints
and documents why they were not used.

The checker reads the sealed key and the AI-pilot outputs ONLY to search for them; it prints no assignment.

Usage: python -E -P leakage_check_v1_packet.py --repo <REPO> --w132 <path to 132 csv> --sealed <sealed_key.json>
Writes OUT/tools/I/leakage_check_result.json ; exit code 0 = PASS, 1 = FAIL.
"""
from __future__ import annotations

import argparse
import csv
import glob
import hashlib
import json
import os
import re
import sys

from openpyxl import load_workbook

OUT_REL = os.path.join("claude", "r14_merge_v1_8_1_study1_freeze")
CATS = ["M1", "M2", "M3", "M4", "M5", "LONGITUDINAL", "DATA_GATED", "EVIDENCE_ONLY", "NONE", "AMBIGUOUS",
        "OPTIONAL_DATA_GATED", "OTHER", "INSUFFICIENT_INFORMATION"]
CAT_RE = re.compile(r"\b(" + "|".join(CATS) + r")\b")
GID_RE = re.compile(r"\bG(0[1-9]|1[0-6])\b")
# researcher-adjudication vocabulary (41 / 53A / manuscript Table 3-4) and AI-pilot vocabulary (54C)
FORBIDDEN = [r"\bI[1-5]\b", r"\banchor\b", r"\bMerge\b", r"\bKeep\b", r"static core", r"static bridge",
             r"\bbridge\b", r"researcher", r"\bsealed\b", r"\bkey\b", r"\b41_", r"\b53A\b", r"\b54C\b", r"\b132\b",
             r"convergence", r"\bpilot\b", r"AI coder", r"Optional module", r"contested", r"sensitivity",
             r"爭議", r"研究者", r"裁決", r"答案", r"O/E", r"q\s*=", r"p_BH", r"\bobserved\s*=", r"\bdynamic\b",
             r"\bstatic\b", r"evidence-only"]
# mechanism-hint vocabulary that must not appear in neutral item text (pattern / evidence fields)
HINT_EN = ["authority", "accountab", "approve", "sign-off", "sign off", "veto", "stop", "governance", "policy",
           "routing", "route", "recovery", "rollback", "exception", "workflow", "handoff", "configuration",
           "division of work", "override", "delegat", "assurance", "validation", "monitor", "audit", "drift",
           "longitudinal", "feedback", "learning", "next period", "data-gated", "threshold", "label", "score",
           "mechanism", "interface", "evidence layer", "frontier", "pilot"]
HINT_ZH = ["誰", "負責", "問責", "簽核", "停止", "回復", "例外", "流程", "門檻", "回饋", "下一期", "試點", "分工",
           "配置決策", "組織配置", "技術應用前沿", "嵌入", "審", "驗證", "最終權", "授權"]
ITEM_TEXT_COLS = {"Literature_pattern", "Evidence_description"}
FAMILY_COLS = {"AI_family", "Decision_family", "AI_family_definition", "Decision_family_definition",
               "AI_family_ZH", "Decision_family_ZH"}
FAMILY_LABELS = []   # filled in main() from 54A
RESP_PREFIX = ("Q1_", "Q2_", "Q3_", "Q4_", "Q5_")
NETWORK_NUM_COLS = ("E_analytic", "Observed", "O_E_Ratio", "q_primary_BH", "q_exact", "q_document_preserving",
                    "q_raw_record")


def sha(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()


def shingles(text, n=6):
    w = re.findall(r"[a-z0-9']+", text.lower())
    return {" ".join(w[i:i + n]) for i in range(len(w) - n + 1)}


def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    return rows


def check_workbook(path, ai_shingles, net_numbers, findings, is_coder_file):
    wb = load_workbook(path)
    name = os.path.basename(path)
    coder_sheets = [s for s in wb.sheetnames if s.startswith("CODER_")]
    if is_coder_file and len(coder_sheets) != 1:
        findings.append((name, "PER_CODER_FILE_HAS_NOT_EXACTLY_ONE_CODER_SHEET", str(coder_sheets)))
    nonempty_resp = 0
    style_sig = {}
    for ws in wb.worksheets:
        rows = sheet_rows(ws)
        if not rows:
            continue
        head_idx = 1 if ws.title == "ATTESTATION" else 0
        head = [str(h) if h is not None else "" for h in rows[head_idx]] if len(rows) > head_idx else []
        for ri, row in enumerate(rows):
            for ci, v in enumerate(row):
                if v is None:
                    continue
                s = str(v)
                col = head[ci] if ci < len(head) else ""
                where = f"{ws.title}!R{ri + 1}C{ci + 1}"
                if ws.title == "ADJUDICATION_EMPTY" and ci == 0:
                    pass
                # response cells must be empty
                if ri > head_idx and (col.startswith(RESP_PREFIX) or (ws.title == "ATTESTATION" and ci in (1, 3, 4, 5, 6))):
                    nonempty_resp += 1
                    findings.append((name, "NON_EMPTY_RESPONSE_CELL", where))
                if ws.title == "ADJUDICATION_EMPTY" and ri > 0 and ci > 0:
                    findings.append((name, "NON_EMPTY_ADJUDICATION_CELL", where))
                # G ID together with a category code in one cell = assignment leak
                if GID_RE.search(s) and CAT_RE.search(s):
                    findings.append((name, "G_ID_AND_CATEGORY_IN_SAME_CELL", where))
                # category codes may appear only in instruction / definition sheets
                if ws.title.startswith(("CODER_", "ITEM_INFO", "ADJUDICATION")) and CAT_RE.search(s) and ri > 0:
                    findings.append((name, "CATEGORY_CODE_IN_ITEM_ROW", f"{where}: {CAT_RE.search(s).group(0)}"))
                for pat in FORBIDDEN:
                    if re.search(pat, s, flags=re.I):
                        findings.append((name, "FORBIDDEN_TOKEN", f"{where}: /{pat}/"))
                if col in ITEM_TEXT_COLS and ri > 0:
                    stripped = s
                    for fam in FAMILY_LABELS:          # family names are the stimulus; overlap reported separately
                        stripped = stripped.replace(fam, "<FAMILY>")
                    low = stripped.lower()
                    for h in HINT_EN:
                        if h in low:
                            findings.append((name, "MECHANISM_HINT_IN_ITEM_TEXT", f"{where}: '{h}'"))
                    for h in HINT_ZH:
                        if h in stripped:
                            findings.append((name, "MECHANISM_HINT_IN_ITEM_TEXT", f"{where}: '{h}'"))
                hit = shingles(s) & ai_shingles
                if hit:
                    findings.append((name, "AI_PILOT_TEXT_OVERLAP", f"{where}: {sorted(hit)[:2]}"))
                for num in net_numbers:
                    if num in s:
                        findings.append((name, "NETWORK_NUMBER_PRESENT", f"{where}: {num}"))
                # style signature for G04/G10/G11 equal-marking test
                if ws.title.startswith("CODER_") and ci == 1 and ri > 0:
                    cells = list(ws.iter_rows(min_row=ri + 1, max_row=ri + 1))[0]
                    sig = tuple((c.fill.fgColor.rgb, c.font.b, c.font.i, c.font.color.rgb if c.font.color else None)
                                for c in cells)
                    style_sig[(ws.title, s)] = (sig, sum(1 for c in cells if c.value is not None))
    # G04/G10/G11 must be styled and populated exactly like every other item
    for sheet in {k[0] for k in style_sig}:
        sigs = {g: v for (sh, g), v in style_sig.items() if sh == sheet}
        ref = {v for g, v in sigs.items() if g not in ("G04", "G10", "G11")}
        for g in ("G04", "G10", "G11"):
            if g in sigs and (len(ref) != 1 or sigs[g] not in ref):
                findings.append((name, "SPECIAL_MARKING_OF_SENSITIVITY_ITEM", f"{sheet}: {g}"))
        if len(ref) > 1:
            findings.append((name, "ITEM_ROWS_NOT_UNIFORM", sheet))
    return nonempty_resp, wb.sheetnames


def screen_132(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp950", "big5"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    rows = list(csv.reader(text.splitlines()))
    head = rows[0]
    ci = next(i for i, h in enumerate(head) if h.startswith("Literature interpretation"))
    out = []
    for r in rows[1:]:
        t = r[ci]
        hits = [h for h in HINT_ZH if h in t] + [h for h in HINT_EN if h in t.lower()]
        fam = [h for h in hits if h in ("問責", "審", "流程") and (h in r[4] or h in r[2])]
        out.append(dict(gap=r[0], encoding=enc, hint_terms=hits,
                        hint_terms_only_from_family_name=sorted(set(fam)) if set(hits) <= set(fam) and hits else [],
                        verdict="LEAKS_MECHANISM_HINT" if [h for h in hits if h not in fam] else
                        ("FAMILY_NAME_OVERLAP_ONLY" if hits else "NO_HINT_TERM_FOUND")))
    return out


def lexical_overlap(repo):
    """Inherent stimulus overlap: family names that share a word with a category name."""
    cat_words = {"authority", "accountability", "workflow", "routing", "recovery", "configuration", "response",
                 "evaluation", "assurance", "longitudinal", "evidence"}
    out = []
    with open(os.path.join(repo, "54A_V1_BLIND_PACKET_A_v2_2.csv"), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            for fam in (r["ai_family"], r["decision_family"]):
                w = {x.lower() for x in re.findall(r"[A-Za-z]+", fam)} & cat_words
                if w and fam not in [o["family"] for o in out]:
                    out.append(dict(family=fam, shared_words=sorted(w)))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True)
    ap.add_argument("--w132", required=True)
    ap.add_argument("--sealed", required=True)
    a = ap.parse_args()
    out = os.path.join(a.repo, OUT_REL)
    ai_sh = set()
    ai_files = [os.path.join(a.repo, "v1_blind", f) for f in ("ai_coder_A_output.csv", "ai_coder_B_output.csv")]
    for p in ai_files:
        with open(p, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                ai_sh |= shingles(r["rationale"]) | shingles(r["enterprise_decision_interface"])
    # AI rationales partly quote the 53B codebook, which was frozen BEFORE any AI output (manifest stage 1).
    # Shingles that already occur in that pre-coding codebook text are not AI-pilot content and are excluded.
    pre = set()
    for p in ("53B_V1_BLIND_CODING_PROTOCOL_v2_2.md", os.path.join("v1_human_packet", "CODER_CLARIFICATIONS_v2_2_1.md")):
        pre |= shingles(open(os.path.join(a.repo, p), encoding="utf-8").read())
    n_codebook_shared = len(ai_sh & pre)
    ai_sh -= pre
    net = set()
    with open(os.path.join(a.repo, "11_Priority_Gaps_Frozen_Export_G01_G16_v1_0_20261009.csv"), encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            for c in NETWORK_NUM_COLS:
                v = r[c]
                if re.fullmatch(r"[0-9.]+", v) and len(v.rstrip("0").rstrip(".")) >= 4:
                    net.add(v)
                    try:
                        net.add(f"{float(v):.2f}")
                    except ValueError:
                        pass
    net = {n for n in net if len(n) >= 4 and n not in ("0.00",)}
    with open(os.path.join(a.repo, "54A_V1_BLIND_PACKET_A_v2_2.csv"), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            for fam in (r["ai_family"], r["decision_family"]):
                if fam not in FAMILY_LABELS:
                    FAMILY_LABELS.append(fam)
    findings = []
    files = sorted(glob.glob(os.path.join(out, "20_V1_HUMAN_CONSTRUCT_VALIDATION_PACKET*.xlsx")))
    per_file = {}
    for p in files:
        is_coder = "_CODER_" in os.path.basename(p)
        n_resp, sheets = check_workbook(p, ai_sh, net, findings, is_coder)
        per_file[os.path.basename(p)] = dict(sha256=sha(p), sheets=sheets, coder_facing=True,
                                             distributed_to_coders=is_coder, nonempty_response_cells=n_resp)
    # sealed key cross-check: packet G<->family mapping equals the sealed case map (no assignment printed)
    sk = json.load(open(a.sealed, encoding="utf-8"))
    pk = {}
    with open(os.path.join(a.repo, "54A_V1_BLIND_PACKET_A_v2_2.csv"), encoding="utf-8") as f:
        packet_rows = {r["case_id"]: (r["ai_family"], r["decision_family"]) for r in csv.DictReader(f)}
    wb = load_workbook(os.path.join(out, "20_V1_HUMAN_CONSTRUCT_VALIDATION_PACKET.xlsx"))
    for row in list(wb["ITEM_INFO"].iter_rows(values_only=True))[1:]:
        pk[row[1]] = (row[2], row[5])
    mapping_ok = all(packet_rows[sk["case"][g]] == pk[g] for g in pk)
    # codebook / protocol must not pair a G ID with a category in one line (study-team docs; assignments excluded)
    doc_findings = []
    for fn in ("21_V1_HUMAN_CONSTRUCT_VALIDATION_CODEBOOK.md",):
        p = os.path.join(out, fn)
        if os.path.exists(p):
            for i, line in enumerate(open(p, encoding="utf-8"), 1):
                if GID_RE.search(line) and CAT_RE.search(line):
                    doc_findings.append((fn, "G_ID_AND_CATEGORY_IN_SAME_LINE", f"line {i}"))
    res = dict(
        checker="leakage_check_v1_packet.py", files=per_file,
        ai_pilot_sources=[dict(path=os.path.relpath(p, a.repo).replace("\\", "/"), sha256=sha(p)) for p in ai_files],
        ai_pilot_shingles_checked=len(ai_sh), ai_pilot_shingles_excluded_as_53B_codebook_text=n_codebook_shared,
        network_numbers_checked=len(net),
        sealed_key_sha256=sha(a.sealed), packet_mapping_equals_sealed_case_map=mapping_ok,
        findings=[dict(file=f, check=c, where=w) for f, c, w in findings + doc_findings],
        lexical_overlap_family_vs_category_names=lexical_overlap(a.repo),
        screen_132_literature_interpretation=screen_132(a.w132), w132_sha256=sha(a.w132),
    )
    res["verdict"] = "PASS" if not res["findings"] and mapping_ok else "FAIL"
    tdir = os.path.join(out, "tools", "I")
    os.makedirs(tdir, exist_ok=True)
    with open(os.path.join(tdir, "leakage_check_result.json"), "w", encoding="utf-8") as f:
        json.dump(res, f, indent=1, ensure_ascii=False)
    print(res["verdict"], len(res["findings"]), "findings; mapping_ok =", mapping_ok)
    for x in res["findings"][:40]:
        print(x)
    return 0 if res["verdict"] == "PASS" else 1


if __name__ == "__main__":
    sys.exit(main())
