"""V1 human construct-validation analyzer (v1.8.1, Q1-Q5 instrument; protocol 19 sections 9-11).

Deterministic. Reads RETURNED human coder workbooks (20_..._CODER_k.xlsx), validates them against the codebook
(21) enums and consistency rules, and computes ONLY the metrics predeclared in protocol 19:
  Q2 (primary, nominal) and Q1 (secondary, nominal): raw agreement and Cohen's kappa per coder pair, kappa
  bootstrap 95% CI, Krippendorff's alpha (nominal) with bootstrap 95% CI (5,000 resamples, seed 20261011);
  Q2 classification RELIABLE / TENTATIVE / UNRESOLVED (19 section 10.3; 53B C1/C2 .800/.667 = RESEARCH_CONVENTION);
  Q1 unclassified (NO_SOURCE_THRESHOLD); derived primary_53B/status_53B continuity statistics (OA-2 pending);
  Jaccard of {Q2 u Q4}; per-G agreement; confusion matrices; category frequency; UNCERTAIN/AMBIGUOUS/NONE rates;
  Q3 confidence descriptive only (never in kappa/alpha, never a weight); G04/G10/G11 sensitivity (16 vs 13 items).
Reliability functions are imported UNCHANGED from REPO/v1_blind/reliability_v2_2.py (hash-checked).

Comparison with the researcher key only with --unseal, which first verifies sha256(sealed_key.json) against
V1_SEALED/RESEARCHER_ASSIGNMENT_COMMITMENT.md, and requires a hash-frozen adjudication file.

Exit status / JSON "status":
  NO_HUMAN_DATA            every response cell in every supplied sheet is empty -> nothing is computed (exit 0)
  INPUT_INVALID            validation errors (exit 2)
  RELIABILITY_COMPUTED     stage R done; key sealed (exit 0)
  KEY_COMPARED             --unseal done (exit 0)

Usage:
  python -E -P v1_construct_analyzer.py SHEET1.xlsx SHEET2.xlsx [SHEET3.xlsx ...] --out DIR
         [--manifest V1_SEALED/FREEZE_MANIFEST_V1_8_1.txt] [--repo REPO]
         [--unseal --key <path>/sealed_key.json --adjudication adj.csv]
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import os
import re
import sys
from collections import Counter
from itertools import combinations

from openpyxl import load_workbook

VERSION = "v1_construct_analyzer_v1_8_1"
ITEMS = [f"G{i:02d}" for i in range(1, 17)]
SENSITIVITY = ("G04", "G10", "G11")
Q1 = {"YES", "NO", "UNCERTAIN"}
Q2 = {"M1", "M2", "M3", "M4", "LONGITUDINAL", "DATA_GATED", "EVIDENCE_ONLY", "NONE", "AMBIGUOUS"}
Q3 = {"HIGH", "MEDIUM", "LOW"}
Q4 = {"M1", "M2", "M3", "M4", "LONGITUDINAL", "DATA_GATED"}
Q2_IF_YES = Q2 - {"EVIDENCE_ONLY"}
RESP = ["Q1_identifiable_interface", "Q2_best_category", "Q3_confidence", "Q4_secondary_category_optional",
        "Q5_rationale"]
REL_SHA256 = "866e9005da8a20e996c93b5c8f6bee916992f24d60c6a96e8f1dcda4556848e7"   # v1_blind/reliability_v2_2.py
SEED, B = 20261011, 5000                                                         # 53B section 7 line 86
LO, HI = 0.667, 0.800                                                            # 53B C1/C2 (RESEARCH_CONVENTION)
OUT_REL = os.path.join("claude", "r14_merge_v1_8_1_study1_freeze")
COMMIT_REL = os.path.join(OUT_REL, "V1_SEALED", "RESEARCHER_ASSIGNMENT_COMMITMENT.md")


class InputInvalid(ValueError):
    pass


def sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def default_repo():
    here = os.path.dirname(os.path.abspath(__file__))
    cand = os.path.abspath(os.path.join(here, "..", "..", "..", ".."))       # OUT/tools/I -> REPO
    if os.path.exists(os.path.join(cand, "v1_blind", "reliability_v2_2.py")):
        return cand
    return os.path.abspath(os.path.join(here, "..", "..", "repo"))           # working copy location


def load_rel(repo):
    p = os.path.join(repo, "v1_blind", "reliability_v2_2.py")
    if sha(p) != REL_SHA256:
        raise SystemExit(f"reliability_v2_2.py hash mismatch: {p}")
    spec = importlib.util.spec_from_file_location("reliability_v2_2", p)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ------------------------------------------------------------------------------------------------ input
def _cell(v):
    return "" if v is None else str(v).strip()


def read_sheet(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    names = [s for s in wb.sheetnames if s.startswith("CODER_")]
    if len(names) != 1:
        raise InputInvalid(f"{os.path.basename(path)}: expected exactly one CODER_* sheet, found {names}")
    rows = list(wb[names[0]].iter_rows(values_only=True))
    head = [_cell(h) for h in rows[0]]
    missing = [h for h in ["G_ID"] + RESP if h not in head]
    if missing:
        raise InputInvalid(f"{os.path.basename(path)}: missing columns {missing}")
    ix = {h: head.index(h) for h in ["G_ID"] + RESP}
    data = {}
    for r in rows[1:]:
        if r is None or all(v is None for v in r):
            continue
        g = _cell(r[ix["G_ID"]])
        if g in data:
            raise InputInvalid(f"{os.path.basename(path)}: duplicate {g}")
        data[g] = {h: _cell(r[ix[h]]) for h in RESP}
    return dict(path=path, sheet=names[0], sha256=sha(path), rows=data)


def is_empty(sheet):
    return all(not any(v for v in row.values()) for row in sheet["rows"].values())


def validate(sheet):
    errs, name = [], os.path.basename(sheet["path"])
    if sorted(sheet["rows"]) != ITEMS:
        errs.append(f"{name}: item set is not G01-G16")
    for g, r in sheet["rows"].items():
        q1, q2, q3, q4, q5 = (r[h] for h in RESP)
        if q1 not in Q1:
            errs.append(f"{name} {g}: Q1 '{q1}' not in {sorted(Q1)}")
        if q2 not in Q2:
            errs.append(f"{name} {g}: Q2 '{q2}' not in codebook enum")
        if q3 not in Q3:
            errs.append(f"{name} {g}: Q3 '{q3}' not in {sorted(Q3)}")
        if q4 and q4 not in Q4:
            errs.append(f"{name} {g}: Q4 '{q4}' not in {sorted(Q4)}")
        if q4 and q4 == q2:
            errs.append(f"{name} {g}: Q4 equals Q2")
        if not q5:
            errs.append(f"{name} {g}: Q5 rationale empty")
        if q1 == "YES" and q2 not in Q2_IF_YES:
            errs.append(f"{name} {g}: Q1=YES requires Q2 != EVIDENCE_ONLY")
        if q1 == "NO" and q2 != "EVIDENCE_ONLY":
            errs.append(f"{name} {g}: Q1=NO requires Q2=EVIDENCE_ONLY")
    return errs


def manifest_hashes(path):
    out = set()
    with open(path, encoding="utf-8") as f:
        for line in f:
            m = re.match(r"^([0-9a-f]{64})\s", line.strip() + " ")
            if m:
                out.add(m.group(1))
    return out


# ------------------------------------------------------------------------------------------------ metrics
def derived(q2, q4):
    """Codebook 21 section 5.2."""
    if q2 in ("M1", "M2", "M3", "M4"):
        return q2, "static"
    return {"LONGITUDINAL": ("M5", "dynamic"),
            "DATA_GATED": (q4 if q4 in ("M1", "M2", "M3", "M4") else "OPTIONAL_DATA_GATED", "optional"),
            "EVIDENCE_ONLY": ("EVIDENCE_ONLY", "evidence-only"),
            "NONE": ("OTHER", "static"),
            "AMBIGUOUS": ("INSUFFICIENT_INFORMATION", "insufficient")}[q2]


def label_3(alpha, lo):
    """Protocol 19 section 10.3: RELIABLE / TENTATIVE / UNRESOLVED from 53B C1/C2 only."""
    if alpha != alpha:
        return "UNRESOLVED", "UNDEFINED"
    sub = "RELIABLE" if alpha >= HI else "TENTATIVE" if alpha >= LO else "NOT_ACCEPTABLE"
    imprecise = lo != lo or lo < LO
    if sub == "NOT_ACCEPTABLE" or imprecise:
        return "UNRESOLVED", sub + ("+IMPRECISE" if imprecise else "")
    return sub, sub


def kappa_ci(rel, a, b):
    rng = __import__("random").Random(SEED)
    n, vals = len(a), []
    for _ in range(B):
        idx = [rng.randrange(n) for _ in range(n)]
        k, _, _ = rel.cohen_kappa([a[i] for i in idx], [b[i] for i in idx])
        if k == k:
            vals.append(k)
    vals.sort()
    if not vals:
        return [float("nan"), float("nan")], 0
    return [vals[int(0.025 * len(vals))], vals[int(0.975 * len(vals)) - 1]], len(vals)


def confusion(a, b):
    labels = sorted(set(a) | set(b))
    m = {x: {y: 0 for y in labels} for x in labels}
    for x, y in zip(a, b):
        m[x][y] += 1
    return dict(labels=labels, matrix=m)


def field_stats(rel, coders, items, getter, classify):
    data = [[getter(c, g) for c in coders] for g in items]
    alpha = rel.kripp_alpha_nominal(data)
    lo, hi, nb = rel.bootstrap_alpha(data)
    out = dict(n_items=len(items), alpha_nominal=alpha, alpha_ci95=[lo, hi], alpha_boot_valid=nb, pairs=[])
    if classify == "Q2":
        out["classification"], out["classification_sub"] = label_3(alpha, lo)
    elif classify == "53B":
        out["c1_c3_class_53B"] = rel.classify(alpha, lo)
    else:
        out["classification"] = "NOT_CLASSIFIED (NO_SOURCE_THRESHOLD)"
    for (i, ci), (j, cj) in combinations(enumerate(coders), 2):
        a, b = [getter(ci, g) for g in items], [getter(cj, g) for g in items]
        k, po, pe = rel.cohen_kappa(a, b)
        kci, kb = kappa_ci(rel, a, b)
        out["pairs"].append(dict(pair=[ci["sheet"], cj["sheet"]], agreements=sum(x == y for x, y in zip(a, b)),
                                 raw_agreement=po, cohen_kappa=k, kappa_ci95=kci, kappa_boot_valid=kb,
                                 confusion=confusion(a, b)))
    return out


def q(field):
    return lambda c, g: c["rows"][g][field]


def reliability(rel, coders, items):
    d1 = lambda c, g: derived(c["rows"][g][RESP[1]], c["rows"][g][RESP[3]])[0]
    d2 = lambda c, g: derived(c["rows"][g][RESP[1]], c["rows"][g][RESP[3]])[1]
    return dict(
        Q2_primary=field_stats(rel, coders, items, q(RESP[1]), "Q2"),
        Q1_secondary=field_stats(rel, coders, items, q(RESP[0]), None),
        derived_primary_53B_C1=field_stats(rel, coders, items, d1, "53B"),
        derived_status_53B_C3=field_stats(rel, coders, items, d2, "53B"),
    )


def descriptives(coders):
    per_g = {}
    for g in ITEMS:
        q2s = [c["rows"][g][RESP[1]] for c in coders]
        q1s = [c["rows"][g][RESP[0]] for c in coders]
        pairs = list(combinations(range(len(coders)), 2))
        per_g[g] = dict(
            Q2={c["sheet"]: c["rows"][g][RESP[1]] for c in coders},
            Q1={c["sheet"]: c["rows"][g][RESP[0]] for c in coders},
            Q2_modal=Counter(q2s).most_common(1)[0][0], Q2_n_distinct=len(set(q2s)),
            Q2_pair_agreement_share=sum(q2s[i] == q2s[j] for i, j in pairs) / len(pairs),
            Q2_all_agree=len(set(q2s)) == 1, Q1_all_agree=len(set(q1s)) == 1,
            status="AGREED" if len(set(q2s)) == 1 else "TO_ADJUDICATION",
            sensitivity_target=g in SENSITIVITY)
    freq = {c["sheet"]: dict(Q2=dict(Counter(c["rows"][g][RESP[1]] for g in ITEMS)),
                             Q4=dict(Counter(c["rows"][g][RESP[3]] for g in ITEMS if c["rows"][g][RESP[3]])))
            for c in coders}
    freq["POOLED"] = dict(Q2=dict(Counter(c["rows"][g][RESP[1]] for c in coders for g in ITEMS)))

    def rates(cs):
        n = len(cs) * len(ITEMS)
        return dict(Q1_UNCERTAIN=sum(c["rows"][g][RESP[0]] == "UNCERTAIN" for c in cs for g in ITEMS) / n,
                    Q2_AMBIGUOUS=sum(c["rows"][g][RESP[1]] == "AMBIGUOUS" for c in cs for g in ITEMS) / n,
                    Q2_NONE=sum(c["rows"][g][RESP[1]] == "NONE" for c in cs for g in ITEMS) / n)
    amb = {c["sheet"]: rates([c]) for c in coders}
    amb["POOLED"] = rates(coders)
    none_desc = [dict(coder=c["sheet"], item=g, Q5=c["rows"][g][RESP[4]]) for c in coders for g in ITEMS
                 if c["rows"][g][RESP[1]] == "NONE"]
    jac = []
    for g in ITEMS:
        sets = [{c["rows"][g][RESP[1]]} | ({c["rows"][g][RESP[3]]} - {""}) for c in coders]
        jac += [len(a & b) / len(a | b) for a, b in combinations(sets, 2)]
    conf = {}
    order = {"LOW": 0, "MEDIUM": 1, "HIGH": 2}
    for c in coders:
        def med(gs):
            v = sorted(order[c["rows"][g][RESP[2]]] for g in gs)
            return None if not v else ["LOW", "MEDIUM", "HIGH"][v[(len(v) - 1) // 2]]
        agreed = [g for g in ITEMS if per_g[g]["Q2_all_agree"]]
        conf[c["sheet"]] = dict(distribution=dict(Counter(c["rows"][g][RESP[2]] for g in ITEMS)),
                                median_lower=med(ITEMS), median_lower_agreed=med(agreed),
                                median_lower_disagreed=med([g for g in ITEMS if g not in agreed]))
    return dict(per_item=per_g, category_frequency=freq, uncertainty_ambiguity_rates=amb,
                none_descriptions=none_desc, jaccard_Q2_union_Q4_mean=sum(jac) / len(jac),
                Q3_confidence_descriptive_only=conf)


# ------------------------------------------------------------------------------------------------ unseal
def committed_key_hash(repo):
    txt = open(os.path.join(repo, COMMIT_REL), encoding="utf-8").read()
    m = re.search(r"\| sealed_key\.json \|[^|]*\| `([0-9a-f]{64})` \|", txt)
    if not m:
        raise SystemExit("commitment file has no sealed_key.json hash")
    return m.group(1)


def flatten_key(entry):
    """Codebook 21 section 5.3."""
    primary, _secondary, status = entry
    if status == "dynamic":
        return "LONGITUDINAL"
    if status == "optional":
        return "DATA_GATED"
    if status == "evidence-only":
        return "EVIDENCE_ONLY"
    return primary


def load_adjudication(path, per_item):
    final = {g: dict(Q2=v["Q2_modal"], Q4="", route="AGREED") for g, v in per_item.items() if v["Q2_all_agree"]}
    with open(path, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            g = r["G_ID"].strip()
            if g not in ITEMS or r["Final_Q2"].strip() not in Q2:
                raise InputInvalid(f"adjudication row invalid: {r}")
            final[g] = dict(Q2=r["Final_Q2"].strip(), Q4=r.get("Final_Q4", "").strip(), route=r["Route"].strip())
    missing = [g for g in ITEMS if g not in final]
    if missing:
        raise InputInvalid(f"adjudication missing items {missing}")
    return final


def unseal(repo, key_path, adj_path, manifest, per_item):
    want = committed_key_hash(repo)
    got = sha(key_path)
    if got != want:
        raise SystemExit(f"STOP: sealed_key.json sha256 {got} != committed {want}; no comparison made")
    if not manifest or sha(adj_path) not in manifest_hashes(manifest):
        raise SystemExit("STOP: adjudication file is not hash-frozen in the manifest; key stays sealed")
    final = load_adjudication(adj_path, per_item)
    key = json.load(open(key_path, encoding="utf-8"))["key"]
    rows = []
    for g in ITEMS:
        fk = flatten_key(key[g])
        fp, fs = derived(final[g]["Q2"], final[g]["Q4"])
        rows.append(dict(item=g, final_Q2=final[g]["Q2"], route=final[g]["route"], key_flat=fk,
                         flat_match=final[g]["Q2"] == fk, key_primary=key[g][0], key_status=key[g][2],
                         primary_53B_match=fp == key[g][0], status_53B_match=fs == key[g][2],
                         sensitivity_target=g in SENSITIVITY, model_change_class="[TO BE CLASSIFIED (53B s11)]"
                         if final[g]["Q2"] != fk else "NO_CHANGE"))
    with open(manifest, "a", encoding="utf-8") as f:
        f.write(f"# KEY_OPENED adjudication={sha(adj_path)} by {VERSION}\n")
    return dict(key_sha256=got, adjudication_sha256=sha(adj_path), rows=rows,
                flat_agreement=sum(r["flat_match"] for r in rows),
                flat_agreement_13=sum(r["flat_match"] for r in rows if not r["sensitivity_target"]))


# ------------------------------------------------------------------------------------------------ output
def fmt(x):
    return "n/a" if x is None or x != x else f"{x:.3f}"


def render_md(res):
    L = [f"# V1 construct-validation analysis ({VERSION})", "", f"**Status: `{res['status']}`**", ""]
    if res["status"] == "NO_HUMAN_DATA":
        L += ["All response cells in every supplied coder sheet are empty. Nothing was computed.", "",
              "| Sheet | File | SHA-256 |", "|---|---|---|"]
        L += [f"| {s['sheet']} | {os.path.basename(s['path'])} | `{s['sha256']}` |" for s in res["inputs"]]
        return "\n".join(L) + "\n"
    if res["status"] == "INPUT_INVALID":
        return "\n".join(L + ["Validation errors:", ""] + [f"- {e}" for e in res["errors"]]) + "\n"
    for scope in ("all_16", "excluding_G04_G10_G11"):
        r = res["reliability"][scope]
        L += [f"## Reliability — {scope}", "", "| Field | alpha (nominal) [95% CI] | class | pair | agree | kappa [95% CI] |",
              "|---|---|---|---|---|---|"]
        for name, f in r.items():
            cls = f.get("classification", f.get("c1_c3_class_53B"))
            for p in f["pairs"]:
                L.append(f"| {name} | {fmt(f['alpha_nominal'])} [{fmt(f['alpha_ci95'][0])}, {fmt(f['alpha_ci95'][1])}] | "
                         f"{cls} | {p['pair'][0]}–{p['pair'][1]} | {p['agreements']}/{f['n_items']} | "
                         f"{fmt(p['cohen_kappa'])} [{fmt(p['kappa_ci95'][0])}, {fmt(p['kappa_ci95'][1])}] |")
        L.append("")
    L += ["## Per-item agreement (Q2)", "", "| Item | codes | all agree | status |", "|---|---|---|---|"]
    for g, v in res["descriptives"]["per_item"].items():
        L.append(f"| {g} | {', '.join(f'{k}:{x}' for k, x in v['Q2'].items())} | {v['Q2_all_agree']} | {v['status']} |")
    L += ["", "Confidence (Q3) is descriptive only and is never used in kappa/alpha. Full detail in the JSON.", ""]
    if res.get("key_comparison"):
        k = res["key_comparison"]
        L += ["## Key comparison (after --unseal)", "", f"Flat agreement with key: {k['flat_agreement']}/16; "
              f"excluding G04/G10/G11: {k['flat_agreement_13']}/13.", ""]
    return "\n".join(L) + "\n"


def run(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("sheets", nargs="+")
    ap.add_argument("--out", required=True)
    ap.add_argument("--repo", default=None)
    ap.add_argument("--manifest", default=None)
    ap.add_argument("--unseal", action="store_true")
    ap.add_argument("--key")
    ap.add_argument("--adjudication")
    a = ap.parse_args(argv)
    repo = a.repo or default_repo()
    os.makedirs(a.out, exist_ok=True)
    res = dict(analyzer=VERSION, seed=SEED, bootstrap_reps=B, thresholds=dict(
        alpha_reliable=HI, alpha_tentative=LO, ci_lower_imprecise=LO, provenance="RESEARCH_CONVENTION (53B C1/C2)"))
    code = 0
    try:
        sheets = [read_sheet(p) for p in a.sheets]
        res["inputs"] = [dict(path=s["path"], sheet=s["sheet"], sha256=s["sha256"]) for s in sheets]
        if all(is_empty(s) for s in sheets):
            res["status"] = "NO_HUMAN_DATA"
        else:
            errs = []
            if len(sheets) < 2:
                errs.append("at least 2 human coder sheets are required")
            if len({s["sheet"] for s in sheets}) != len(sheets):
                errs.append("duplicate coder sheet names")
            for s in sheets:
                errs += validate(s)
            for s1, s2 in combinations(sheets, 2):
                if s1["sha256"] == s2["sha256"] or s1["rows"] == s2["rows"]:
                    errs.append(f"identical content: {s1['sheet']} and {s2['sheet']}")
            if not a.manifest:
                errs.append("--manifest required: sheets must be hash-frozen (protocol 19 P4) before analysis")
            else:
                mh = manifest_hashes(a.manifest)
                errs += [f"{s['sheet']}: sha256 not in manifest (not frozen)" for s in sheets if s["sha256"] not in mh]
            if errs:
                raise InputInvalid("\n".join(errs))
            rel = load_rel(repo)
            sub = [g for g in ITEMS if g not in SENSITIVITY]
            res["reliability"] = dict(all_16=reliability(rel, sheets, ITEMS),
                                      excluding_G04_G10_G11=reliability(rel, sheets, sub))
            res["descriptives"] = descriptives(sheets)
            res["status"] = "RELIABILITY_COMPUTED"
            if a.unseal:
                if not (a.key and a.adjudication):
                    raise InputInvalid("--unseal requires --key and --adjudication")
                res["key_comparison"] = unseal(repo, a.key, a.adjudication, a.manifest, res["descriptives"]["per_item"])
                res["status"] = "KEY_COMPARED"
    except InputInvalid as e:
        res["status"], res["errors"], code = "INPUT_INVALID", str(e).split("\n"), 2
    with open(os.path.join(a.out, "V1_CONSTRUCT_RESULTS.json"), "w", encoding="utf-8") as f:
        json.dump(res, f, indent=1, ensure_ascii=False, default=str)
    with open(os.path.join(a.out, "V1_CONSTRUCT_RESULTS.md"), "w", encoding="utf-8") as f:
        f.write(render_md(res))
    print(res["status"])
    return code


if __name__ == "__main__":
    sys.exit(run(sys.argv[1:]))
