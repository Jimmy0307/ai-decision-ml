"""Build the V1 human construct-validation packet (v1.8.1, sub-agent I).

PREPARES the instrument only. It never writes a response, never simulates a coder, and never reads the
sealed researcher key or the AI-pilot outputs (those are read only by the separate leakage checker, to
prove they are absent).

Inputs (read-only, REPO tracked files):
  54A_V1_BLIND_PACKET_A_v2_2.csv                       neutral family definitions + neutral pair text
  11_Priority_Gaps_Frozen_Export_G01_G16_v1_0_20261009.csv  G ID <-> family-code pair (no mechanism data)

Outputs (OUT = REPO/claude/r14_merge_v1_8_1_study1_freeze):
  20_V1_HUMAN_CONSTRUCT_VALIDATION_PACKET.xlsx                 coordinator master (all coder sheets, empty)
  20_V1_HUMAN_CONSTRUCT_VALIDATION_PACKET_CODER_<k>.xlsx       one file per coder (only that coder's sheet)
  tools/I/item_order_log.json                                  seeds and per-coder item orders

Usage: python -E -P build_v1_construct_packet.py --repo <REPO> [--k 3] [--mask-gid]
  (-E -P instead of -I: openpyxl is installed in the user site-packages, which -I disables.)
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

BASE_SEED = 20261014          # item-order seed; coder k uses BASE_SEED + k
OUT_REL = os.path.join("claude", "r14_merge_v1_8_1_study1_freeze")
PREFIX = "20_V1_HUMAN_CONSTRUCT_VALIDATION_PACKET"

Q1_OPTS = ["YES", "NO", "UNCERTAIN"]
Q2_OPTS = ["M1", "M2", "M3", "M4", "LONGITUDINAL", "DATA_GATED", "EVIDENCE_ONLY", "NONE", "AMBIGUOUS"]
Q3_OPTS = ["HIGH", "MEDIUM", "LOW"]
Q4_OPTS = ["M1", "M2", "M3", "M4", "LONGITUDINAL", "DATA_GATED"]

# 54A English family label -> 11-CSV family code (exact label match; verified below for all 16 rows)
AI_LABEL = {
    "Human Augmentation": "HUMAN_AUGMENTATION",
    "Prediction and Risk Estimation": "PREDICTION_RISK_ESTIMATION",
    "Coordination and Execution": "COORDINATION_EXECUTION",
    "Generation and Synthesis": "GENERATION_SYNTHESIS",
    "Explanation and Assurance": "EXPLANATION_ASSURANCE",
    "Evaluation and Audit": "EVALUATION_ASSESSMENT",
    "Recommendation and Optimization": "RECOMMENDATION_OPTIMIZATION",
}
DEC_LABEL = {
    "Prediction and Diagnosis": "PREDICTION_DIAGNOSIS",
    "Human Review and Accountability": "HUMAN_REVIEW_ACCOUNTABILITY",
    "Risk Assessment": "RISK_ASSESSMENT",
    "Allocation and Planning": "ALLOCATION_PLANNING",
    "Evaluation and Selection": "EVALUATION_SELECTION",
    "Governance and Oversight": "GOVERNANCE_OVERSIGHT",
    "Information Acquisition": "INFORMATION_ACQUISITION",
    "General Decision Process": "GENERAL_DECISION_PROCESS",
}
# Chinese family labels as printed in the v1.8 manuscript Table 3-4A family column (family names only)
ZH = {
    "HUMAN_AUGMENTATION": "人機增能", "PREDICTION_RISK_ESTIMATION": "預測與風險估計",
    "COORDINATION_EXECUTION": "協調與執行", "GENERATION_SYNTHESIS": "生成與合成",
    "EXPLANATION_ASSURANCE": "解釋與保證", "EVALUATION_ASSESSMENT": "評估與審核",
    "RECOMMENDATION_OPTIMIZATION": "推薦與最佳化",
    "PREDICTION_DIAGNOSIS": "預測與診斷", "HUMAN_REVIEW_ACCOUNTABILITY": "人工複核與問責",
    "RISK_ASSESSMENT": "風險評估", "ALLOCATION_PLANNING": "配置與規劃", "EVALUATION_SELECTION": "評估與選擇",
    "GOVERNANCE_OVERSIGHT": "治理與監督", "INFORMATION_ACQUISITION": "資訊取得",
    "GENERAL_DECISION_PROCESS": "一般決策流程",
}

PATTERN_TEXT = ("Literature pattern: DEPLETED. In this study's literature network, the number of published "
                "documents that directly link AI functions of type '{ai}' to decision functions of type '{dec}' "
                "is below the number expected under a statistical null model, and this holds under all four "
                "pre-specified analysis methods. This is a research lead only. It does not mean that the "
                "combination is impossible, ineffective, unimportant, or absent from practice.")

README_ROWS = [
    "V1 HUMAN CONSTRUCT-VALIDATION CODING — INSTRUCTIONS (identical for every coder)",
    "",
    "Work alone. Expected time: 45-60 minutes. Do not discuss the items with anyone until you have submitted.",
    "Do not consult the project repository, any manuscript or report of this project, any other coder's sheet, "
    "or any AI assistant while coding. If you have already read project documents that describe how these items "
    "are mapped, tell the coordinator before you start.",
    "",
    "Each item shows one AI function family and one decision function family, with their definitions, and the "
    "same standard description of a literature pattern. Every item is answered with the same five questions.",
    "",
    "Q1  Does this pairing correspond to an identifiable enterprise decision interface?  YES / NO / UNCERTAIN",
    "     An enterprise decision interface is a decision that an enterprise must make about how this kind of AI "
    "function feeds this kind of decision function, and that is specific to this pairing.",
    "Q2  Best category (answer for every item; see sheet CATEGORY_DEFINITIONS and the consistency rules there).",
    "Q3  Your confidence in Q2:  HIGH / MEDIUM / LOW.",
    "Q4  Optional secondary category. Leave blank unless a second category is clearly necessary.",
    "Q5  Short rationale (1-3 sentences). If Q2 is NONE or DATA_GATED or AMBIGUOUS, say what is missing, which data, "
    "or which categories compete.",
    "",
    "Use the drop-down lists in the coloured response columns. Do not change the item rows, IDs or order. "
    "Fill sheet ATTESTATION, save the file, and return it to the coordinator.",
]

CATEGORY_ROWS = [
    ("Q1", "YES", "You can name a decision that an enterprise must make about how this kind of AI function feeds this "
     "kind of decision function, and that decision is specific to this pairing.", ""),
    ("Q1", "NO", "The only decisions you can name would apply equally to almost any AI function feeding this decision "
     "function, or no enterprise decision is involved.", "If Q1 = NO, set Q2 = EVIDENCE_ONLY."),
    ("Q1", "UNCERTAIN", "You cannot decide between YES and NO.", "Q2 may be any category; AMBIGUOUS is allowed."),
    ("Q2/Q4", "M1  Authority & Accountability (授權與問責; = 53B Governance & Decision Rights / authority-to-act)",
     "Use when the main issue of the connection is who may approve, reject, stop or change what the AI output "
     "is allowed to decide; formal versus effective control; accountability; enterprise policy boundaries.",
     "Do not use when the main issue is how outputs are routed or technically validated."),
    ("Q2/Q4", "M2  Workflow Routing & Recovery",
     "Use when the main issue of the connection is how the AI output enters the next decision or action step; "
     "handoffs across systems or units; exception handling; rollback and recovery.",
     "Do not use when the main issue is who holds authority."),
    ("Q2/Q4", "M3  Human-AI Configuration & Response",
     "Use when the main issue of the connection is the sequence and division of work between people and AI "
     "(who proposes, who decides, aggregation, delegation, override) and of how users actually respond (use, verify, "
     "modify, escalate, reject, bypass).",
     "Do not use when the main issue is formal authority or validation evidence."),
    ("Q2/Q4", "M4  Evaluation & Assurance",
     "Use when the main issue of the connection is pre-deployment validation, monitoring, revalidation, "
     "auditability, explanation usability, drift, or evidence that the AI output is fit for the decision.",
     "Do not use when the main issue is who decides or how work is sequenced."),
    ("Q2/Q4", "LONGITUDINAL  Longitudinal extension",
     "Use when the enterprise decision can only be studied with outcomes observed over more than one period: results "
     "of one cycle change later resource allocation, policy or information acquisition. A single-period choice is "
     "not enough to represent it.",
     "Do not use when a single-period configuration choice is enough. If you use it, you may put the M1-M4 "
     "category involved in Q4."),
    ("Q2/Q4", "DATA_GATED  Data-gated module",
     "Use when an enterprise decision exists only if a specific kind of data is available. Name the data in Q5.",
     "Do not use when a category applies without special data. If you use it, you may put the M1-M4 category "
     "involved in Q4."),
    ("Q2", "EVIDENCE_ONLY  Evidence-layer only",
     "Use when the pairing describes a research or technology pattern but implies no distinct enterprise decision "
     "beyond what any AI function feeding this decision function would raise.",
     "Use together with Q1 = NO. Do not use when a specific category clearly applies."),
    ("Q2", "NONE  None of the listed categories",
     "Use when a distinct enterprise decision interface exists (Q1 = YES) but none of M1, M2, M3, M4, LONGITUDINAL or "
     "DATA_GATED can express it. Describe the decision in Q5.",
     "Do not use when a listed category can express it, even imperfectly."),
    ("Q2", "AMBIGUOUS",
     "Use when two or more categories fit equally well and you cannot choose one, or the definitions do not allow a "
     "defensible single choice. Name the competing categories in Q5.",
     "Do not use as a default; choose a category whenever one fits better."),
    ("Rule", "Consistency", "Q1 = YES -> Q2 is one of M1, M2, M3, M4, LONGITUDINAL, DATA_GATED, NONE, AMBIGUOUS.  "
     "Q1 = NO -> Q2 = EVIDENCE_ONLY.  Q1 = UNCERTAIN -> any Q2.", ""),
    ("Rule", "Order of tests for Q2",
     "(a) Can the decision be represented only with outcomes over more than one period? -> LONGITUDINAL. "
     "(b) Does the decision exist only if a named dataset is available? -> DATA_GATED. "
     "(c) Is there a distinct single-period decision? -> the M1-M4 category whose 'use when' text fits the main issue "
     "(or NONE). (d) Otherwise -> EVIDENCE_ONLY with Q1 = NO. Use AMBIGUOUS only if you cannot choose.", ""),
    ("Rule", "Residual decision family",
     "'General Decision Process' is a residual category. Code it like any other family if a category fits; otherwise "
     "use EVIDENCE_ONLY or AMBIGUOUS.", ""),
    ("Rule", "Code the enterprise decision",
     "Code the enterprise decision, not the technology, not the topic of the literature, and not shared words "
     "between a family name and a category name.", ""),
    ("Rule", "Q3 confidence", "HIGH / MEDIUM / LOW refers to Q2 only. It is descriptive and never weights agreement.", ""),
]

ITEM_HEAD = ["G_ID", "AI_family", "AI_family_ZH", "AI_family_definition", "Decision_family", "Decision_family_ZH",
             "Decision_family_definition", "Literature_pattern", "Evidence_description"]
RESP_HEAD = ["Q1_identifiable_interface", "Q2_best_category", "Q3_confidence", "Q4_secondary_category_optional",
             "Q5_rationale"]
RESP_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def load_items(repo):
    gaps = {}
    with open(os.path.join(repo, "11_Priority_Gaps_Frozen_Export_G01_G16_v1_0_20261009.csv"), encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            gaps[(r["AI_Family"], r["Decision_Family"])] = r["Gap_ID"]
    if len(gaps) != 16:
        raise SystemExit(f"expected 16 unique family pairs in 11-CSV, got {len(gaps)}")
    items = {}
    with open(os.path.join(repo, "54A_V1_BLIND_PACKET_A_v2_2.csv"), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            pair = (AI_LABEL[r["ai_family"]], DEC_LABEL[r["decision_family"]])
            gid = gaps[pair]  # KeyError -> unmapped pair: fail loudly
            if gid in items:
                raise SystemExit(f"duplicate mapping for {gid}")
            items[gid] = dict(
                G_ID=gid, AI_family=r["ai_family"], AI_family_ZH=ZH[pair[0]],
                AI_family_definition=r["ai_family_definition"], Decision_family=r["decision_family"],
                Decision_family_ZH=ZH[pair[1]], Decision_family_definition=r["decision_family_definition"],
                Literature_pattern=PATTERN_TEXT.format(ai=r["ai_family"], dec=r["decision_family"]),
                Evidence_description=r["neutral_pair_description"] + " " + r["network_evidence"],
            )
    if sorted(items) != [f"G{i:02d}" for i in range(1, 17)]:
        raise SystemExit(f"item set mismatch: {sorted(items)}")
    return items


def orders(k):
    base = [f"G{i:02d}" for i in range(1, 17)]
    out = {}
    for c in range(1, k + 1):
        o = list(base)
        random.Random(BASE_SEED + c).shuffle(o)
        out[f"CODER_{c}" if c <= 2 else f"CODER_{c}_OPTIONAL"] = o
    return out


def _widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def add_readme(wb):
    ws = wb.create_sheet("README_INSTRUCTIONS")
    for i, t in enumerate(README_ROWS, 1):
        ws.cell(i, 1, t).alignment = WRAP
    ws.cell(1, 1).font = HEAD_FONT
    _widths(ws, {"A": 150})


def add_categories(wb):
    ws = wb.create_sheet("CATEGORY_DEFINITIONS")
    for j, h in enumerate(["Question", "Category", "Use when (operational definition)", "Do not use when / rule"], 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i, row in enumerate(CATEGORY_ROWS, 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v).alignment = WRAP
    _widths(ws, {"A": 9, "B": 36, "C": 90, "D": 60})


def add_item_info(wb, items, order):
    ws = wb.create_sheet("ITEM_INFO")
    head = ["Item_order"] + ITEM_HEAD
    for j, h in enumerate(head, 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i, gid in enumerate(order, 2):
        ws.cell(i, 1, i - 1)
        for j, h in enumerate(ITEM_HEAD, 2):
            ws.cell(i, j, items[gid][h]).alignment = WRAP
    _widths(ws, {"A": 6, "B": 7, "C": 26, "D": 14, "E": 50, "F": 28, "G": 14, "H": 50, "I": 60, "J": 60})


def add_coder_sheet(wb, name, items, order):
    ws = wb.create_sheet(name)
    head = ["Item_order", "G_ID", "AI_family", "Decision_family", "AI_family_definition",
            "Decision_family_definition", "Literature_pattern"] + RESP_HEAD
    for j, h in enumerate(head, 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i, gid in enumerate(order, 2):
        it = items[gid]
        vals = [i - 1, gid, f"{it['AI_family']} ({it['AI_family_ZH']})",
                f"{it['Decision_family']} ({it['Decision_family_ZH']})",
                it["AI_family_definition"], it["Decision_family_definition"], it["Literature_pattern"]]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v).alignment = WRAP
        for j in range(len(vals) + 1, len(head) + 1):
            c = ws.cell(i, j)          # response cell: left EMPTY
            c.fill = RESP_FILL
            c.alignment = WRAP
    n = len(order) + 1
    first = len(head) - len(RESP_HEAD) + 1           # column index of Q1

    def col(idx):
        return ws.cell(1, idx).column_letter

    for opts, idx, blank in ((Q1_OPTS, first, False), (Q2_OPTS, first + 1, False),
                             (Q3_OPTS, first + 2, False), (Q4_OPTS, first + 3, True)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid value",
                            error="Choose a value from the drop-down list.")
        dv.add(f"{col(idx)}2:{col(idx)}{n}")
        ws.add_data_validation(dv)
    dv5 = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="600", allow_blank=True,
                         showErrorMessage=True, error="Keep the rationale to 1-3 sentences (max 600 characters).")
    dv5.add(f"{col(first + 4)}2:{col(first + 4)}{n}")
    ws.add_data_validation(dv5)
    ws.freeze_panes = "C2"
    _widths(ws, {"A": 6, "B": 7, "C": 30, "D": 30, "E": 45, "F": 45, "G": 55,
                 col(first): 14, col(first + 1): 18, col(first + 2): 12, col(first + 3): 18, col(first + 4): 60})


def add_attestation(wb, sheet_names):
    ws = wb.create_sheet("ATTESTATION")
    ws.cell(1, 1, "Attestation (one row per coder sheet in this file). All answer cells are empty until the coder fills them.").font = HEAD_FONT
    head = ["Coder_sheet", "Coder_ID (assigned by coordinator)", "Statement", "I_attest (YES/NO)", "Name", "Date",
            "Signature"]
    for j, h in enumerate(head, 1):
        ws.cell(2, j, h).font = HEAD_FONT
    stmt = ("I have not read any project document, manuscript or report that describes how these items are mapped "
            "to categories, and I did not consult the project repository, other coders or an AI assistant. "
            "I coded independently and did not discuss the items with anyone before submitting.")
    for i, s in enumerate(sheet_names, 3):
        ws.cell(i, 1, s)
        ws.cell(i, 3, stmt).alignment = WRAP
        for j in range(2, 8):
            if j != 3:
                ws.cell(i, j).fill = RESP_FILL
    dv = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
    dv.add(f"D3:D{2 + len(sheet_names)}")
    ws.add_data_validation(dv)
    _widths(ws, {"A": 20, "B": 22, "C": 80, "D": 14, "E": 20, "F": 12, "G": 20})


def add_adjudication(wb):
    """Coordinator/adjudicator only (master file). Empty; filled after coder sheets are hash-frozen."""
    ws = wb.create_sheet("ADJUDICATION_EMPTY")
    head = ["G_ID", "Coder_Q1_values (paste after freeze)", "Coder_Q2_values (paste after freeze)",
            "Disagreement_type", "Route (CODER_CONSENSUS / INDEPENDENT_ADJUDICATOR / AUTHOR_ADJUDICATED)",
            "Final_Q1", "Final_Q2", "Final_Q4", "Adjudicator_role", "Evidence_used", "Rationale"]
    for j, h in enumerate(head, 1):
        ws.cell(1, j, h).font = HEAD_FONT
    for i in range(1, 17):
        ws.cell(i + 1, 1, f"G{i:02d}")
    _widths(ws, {c: 22 for c in "ABCDEFGHIJK"})


def build(repo, k, mask_gid):
    items = load_items(repo)
    if mask_gid:
        for gid, it in items.items():
            it["G_ID"] = "ITEM_" + gid[1:]
    ords = orders(k)
    out = os.path.join(repo, OUT_REL)
    canon = [f"G{i:02d}" for i in range(1, 17)]
    # master
    wb = Workbook()
    wb.remove(wb.active)
    add_readme(wb)
    add_categories(wb)
    add_item_info(wb, items, canon)
    for name, o in ords.items():
        add_coder_sheet(wb, name, items, o)
    add_attestation(wb, list(ords))
    add_adjudication(wb)
    paths = [os.path.join(out, PREFIX + ".xlsx")]
    wb.save(paths[0])
    # per coder
    for name, o in ords.items():
        wb = Workbook()
        wb.remove(wb.active)
        add_readme(wb)
        add_categories(wb)
        add_item_info(wb, items, o)
        add_coder_sheet(wb, name, items, o)
        add_attestation(wb, [name])
        p = os.path.join(out, f"{PREFIX}_{name}.xlsx")
        wb.save(p)
        paths.append(p)
    log = dict(base_seed=BASE_SEED, rule="random.Random(BASE_SEED + k).shuffle(['G01'..'G16']) for coder k (Python stdlib)",
               orders=ords, item_sources=["54A_V1_BLIND_PACKET_A_v2_2.csv (family definitions, neutral pair text)",
                                          "11_Priority_Gaps_Frozen_Export_G01_G16_v1_0_20261009.csv (G ID <-> family pair)"],
               mask_gid=mask_gid, outputs=[os.path.relpath(p, repo).replace("\\", "/") for p in paths])
    tdir = os.path.join(out, "tools", "I")
    os.makedirs(tdir, exist_ok=True)
    with open(os.path.join(tdir, "item_order_log.json"), "w", encoding="utf-8") as f:
        json.dump(log, f, indent=1, ensure_ascii=False)
    print(json.dumps(log, indent=1, ensure_ascii=False))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True)
    ap.add_argument("--k", type=int, default=3)
    ap.add_argument("--mask-gid", action="store_true")
    a = ap.parse_args()
    if a.k < 2:
        raise SystemExit("design requires at least 2 coders")
    build(a.repo, a.k, a.mask_gid)
