"""S1b -> S1c DIFF GATE (cell-by-cell, all sheets, validations, sheet/row/column properties, docProps, raw zip parts).

Allowed differences ONLY:
  M1_M4_DEFINITION_OR_DISTINCTION   类別定義 sheet, rows M1..M4, columns B/C
  EVIDENCE_ONLY_DEFINITION          类別定義 sheet, row 僅屬證據層, column B
  SHARED_WORD_REMINDER_ADDED (FLAG) 类別定義 sheet, row 判斷對象, column B
  ATTESTATION_HEADER                作答聲明 sheet, A1
  ATTESTATION_ITEM_1                作答聲明 sheet, column B, rows holding S1b statement 1
  METADATA                          docProps created/modified timestamps; raw parts docProps/core.xml and the sheet
                                    XML of 類別定義 / 作答聲明 (which carry the cell changes above)
Anything else -> verdict STOP (exit 1). Writes OUT/S1C_PATCH_DIFF.md and OUT/tools/s1b_to_s1c_diff_gate_result.json.

Usage: python -E -X utf8 -P s1b_to_s1c_diff_gate.py --repo REPO
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import zipfile

from openpyxl import load_workbook

sys.dont_write_bytecode = True
S1B = os.path.join("claude", "r14_merge_v1_8_2b_v1_s1b_zh_tw")
S1C = os.path.join("claude", "r14_merge_v1_8_2c_v1_s1c_zh_tw")
PAIRS = [(f"coder_distribution/CODER_{k}_ZH_TW.xlsx", f"coder_distribution/CODER_{k}_ZH_TW.xlsx") for k in (1, 2, 3)] + \
        [("20_MASTER_PACKET_S1B_ZH_TW.xlsx", "20_MASTER_PACKET_S1C_ZH_TW.xlsx")]
CATS, ATT = "類別定義", "作答聲明"
M_LABELS = ("M1｜權責與問責", "M2｜流程路由與異常復原", "M3｜人機配置與回應", "M4｜評估與保證")
S1B_ATTEST_1 = "我對本研究的了解全部來自本問卷，作答前與作答期間都未查閱與本研究相關的線上資料。"


def style(c):
    return (c.fill.fgColor.rgb, c.fill.fill_type, c.font.b, c.font.i, c.font.sz, c.font.name,
            c.font.color.rgb if c.font.color else None, c.alignment.wrap_text, c.alignment.vertical,
            c.alignment.horizontal, c.number_format, c.border.left.style, c.protection.locked)


def dvs(ws):
    return sorted((str(d.sqref), d.type, d.formula1, d.formula2, d.operator, d.errorTitle, d.error, d.promptTitle,
                   d.prompt, d.showErrorMessage, d.showInputMessage, d.allow_blank) for d in ws.data_validations.dataValidation)


def classify(sheet, row_key, col, addr, old_v):
    if sheet == CATS and row_key in M_LABELS and col in ("B", "C"):
        return "M1_M4_DEFINITION_OR_DISTINCTION"
    if sheet == CATS and row_key == "僅屬證據層" and col == "B":
        return "EVIDENCE_ONLY_DEFINITION"
    if sheet == CATS and row_key == "判斷對象" and col == "B":
        return "SHARED_WORD_REMINDER_ADDED (FLAG)"
    if sheet == ATT and addr == "A1":
        return "ATTESTATION_HEADER"
    if sheet == ATT and col == "B" and old_v == S1B_ATTEST_1:
        return "ATTESTATION_ITEM_1"
    return None


def compare(pb, pc, D):
    name = os.path.basename(pc)
    wb, wc = load_workbook(pb), load_workbook(pc)
    if wb.sheetnames != wc.sheetnames:
        D.append(dict(file=name, sheet="*", cell="sheetnames", old=wb.sheetnames, new=wc.sheetnames, category=None))
        return
    for sn in wb.sheetnames:
        a, b = wb[sn], wc[sn]
        if (a.max_row, a.max_column) != (b.max_row, b.max_column):
            D.append(dict(file=name, sheet=sn, cell="dimensions", old=[a.max_row, a.max_column],
                          new=[b.max_row, b.max_column], category=None))
        for r in range(1, max(a.max_row, b.max_row) + 1):
            row_key = a.cell(r, 1).value
            for c in range(1, max(a.max_column, b.max_column) + 1):
                ca, cb = a.cell(r, c), b.cell(r, c)
                if ca.value != cb.value:
                    D.append(dict(file=name, sheet=sn, cell=ca.coordinate, old=ca.value, new=cb.value,
                                  category=classify(sn, row_key, ca.column_letter, ca.coordinate, ca.value)))
                if style(ca) != style(cb):
                    D.append(dict(file=name, sheet=sn, cell=ca.coordinate + " (style)", old=str(style(ca)),
                                  new=str(style(cb)), category=None))
            if (a.row_dimensions[r].height, a.row_dimensions[r].hidden) != (b.row_dimensions[r].height,
                                                                            b.row_dimensions[r].hidden):
                D.append(dict(file=name, sheet=sn, cell=f"row {r} (height/hidden)", old=None, new=None, category=None))
        cw = lambda ws: {k: (v.width, v.hidden) for k, v in ws.column_dimensions.items()}  # noqa: E731
        for label, fa, fb in (("column widths", cw(a), cw(b)), ("freeze panes", a.freeze_panes, b.freeze_panes),
                              ("data validations", dvs(a), dvs(b)), ("merged cells", str(a.merged_cells),
                                                                     str(b.merged_cells)),
                              ("sheet state", a.sheet_state, b.sheet_state),
                              ("conditional formatting", len(list(a.conditional_formatting)),
                               len(list(b.conditional_formatting)))):
            if fa != fb:
                D.append(dict(file=name, sheet=sn, cell=label, old=str(fa), new=str(fb), category=None))
    pa, pc_ = wb.properties, wc.properties
    for k in ("creator", "lastModifiedBy", "title", "subject", "keywords", "description", "category", "language",
              "identifier", "revision", "version", "contentStatus", "lastPrinted", "created", "modified"):
        va, vb = getattr(pa, k), getattr(pc_, k)
        if va != vb:
            D.append(dict(file=name, sheet="docProps", cell=k, old=str(va), new=str(vb),
                          category="METADATA" if k in ("created", "modified") else None))
    # raw zip parts
    with zipfile.ZipFile(pb) as za, zipfile.ZipFile(pc) as zb:
        if sorted(za.namelist()) != sorted(zb.namelist()):
            D.append(dict(file=name, sheet="zip", cell="part list", old=sorted(za.namelist()),
                          new=sorted(zb.namelist()), category=None))
        sheet_xml = {f"xl/worksheets/sheet{i + 1}.xml": s for i, s in enumerate(wb.sheetnames)}
        for p in za.namelist():
            if p in zb.namelist() and za.read(p) != zb.read(p):
                ok = p == "docProps/core.xml" or sheet_xml.get(p) in (CATS, ATT)
                D.append(dict(file=name, sheet="zip", cell=p + " (raw bytes)", old="", new="",
                              category="METADATA (raw part carrying the allowed changes)" if ok else None))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True)
    a = ap.parse_args()
    D = []
    for b_rel, c_rel in PAIRS:
        compare(os.path.join(a.repo, S1B, b_rel), os.path.join(a.repo, S1C, c_rel), D)
    bad = [d for d in D if d["category"] is None]
    verdict = "PASS" if not bad else "STOP"
    out = os.path.join(a.repo, S1C)
    with open(os.path.join(out, "tools", "s1b_to_s1c_diff_gate_result.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump(dict(verdict=verdict, n_differences=len(D), n_disallowed=len(bad), differences=D), f, indent=1,
                  ensure_ascii=False, default=str)
    L = ["# S1C_PATCH_DIFF — S1b → S1c 逐格差異（腳本產生）", "",
         "腳本：`tools/s1b_to_s1c_diff_gate.py`；比較 S1b（`claude/r14_merge_v1_8_2b_v1_s1b_zh_tw/`，唯讀）與 S1c 的四個活頁簿："
         "所有工作表的每個儲存格（值與樣式）、列高／隱藏、欄寬、凍結窗格、資料驗證（範圍、類型、選項、錯誤與提示訊息）、"
         "合併儲存格、條件式格式、工作表狀態、文件屬性，以及 zip 內每個部件的原始位元組。", "",
         f"**DIFF GATE：{verdict}** — 差異 {len(D)} 項，其中不允許的 {len(bad)} 項。", "",
         "允許類別：M1_M4_DEFINITION_OR_DISTINCTION、EVIDENCE_ONLY_DEFINITION、SHARED_WORD_REMINDER_ADDED (FLAG)、"
         "ATTESTATION_HEADER、ATTESTATION_ITEM_1、METADATA。", "",
         "| 檔案 | 工作表 | 儲存格 | 類別 | S1b（舊） | S1c（新） |", "|---|---|---|---|---|---|"]
    esc = lambda v: "" if v is None else str(v).replace("|", "\\|").replace("\n", " ")  # noqa: E731
    for d in D:
        L.append(f"| {d['file']} | {d['sheet']} | {d['cell']} | {d['category'] or '**NOT ALLOWED → STOP**'} | "
                 f"{esc(d['old'])} | {esc(d['new'])} |")
    with open(os.path.join(out, "S1C_PATCH_DIFF.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(L) + "\n")
    print(verdict, len(D), len(bad))
    for d in bad[:30]:
        print(d)
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main())
