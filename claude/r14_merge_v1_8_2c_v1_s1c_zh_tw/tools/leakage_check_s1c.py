"""Leakage + blinding + workbook-structure audit for the S1c (ZH_TW) packet. Derived from leakage_check_s1b.py.

S1c additions: exact-sentence whitelist for the two coordinator-mandated attestation texts (they contain the
listed terms 預先設定 / 外部); S1C_REQUIRED_TEXT checks (mandated sentences present verbatim: time, reminder,
attestation header, attestation statement 1, M1-M4 and EVIDENCE_ONLY definitions); M4_RISK_KEYWORD check
(「風險」 must not appear in any M4 text: label, definition, distinction). Everything else as in S1b.
Required result: 0 findings.

Differences from v182: Chinese + English forbidden-term lists (researcher decision items 4 and 7); a Latin-token
whitelist for coder-facing text (no mixed Chinese-English); the sealed item map is NOT opened, so row-uniformity is
checked on ALL item rows (which covers any subset); orders checked against each other, the identity and every S1
order; per-ID content checked against S1C_ITEM_CONTENT_CORRESPONDENCE.csv; AI-pilot shingle overlap dropped
(coder text contains no English prose; the Latin whitelist is stricter). Also writes a shared-term priming report
(record only, never a finding).

Scans
  (1) coder_distribution/CODER_1..3_ZH_TW.xlsx and 20_MASTER_PACKET_S1C_ZH_TW.xlsx
      - raw XML, every part (incl. docProps/core.xml, app.xml, workbook.xml defined names, sheet XML data
        validations): forbidden parts (comments, threaded comments, customXml, VBA, external links, drawings),
        G-ID tokens (cell-reference attributes stripped, so G10..G16 cell refs are not false positives),
        defined names, formulas, hidden sheets/rows/cols, forbidden EN/ZH terms, docProps neutrality;
      - cells (openpyxl): forbidden EN/ZH terms (exact whitelisted sentences excepted), Latin whitelist, G IDs,
        32/47 counts, network numbers, category labels/codes inside item rows, empty response/attestation/
        adjudication cells, uniform per-item formatting, identical literature-pattern text, neutral unique IDs,
        sheet set, data-validation text, no conditional formatting;
      - orders: three coder orders pairwise different, not identity, different from every S1 order.
  (2) every text file in OUT: ITEM-ID with G-ID on one line, >= 5 item IDs on one line (order lists),
      plaintext order-log / map files.
Writes OUT/tools/leakage_check_result_s1c.json.

Usage: python -E -X utf8 -P leakage_check_s1c.py --repo REPO
"""
from __future__ import annotations

import argparse
import csv
import glob
import hashlib
import json
import os
import re
import sys
import zipfile
from itertools import combinations

from openpyxl import load_workbook

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_v1_packet_s1c as BLD  # noqa: E402  (constants only)
import s1b_label_map as LM  # noqa: E402

OUT_REL = BLD.OUT_REL
S1_REL = BLD.S1_REL
GID_RE = re.compile(r"(?<![A-Za-z0-9_])G(0[1-9]|1[0-6])(?![0-9])")
ITEM_RE = re.compile(r"ITEM-A(0[1-9]|1[0-6])")
ATTR_STRIP = re.compile(r'\s(r|ref|sqref|topLeftCell|activeCell|spans|xr:uid|uid)="[^"]*"')
FORBIDDEN_EN = [r"enrich", r"deplet", r"robust", r"q[\s_-]*value", r"\bq\s*=", r"signific", r"sensitiv", r"\bV9\b",
                r"v1[._]8", r"\b1\.8\.\d", r"researcher", r"mapping", r"(?<!co)pilot", r"\bgain\b", r"\bloss\b",
                r"null model", r"github", r"\brepo(sitory)?\b", r"manuscript", r"\bthesis\b", r"\bsearch\b",
                r"\bexternal\b", r"\b53[AB]\b", r"\b54[ABC]\b", r"sealed", r"\bkey\b", r"\bO/E\b", r"p_BH",
                r"Gap_ID", r"G_ID", r"over-?represent", r"under-?represent", r"\bcore\b"]
FORBIDDEN_ZH = ["過度呈現", "不足呈現", "穩健", "顯著", "敏感度", "敏感性", "研究者", "對應表", "映射", "試點", "試行",
                "前導", "先導", "增益", "損失", "版本", "程式", "儲存庫", "資料庫", "原始碼", "公開", "網址", "連結",
                "外部", "上網", "網路搜尋", "手稿", "論文", "虛無", "零模型", "四種", "預先設定", "統計", "檢定",
                "裁決", "答案", "爭議", "核心集合", "q值", "Ｑ值"]
QVAL_ZH = re.compile(r"[qQｑＱ]\s*值")
COUNT_RE = re.compile(r"(?<![0-9])(32|47)(?![0-9])")
# exact coder-facing strings that legitimately contain a listed term (each allowed only as this exact cell text)
WHITELIST_CELLS = {
    BLD.NO_SEARCH_TEXT: {"外部"},
    BLD.ATTEST_ITEMS[1]: {"論文"},
    BLD.ATTEST_ITEMS[2]: {"研究者"},
    BLD.ATTEST_PREAMBLE: {"預先設定"},
    BLD.ATTEST_ITEMS[0]: {"外部"},
}
# coordinator-mandated S1c texts, written here independently of the builder (verbatim from the S1c instruction)
REQ_TIME = "預計作答時間：約 60–120 分鐘，視閱讀速度而異。"
REQ_REMINDER = "功能類型名稱與機制類別有相同或相近字詞時，不得只因字詞相同就選擇該類別。"
REQ_HEADER = ("若有任何一項選擇『不符合』，研究團隊會依預先設定的規則將該份作答標記並分開分析，"
              "不會在未標示的情況下與其他編碼者的主要信度分析合併。")
REQ_ATTEST_1 = "在作答前與作答期間，我沒有另外查閱與本研究內容相關的外部資料；我的判斷只依據本問卷提供的內容。"
REQ_DEFS = {
    "M1｜權責與問責": ("此組合的主要議題是：誰可以核准、拒絕、停止，或改變 AI 輸出被允許作出的決策。包括正式控制與實際控制、問責，以及企業政策所設定的權限邊界。",
                   "若主要議題是 AI 輸出如何進入後續流程，選 M2。若主要議題是人與 AI 如何分工或回應，選 M3。若主要議題是如何確認 AI 輸出是否適合用於該項決策，選 M4。"),
    "M2｜流程路由與異常復原": ("此組合的主要議題是：AI 輸出如何進入下一個決策或行動步驟。包括跨系統或單位的交接、例外處理、回復與復原。",
                      "若主要議題是誰具有正式決策權限，選 M1。若主要議題是人與 AI 的工作分工與互動，選 M3。"),
    "M3｜人機配置與回應": ("此組合的主要議題是：人與 AI 的工作順序與分工，以及使用者如何回應 AI 輸出。例如：誰提出建議、誰做最後決定、如何委派或覆寫，以及使用者是否使用、驗證、修改、升級處理、拒絕或繞過 AI 輸出。",
                    "若主要議題是正式決策權限與問責，選 M1。若主要議題是流程如何轉交、升級或復原，選 M2。若主要議題是 AI 是否經過充分驗證、監控或保證，選 M4。"),
    "M4｜評估與保證": ("此組合的主要議題是：如何確認 AI 輸出適合被用於該項決策。包括：部署前驗證、持續監控、重新驗證、可稽核性、解釋是否足以供決策使用、漂移監測，以及其他能證明 AI 輸出適用於該決策的證據。",
                   "若主要議題是誰有權決定或誰負責，選 M1。若主要議題是人與 AI 如何分工與回應，選 M3。"),
    "僅屬證據層": ("這個配對只呈現研究、技術、資訊或證據上的關係，但沒有形成這個配對所特有的企業決策介面。如果你能想到的問題，是幾乎任何 AI 功能進入這類決策功能時都會同樣遇到的一般性問題，也屬於此類。", None),
}


def s1c_checks(paths, F):
    for p in paths:
        name = os.path.basename(p)
        wb = load_workbook(p)
        readme = [c.value for c in wb[BLD.SHEET_README]["A"]]
        if REQ_TIME not in readme:
            F.append((name, "S1C_REQUIRED_TEXT_MISSING", "time"))
        cats = {r[0]: r for r in wb[BLD.SHEET_CATS].iter_rows(min_row=2, values_only=True)}
        if REQ_REMINDER not in (cats.get("判斷對象", [None, ""])[1] or ""):
            F.append((name, "S1C_REQUIRED_TEXT_MISSING", "shared-word reminder"))
        for lab, (d, dist) in REQ_DEFS.items():
            row = cats.get(lab)
            if row is None or row[1] != d or (dist is not None and row[2] != dist):
                F.append((name, "S1C_REQUIRED_TEXT_MISSING", f"definition/distinction {lab}"))
        m4 = cats.get("M4｜評估與保證", ("", "", ""))
        if any("風險" in (v or "") for v in m4):
            F.append((name, "M4_RISK_KEYWORD", "M4 row"))
        if any("風險" in o for o in LM.options("Q2") + LM.options("Q4") if o.startswith("M4")):
            F.append((name, "M4_RISK_KEYWORD", "M4 drop-down label"))
        wa = wb[BLD.SHEET_ATTEST]
        if REQ_HEADER not in (wa["A1"].value or ""):
            F.append((name, "S1C_REQUIRED_TEXT_MISSING", "attestation header"))
        if "仍會被使用" in (wa["A1"].value or ""):
            F.append((name, "S1C_OLD_HEADER_SENTENCE_PRESENT", "A1"))
        firsts = [wa.cell(r, 2).value for r in range(3, wa.max_row + 1) if wa.cell(r, 2).value in BLD.ATTEST_ITEMS]
        if not firsts or firsts[0] != REQ_ATTEST_1 or len(firsts) % 7 or firsts.count(REQ_ATTEST_1) != len(firsts) // 7:
            F.append((name, "S1C_REQUIRED_TEXT_MISSING", "attestation statement 1"))
LATIN_OK = {"M1", "M2", "M3", "M4", "Q1", "Q2", "Q3", "Q4", "Q5", "AI"}
AI_TOOLS = {"ChatGPT", "Claude", "Gemini", "Copilot"}
LATIN_RE = re.compile(r"[A-Za-z][A-Za-z0-9]*")
CODER_FACING = None  # set in main: sheet names shown to coders
CAT_TOKENS = [l for l in LM.options("Q2")] + ["LONGITUDINAL", "DATA_GATED", "EVIDENCE_ONLY", "NONE", "AMBIGUOUS",
                                              "M1", "M2", "M3", "M4", "M5"]
NET_COLS = ("E_analytic", "Observed", "O_E_Ratio", "q_primary_BH", "q_exact", "q_document_preserving", "q_raw_record")
FUNC = set("與的或和及之在是否為了而且並使能被")  # function characters: bigrams containing them are not content words
BAD_PARTS = ("comments", "threadedcomment", "customxml", "vbaproject", "externallink", "drawing", "printersettings",
             "persons", "metadata")


def sha(p):
    with open(p, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def term_hits(text, allowed=frozenset()):
    hits = [p for p in FORBIDDEN_EN if re.search(p, text, flags=re.I)]
    hits += [z for z in FORBIDDEN_ZH if z in text and z not in allowed]
    if QVAL_ZH.search(text):
        hits.append("q值")
    return hits


def latin_bad(text, attest_tools_ok):
    toks = LATIN_RE.findall(ITEM_RE.sub(" ", text))
    ok = LATIN_OK | (AI_TOOLS if attest_tools_ok else set())
    return [t for t in toks if t not in ok]


# ------------------------------------------------------------------------------------------------ XML level
def xml_scan(path, F, is_coder):
    name = os.path.basename(path)
    with zipfile.ZipFile(path) as z:
        parts = z.namelist()
        for p in parts:
            low = p.lower()
            if any(b in low for b in BAD_PARTS):
                F.append((name, "XML_FORBIDDEN_PART", p))
            try:
                txt = z.read(p).decode("utf-8")
            except UnicodeDecodeError:
                F.append((name, "XML_NON_UTF8_PART", p))
                continue
            body = ATTR_STRIP.sub("", txt)
            for m in GID_RE.finditer(body):
                F.append((name, "XML_GID_TOKEN", f"{p}: ...{body[max(0, m.start() - 30):m.end() + 10]}..."))
            if re.search(r"<definedName[ >]", txt):
                F.append((name, "XML_DEFINED_NAME", p))
            if re.search(r"<f[ >]", txt):
                F.append((name, "XML_FORMULA", p))
            if re.search(r'hidden="(1|true)"|state="(hidden|veryHidden)"', txt):
                F.append((name, "XML_HIDDEN_ELEMENT", p))
            if re.search(r"<conditionalFormatting[ >]", txt):
                F.append((name, "XML_CONDITIONAL_FORMATTING", p))
            if "<extLst" in txt and ("Google" in txt or "x14:dataValidation" in txt):
                F.append((name, "XML_FOREIGN_EXTENSION", p))
            # forbidden terms in every part: text nodes + human-readable attributes (sheet names, validation
            # messages). Tag names / namespaces / package paths (e.g. docProps/core.xml) are not prose.
            attrs = re.findall(r'\s(?:name|error|errorTitle|prompt|promptTitle|title)="([^"]*)"', txt)
            scan = re.sub(r"<[^>]+>", " ", txt) + " " + " ".join(attrs)
            for s in WHITELIST_CELLS:
                scan = scan.replace(s, " ")
            for h in term_hits(scan):
                F.append((name, "XML_FORBIDDEN_TERM", f"{p}: {h}"))
            if p == "docProps/core.xml":
                for tag in ("dc:creator", "cp:lastModifiedBy", "dc:title", "dc:subject", "cp:keywords",
                            "dc:description", "cp:category"):
                    m = re.search(rf"<{tag}>(.*?)</{tag}>", txt)
                    v = m.group(1) if m else ""
                    if is_coder and LATIN_RE.search(v):
                        F.append((name, "DOCPROPS_LATIN", f"{tag}={v}"))
                    if re.search(r"V1|S1|編碼者\d|G\d", v):
                        F.append((name, "DOCPROPS_LEAK", f"{tag}={v}"))
    return len(parts)


# ------------------------------------------------------------------------------------------------ cell level
def row_sig(ws, r):
    cells = ws[r]
    return (tuple((c.fill.fgColor.rgb, c.fill.fill_type, c.font.b, c.font.i, c.font.sz,
                   c.font.color.rgb if c.font.color else None, c.alignment.wrap_text, c.alignment.vertical,
                   c.number_format, c.border.left.style, c.protection.locked) for c in cells),
            ws.row_dimensions[r].height, ws.row_dimensions[r].hidden, sum(c.value is not None for c in cells))


def cell_scan(path, F, net, is_coder, coder_k, content_ref):
    wb = load_workbook(path)
    name = os.path.basename(path)
    facing = set(BLD.SHEET_CODER.values()) | {BLD.SHEET_README, BLD.SHEET_CATS, BLD.SHEET_INFO, BLD.SHEET_ATTEST}
    coder_sheets = [s for s in wb.sheetnames if s.startswith("作答_")]
    if is_coder:
        want = [BLD.SHEET_README, BLD.SHEET_CATS, BLD.SHEET_INFO, BLD.SHEET_CODER[f"CODER_{coder_k}"], BLD.SHEET_ATTEST]
        if wb.sheetnames != want:
            F.append((name, "SHEET_SET", str(wb.sheetnames)))
    else:
        want = [BLD.SHEET_README, BLD.SHEET_CATS, BLD.SHEET_INFO] + [BLD.SHEET_CODER[c] for c in BLD.CODERS] + \
               [BLD.SHEET_ATTEST, "ADJUDICATION_EMPTY", "STUDY_TEAM_CODE_MAP"]
        if wb.sheetnames != want:
            F.append((name, "SHEET_SET", str(wb.sheetnames)))
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            F.append((name, "HIDDEN_SHEET", ws.title))
        if ws.title in facing:
            for t in latin_bad(ws.title, False):
                F.append((name, "LATIN_IN_SHEET_NAME", f"{ws.title}: {t}"))
        if ws.conditional_formatting and len(list(ws.conditional_formatting)):
            F.append((name, "CONDITIONAL_FORMATTING", ws.title))
        for k, d in ws.column_dimensions.items():
            if d.hidden:
                F.append((name, "HIDDEN_COLUMN", f"{ws.title}!{k}"))
        for dv in ws.data_validations.dataValidation:
            if dv.promptTitle or dv.prompt or dv.showInputMessage:
                F.append((name, "DV_INPUT_PROMPT", f"{ws.title}!{dv.sqref}"))
            if not dv.showErrorMessage:
                F.append((name, "DV_NO_ERROR_STOP", f"{ws.title}!{dv.sqref}"))
            for txt in (dv.errorTitle or "", dv.error or "", dv.formula1 or ""):
                for h in term_hits(txt):
                    F.append((name, "DV_FORBIDDEN_TERM", f"{ws.title}!{dv.sqref}: {h}"))
                if ws.title in facing:
                    for t in latin_bad(txt, False):
                        F.append((name, "DV_LATIN_TOKEN", f"{ws.title}!{dv.sqref}: {t}"))
        rows = list(ws.iter_rows(values_only=True))
        is_item_sheet = ws.title.startswith("作答_") or ws.title == BLD.SHEET_INFO
        for ri, row in enumerate(rows):
            for ci, v in enumerate(row):
                if v is None:
                    continue
                s, where = str(v), f"{ws.title}!R{ri + 1}C{ci + 1}"
                allowed = WHITELIST_CELLS.get(s, set())
                for h in term_hits(s, frozenset(allowed)):
                    F.append((name, "FORBIDDEN_TERM", f"{where}: {h}"))
                if GID_RE.search(s):
                    F.append((name, "GID_IN_CELL", where))
                if not isinstance(v, (int, float)) and COUNT_RE.search(s):
                    F.append((name, "COUNT_32_47", where))
                for num in net:
                    if num in s:
                        F.append((name, "NETWORK_NUMBER", f"{where}: {num}"))
                if ws.title in facing:
                    tools_ok = ws.title == BLD.SHEET_ATTEST and s == BLD.ATTEST_ITEMS[5]
                    for t in latin_bad(s, tools_ok):
                        F.append((name, "LATIN_TOKEN", f"{where}: {t}"))
                if is_item_sheet and ri > 0 and ci < len(BLD.CODER_HEAD):
                    for tok in CAT_TOKENS:
                        if tok in s:
                            F.append((name, "CATEGORY_IN_ITEM_ROW", f"{where}: {tok}"))
        # empty response / attestation / adjudication cells
        if ws.title.startswith("作答_"):
            for r in range(2, ws.max_row + 1):
                for c in range(len(BLD.CODER_HEAD) + 1, len(BLD.CODER_HEAD) + 6):
                    if ws.cell(r, c).value is not None:
                        F.append((name, "NON_EMPTY_RESPONSE_CELL", f"{ws.title}!R{r}C{c}"))
        if ws.title == BLD.SHEET_ATTEST:
            for r in range(3, ws.max_row + 1):
                if ws.cell(r, 3).value is not None:
                    F.append((name, "NON_EMPTY_ATTESTATION_CELL", f"R{r}"))
        if ws.title == "ADJUDICATION_EMPTY":
            for r in range(2, ws.max_row + 1):
                for c in range(2, ws.max_column + 1):
                    if ws.cell(r, c).value is not None:
                        F.append((name, "NON_EMPTY_ADJUDICATION_CELL", f"R{r}C{c}"))
        # uniform per-item formatting + identical pattern text + neutral IDs + content per ID
        if is_item_sheet:
            sigs = {row_sig(ws, r) for r in range(2, ws.max_row + 1)}
            if len(sigs) != 1:
                F.append((name, "ITEM_ROWS_NOT_UNIFORM", ws.title))
            pat = {ws.cell(r, len(BLD.CODER_HEAD)).value for r in range(2, ws.max_row + 1)}
            if pat != {BLD.PATTERN_TEXT}:
                F.append((name, "PATTERN_TEXT_NOT_IDENTICAL", ws.title))
            ids = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
            if sorted(ids) != BLD.NEUTRAL or len(set(ids)) != 16:
                F.append((name, "ITEM_IDS_NOT_NEUTRAL_16", ws.title))
            head = [c.value for c in ws[1]]
            for r in range(2, ws.max_row + 1):
                g = {h: ws.cell(r, j + 1).value for j, h in enumerate(head)}
                z = dict(AI_zh=g["AI 功能類型"], AI_def_zh=g["AI 功能類型定義"], Dec_zh=g["決策功能類型"],
                         Dec_def_zh=g["決策功能類型定義"])
                if hashlib.sha256(BLD.zh_text(z).encode("utf-8")).hexdigest() != content_ref.get(g["題目編號"]):
                    F.append((name, "ITEM_CONTENT_MISMATCH", f"{ws.title}: {g['題目編號']}"))
    orders = {ws.title: [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
              for ws in wb.worksheets if ws.title.startswith("作答_")}
    info_order = [wb[BLD.SHEET_INFO].cell(r, 2).value for r in range(2, wb[BLD.SHEET_INFO].max_row + 1)]
    if is_coder and list(orders.values())[0] != info_order:
        F.append((name, "INFO_ORDER_DIFFERS_FROM_CODER_SHEET", ""))
    return orders


def text_scan(out, F):
    n = 0
    for p in glob.glob(os.path.join(out, "**", "*"), recursive=True):
        if os.path.isdir(p) or p.lower().endswith(".xlsx"):
            continue
        base = os.path.basename(p)
        if base in ("ITEM_ID_TO_GID_MAP.json", "ITEM_ORDER_LOG.json", "S1B_ITEM_ORDER_LOG.json") or "SECRET_SEED" in base:
            F.append((base, "PLAINTEXT_SEALED_FILE_IN_OUT", p))
        try:
            txt = open(p, encoding="utf-8").read()
        except (UnicodeDecodeError, OSError):
            continue
        n += 1
        for i, line in enumerate(txt.splitlines(), 1):
            if ITEM_RE.search(line) and GID_RE.search(line):
                F.append((base, "ITEM_GID_PAIR_ON_LINE", f"line {i}"))
            if len(set(ITEM_RE.findall(line))) >= 5:
                F.append((base, "ITEM_ORDER_LIST_ON_LINE", f"line {i}"))
    return n


def priming_report():
    """Record only: CJK bigrams shared between family labels/definitions and category labels/definitions."""
    def grams(t):
        t = re.sub(r"[^一-鿿]", " ", t)
        return {w[i:i + 2] for w in t.split() for i in range(len(w) - 1) if not set(w[i:i + 2]) & FUNC}
    stop = {"決策", "功能", "策功", "這項"}
    fam_lab = {v[1]: k for k, v in list(BLD.AI_FAM.items()) + list(BLD.DEC_FAM.items())}
    fam_def = {v[1]: v[3] for v in list(BLD.AI_FAM.values()) + list(BLD.DEC_FAM.values())}
    cat_lab = {lab: re.sub(r"^M\d｜", "", lab) for lab in LM.options("Q2")}
    cat_def = {row[0]: row[1] for row in BLD.CATEGORY_ROWS[:9]}
    label_label, label_def = [], []
    for f in fam_lab:
        for c, cl in cat_lab.items():
            sh = sorted((grams(f) & grams(cl)) - stop)
            if sh:
                label_label.append(dict(family=f, category=c, shared=sh))
            sh2 = sorted((grams(f) & grams(cat_def[c])) - stop)
            sh3 = sorted((grams(fam_def[f]) & grams(cl)) - stop)
            if sh2 or sh3:
                label_def.append(dict(family=f, category=c, family_label_in_category_definition=sh2,
                                      category_label_in_family_definition=sh3))
    words = sorted({w for x in label_label for w in x["shared"]})
    rule = BLD.CATEGORY_ROWS[[r[0] for r in BLD.CATEGORY_ROWS].index("判斷對象")][1]
    return dict(note="RECORD ONLY (never a finding). CJK bigrams shared between Chinese family labels and Chinese "
                     "category labels / definitions as shown to coders.",
                label_to_label=label_label, label_to_definition=label_def, shared_label_words=words,
                shared_label_words_named_in_coder_rule=[w for w in words if w in rule],
                shared_label_words_not_named_in_coder_rule=[w for w in words if w not in rule])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True)
    a = ap.parse_args()
    out = os.path.join(a.repo, OUT_REL)
    net = set()
    for r in csv.DictReader(open(os.path.join(a.repo, "11_Priority_Gaps_Frozen_Export_G01_G16_v1_0_20261009.csv"),
                                 encoding="utf-8-sig")):
        for c in NET_COLS:
            v = r[c]
            if re.fullmatch(r"[0-9.]+", v) and len(v) >= 4:
                net.add(v)
                net.add(f"{float(v):.2f}")
    net = {x for x in net if len(x) >= 4 and x != "0.00"}
    content_ref = {}
    for r in csv.DictReader(open(os.path.join(out, "S1C_ITEM_CONTENT_CORRESPONDENCE.csv"), encoding="utf-8")):
        content_ref[r["Item_ID"]] = r["S1C_ZH_text_sha256"]
    F = []
    dist = sorted(glob.glob(os.path.join(out, "coder_distribution", "*")))
    s1c_checks(dist + [os.path.join(out, "20_MASTER_PACKET_S1C_ZH_TW.xlsx")], F)
    dist_names = [os.path.basename(p) for p in dist]
    if dist_names != [f"CODER_{k}_ZH_TW.xlsx" for k in (1, 2, 3)]:
        F.append(("coder_distribution", "UNEXPECTED_FILES", str(dist_names)))
    per, coder_orders = {}, {}
    for p in dist + [os.path.join(out, "20_MASTER_PACKET_S1C_ZH_TW.xlsx")]:
        is_coder = "coder_distribution" in p
        k = os.path.basename(p)[6] if is_coder else None
        parts = xml_scan(p, F, is_coder)
        orders = cell_scan(p, F, net, is_coder, k, content_ref)
        if is_coder:
            coder_orders[f"CODER_{k}"] = list(orders.values())[0]
        else:
            master_orders = {f"CODER_{t[-1]}": o for t, o in orders.items()}
        per[os.path.relpath(p, out).replace("\\", "/")] = dict(sha256=sha(p), bytes=os.path.getsize(p),
                                                              xml_parts_scanned=parts, distributed_to_coders=is_coder)
    # orders
    s1_orders = {}
    for c in BLD.CODERS:
        wb = load_workbook(os.path.join(a.repo, S1_REL, "coder_distribution", f"{c}.xlsx"), read_only=True)
        rows = list(wb[c].iter_rows(values_only=True))
        s1_orders[c] = [r[1] for r in rows[1:]]
    tup = {c: tuple(o) for c, o in coder_orders.items()}
    if len(set(tup.values())) != 3:
        F.append(("orders", "CODER_ORDERS_NOT_DISTINCT", ""))
    if tuple(BLD.NEUTRAL) in tup.values():
        F.append(("orders", "ORDER_EQUALS_IDENTITY", ""))
    for c, o in tup.items():
        if o in {tuple(x) for x in s1_orders.values()}:
            F.append(("orders", "ORDER_EQUALS_AN_S1_ORDER", c))
        if master_orders.get(c) != list(o):
            F.append(("orders", "MASTER_ORDER_DIFFERS_FROM_CODER_FILE", c))
    n_text = text_scan(out, F)
    pair_pos_same = {f"{x}-{y}": sum(i == j for i, j in zip(tup[x], tup[y])) for x, y in combinations(sorted(tup), 2)}
    vs_s1_pos_same = {c: sum(i == j for i, j in zip(tup[c], s1_orders[c])) for c in tup}
    res = dict(checker="leakage_check_s1c.py", files=per, text_files_scanned=n_text,
               network_numbers_checked=len(net), forbidden_terms_en=FORBIDDEN_EN, forbidden_terms_zh=FORBIDDEN_ZH,
               latin_whitelist=sorted(LATIN_OK) + ["ITEM-Axx IDs", "AI tool names in attestation statement 6 only"],
               whitelisted_exact_sentences={k: sorted(v) for k, v in WHITELIST_CELLS.items()},
               orders=dict(pairwise_distinct=len(set(tup.values())) == 3,
                           differ_from_every_S1_order=all(o not in {tuple(x) for x in s1_orders.values()}
                                                          for o in tup.values()),
                           same_position_counts_between_coders=pair_pos_same,
                           same_position_counts_vs_same_coder_S1=vs_s1_pos_same),
               coder_distribution_files=dist_names,
               priming_report=priming_report(),
               findings=[dict(file=f, check=c, where=w) for f, c, w in F])
    res["n_findings"] = len(res["findings"])
    res["verdict"] = "PASS" if not F else "FAIL"
    with open(os.path.join(out, "tools", "leakage_check_result_s1c.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump(res, f, indent=1, ensure_ascii=False)
    print(res["verdict"], res["n_findings"])
    for x in res["findings"][:60]:
        print(x)
    return 0 if not F else 1


if __name__ == "__main__":
    sys.exit(main())
