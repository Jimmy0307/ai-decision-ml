"""Write S1D_RESEARCHER_VISUAL_REVIEW.md: everything a coder sees in CODER_1_ZH_TW.xlsx, in the order a coder sees
it (sheet tabs 1..5), with cell addresses, so the researcher can review without opening Excel.

Per sheet: every non-empty cell (address + text), hidden rows/columns (if any), the empty response ranges, every
data-validation range with all drop-down options and its error/prompt messages. Plus document properties. Read-only.

Usage: python -E -X utf8 -P visual_review_s1d.py [--file X.xlsx] [--out FILE.md]
"""
from __future__ import annotations

import argparse
import hashlib
import os
import sys

from openpyxl import load_workbook

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.dirname(HERE)
EXPECTED_ORDER = ["填答說明", "類別定義", "題目資訊", "作答_編碼者", "作答聲明"]


def esc(v):
    return str(v).replace("|", "\\|").replace("\n", "<br>")


def review(path):
    wb = load_workbook(path)
    h = hashlib.sha256(open(path, "rb").read()).hexdigest()
    p = wb.properties
    L = ["# S1D 研究者目視審閱稿（CODER_1_ZH_TW.xlsx）", "",
         "依編碼者看到的順序（工作表 1 填答說明 → 2 類別定義 → 3 題目資訊 → 4 作答頁 → 5 作答聲明），逐格列出每個非空白儲存格"
         "（含位址）、每個資料驗證範圍的全部下拉選項與錯誤／提示訊息，以及文件屬性。三份編碼者檔案除作答頁名稱、"
         "題目順序與聲明表的編碼者編號外內容相同（見 `S1D_BLINDING_AUDIT.md`）。本檔只供研究團隊使用。", "",
         f"- 檔案：`coder_distribution/{os.path.basename(path)}`，SHA-256 `{h}`",
         "- 工作表順序：" + "　".join(f"{i}. {s}" for i, s in enumerate(wb.sheetnames, 1)),
         "- 文件屬性：" + "；".join(f"{k}＝{getattr(p, k) if getattr(p, k) not in (None, '') else '（空白）'}"
                               for k in ("title", "creator", "lastModifiedBy", "subject", "keywords", "description",
                                         "category")) + "；建立／修改時間為建置時間戳記", ""]
    for i, ws in enumerate(wb.worksheets, 1):
        L += [f"## {i}. 工作表「{ws.title}」", ""]
        hid_r = [r for r, d in ws.row_dimensions.items() if d.hidden]
        hid_c = [k for k, d in ws.column_dimensions.items() if d.hidden]
        L.append(f"狀態：{ws.sheet_state}；凍結窗格：{ws.freeze_panes or '無'}；隱藏列：{hid_r or '無'}；隱藏欄：{hid_c or '無'}")
        L += ["", "| 位址 | 內容 |", "|---|---|"]
        empty_fill = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    L.append(f"| {c.coordinate} | {esc(c.value)} |")
                elif c.fill.fill_type == "solid":
                    empty_fill.append(c.coordinate)
        L.append("")
        if empty_fill:
            L += [f"空白作答格（淡黃色，{len(empty_fill)} 格）：{empty_fill[0]} … {empty_fill[-1]}", ""]
        dvs = ws.data_validations.dataValidation
        if dvs:
            L += ["**資料驗證**", "", "| 範圍 | 類型 | 下拉選項／條件 | 錯誤標題 | 錯誤訊息 | 輸入提示 |", "|---|---|---|---|---|---|"]
            for d in dvs:
                cond = "／".join(d.formula1.strip('"').split(",")) if d.type == "list" else f"{d.operator} {d.formula1}"
                prompt = f"{d.promptTitle or ''} {d.prompt or ''}".strip() or "（無）"
                L.append(f"| {d.sqref} | {d.type} | {esc(cond)} | {esc(d.errorTitle or '')} | {esc(d.error or '')} | "
                         f"{esc(prompt)} |")
            L.append("")
    return "\n".join(L) + "\n"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", default=os.path.join(OUT, "coder_distribution", "CODER_1_ZH_TW.xlsx"))
    ap.add_argument("--out", default=os.path.join(OUT, "S1D_RESEARCHER_VISUAL_REVIEW.md"))
    a = ap.parse_args()
    wb = load_workbook(a.file, read_only=True)
    names = wb.sheetnames
    if len(names) != 5 or any(not n.startswith(e) for n, e in zip(names, EXPECTED_ORDER)):
        raise SystemExit(f"unexpected sheet order {names}")
    with open(a.out, "w", encoding="utf-8", newline="\n") as f:
        f.write(review(a.file))
    print("written", a.out)


if __name__ == "__main__":
    sys.exit(main())
