"""S1c -> S1d DIFF GATE. Cell-by-cell (values + styles) for every sheet, data validations, sheet/row/column properties,
docProps and the raw bytes of every zip part. The 填答說明 sheet is compared by CONTENT BLOCKS (difflib alignment of
the row texts), because new rows are inserted there; every aligned unchanged row must be identical in text and
style, and every inserted / replaced row must be one of the allowed S1d texts.

Allowed categories (anything else -> STOP, exit 1):
  VIEWPOINT, INTERFACE_CLARIFICATION, DECISION_SCOPE_CLARIFICATION, PAIRING_SPECIFIC_CLARIFICATION,
  Q1_OPTION_EXPLANATIONS, NON_APPLICABILITY_WARNING, Q2_WORDING (README row + coder-sheet header),
  Q1_UNCERTAIN_VS_Q2_AMBIGUOUS, QUICK_DISTINCTION, DEPLETION_FINAL_SENTENCE (pattern cells: new == old + sentence),
  ATTESTATION_SOCIAL_DESIRABILITY (header: new == old + sentence), METADATA (timestamps; raw parts carrying the above).
Writes OUT/S1D_PATCH_DIFF.md and OUT/tools/s1c_to_s1d_diff_gate_result.json.

Usage: python -E -X utf8 -P s1c_to_s1d_diff_gate.py --repo REPO
"""
from __future__ import annotations

import argparse
import difflib
import json
import os
import sys
import zipfile

from openpyxl import load_workbook

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_v1_packet_s1d as BLD  # noqa: E402  (constants only)

S1C = os.path.join("claude", "r14_merge_v1_8_2c_v1_s1c_zh_tw")
S1D = BLD.OUT_REL
PAIRS = [(f"coder_distribution/CODER_{k}_ZH_TW.xlsx",) * 2 for k in (1, 2, 3)] + \
        [("20_MASTER_PACKET_S1C_ZH_TW.xlsx", "20_MASTER_PACKET_S1D_ZH_TW.xlsx")]
README, INFO, ATT = BLD.SHEET_README, BLD.SHEET_INFO, BLD.SHEET_ATTEST
NEW_README = {BLD.S1D_VIEWPOINT: "VIEWPOINT", BLD.S1D_INTERFACE: "INTERFACE_CLARIFICATION",
              BLD.S1D_SCOPE: "DECISION_SCOPE_CLARIFICATION", BLD.S1D_SPECIFIC: "PAIRING_SPECIFIC_CLARIFICATION",
              BLD.S1D_Q1_YES: "Q1_OPTION_EXPLANATIONS", BLD.S1D_Q1_NO: "Q1_OPTION_EXPLANATIONS",
              BLD.S1D_Q1_UNC: "Q1_OPTION_EXPLANATIONS", BLD.S1D_WARNING: "NON_APPLICABILITY_WARNING",
              BLD.S1D_Q2_README: "Q2_WORDING", BLD.S1D_Q1_VS_Q2: "Q1_UNCERTAIN_VS_Q2_AMBIGUOUS",
              BLD.S1D_QUICK: "QUICK_DISTINCTION"}
OLD_README = {"　　是：你能具體說出一項企業必須做出的決策，而且這項決策是這個組合特有的。": "Q1_OPTION_EXPLANATIONS",
              "　　否：這段內容沒有涉及任何企業決策；或你能說出的決策，幾乎任何 AI 功能進入這個決策功能時都會同樣出現。": "Q1_OPTION_EXPLANATIONS",
              "　　無法確定：你無法在「是」與「否」之間做出判斷。": "Q1_OPTION_EXPLANATIONS",
              "Q2 如果有，這段內容最主要對應哪一種企業決策機制？（單選）": "Q2_WORDING"}
OLD_Q2_HEAD = "Q2 如果有，這段內容最主要對應哪一種企業決策機制？（單選）"


def style(c):
    return (c.fill.fgColor.rgb, c.fill.fill_type, c.font.b, c.font.i, c.font.sz, c.font.name,
            c.font.color.rgb if c.font.color else None, c.alignment.wrap_text, c.alignment.vertical,
            c.alignment.horizontal, c.number_format, c.border.left.style, c.protection.locked)


def dvs(ws):
    return sorted((str(d.sqref), d.type, d.formula1, d.formula2, d.operator, d.errorTitle, d.error, d.promptTitle,
                   d.prompt, d.showErrorMessage, d.showInputMessage, d.allow_blank)
                  for d in ws.data_validations.dataValidation)


def classify_cell(sheet, addr, col, row, old, new):
    if (sheet == INFO or sheet.startswith("作答_")) and col == "G" and row >= 2 and isinstance(old, str) \
            and new == old + BLD.S1D_PATTERN_APPEND:
        return "DEPLETION_FINAL_SENTENCE"
    if sheet.startswith("作答_") and addr == "I1" and old == OLD_Q2_HEAD and new == BLD.Q_TEXT["Q2"]:
        return "Q2_WORDING"
    if sheet == ATT and addr == "A1" and isinstance(old, str) and new == old + BLD.S1D_ATTEST_APPEND:
        return "ATTESTATION_SOCIAL_DESIRABILITY"
    return None


def readme_blocks(name, a, b, D):
    if a.max_column != 1 or b.max_column != 1:
        D.append(dict(file=name, sheet=README, cell="columns", old=a.max_column, new=b.max_column, category=None))
    ra = [(c.value or "", c) for c in a["A"]]
    rb = [(c.value or "", c) for c in b["A"]]
    sm = difflib.SequenceMatcher(None, [t for t, _ in ra], [t for t, _ in rb], autojunk=False)
    ref_style = style(rb[2][1])                                   # an ordinary (non-title) instruction row
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            for (ta, ca), (tb, cb) in zip(ra[i1:i2], rb[j1:j2]):
                if style(ca) != style(cb):
                    D.append(dict(file=name, sheet=README, cell=f"{ca.coordinate}->{cb.coordinate} (style)",
                                  old=str(style(ca)), new=str(style(cb)), category=None))
            continue
        for ta, ca in ra[i1:i2]:
            D.append(dict(file=name, sheet=README, cell=f"{ca.coordinate} (S1c, {op})", old=ta, new="",
                          category=OLD_README.get(ta) if op == "replace" else None))
        for tb, cb in rb[j1:j2]:
            cat = NEW_README.get(tb)
            if cat and style(cb) != ref_style:
                cat = None
            D.append(dict(file=name, sheet=README, cell=f"{cb.coordinate} (S1d, {op})", old="", new=tb, category=cat))
    hs = lambda ws: {(d.height, d.hidden) for d in ws.row_dimensions.values()} - {(None, False)}  # noqa: E731
    if hs(a) or hs(b):
        D.append(dict(file=name, sheet=README, cell="row heights/hidden", old=str(hs(a)), new=str(hs(b)),
                      category=None))


def compare(pa, pb, D):
    name = os.path.basename(pb)
    wa, wb = load_workbook(pa), load_workbook(pb)
    if wa.sheetnames != wb.sheetnames:
        D.append(dict(file=name, sheet="*", cell="sheetnames", old=wa.sheetnames, new=wb.sheetnames, category=None))
        return
    for sn in wa.sheetnames:
        a, b = wa[sn], wb[sn]
        if sn == README:
            readme_blocks(name, a, b, D)
        else:
            if (a.max_row, a.max_column) != (b.max_row, b.max_column):
                D.append(dict(file=name, sheet=sn, cell="dimensions", old=[a.max_row, a.max_column],
                              new=[b.max_row, b.max_column], category=None))
            for r in range(1, max(a.max_row, b.max_row) + 1):
                for c in range(1, max(a.max_column, b.max_column) + 1):
                    ca, cb = a.cell(r, c), b.cell(r, c)
                    if ca.value != cb.value:
                        D.append(dict(file=name, sheet=sn, cell=ca.coordinate, old=ca.value, new=cb.value,
                                      category=classify_cell(sn, ca.coordinate, ca.column_letter, r, ca.value,
                                                             cb.value)))
                    if style(ca) != style(cb):
                        D.append(dict(file=name, sheet=sn, cell=ca.coordinate + " (style)", old=str(style(ca)),
                                      new=str(style(cb)), category=None))
                if (a.row_dimensions[r].height, a.row_dimensions[r].hidden) != \
                        (b.row_dimensions[r].height, b.row_dimensions[r].hidden):
                    D.append(dict(file=name, sheet=sn, cell=f"row {r} (height/hidden)", old=None, new=None,
                                  category=None))
        cw = lambda ws: {k: (v.width, v.hidden) for k, v in ws.column_dimensions.items()}  # noqa: E731
        for label, fa, fb in (("column widths", cw(a), cw(b)), ("freeze panes", a.freeze_panes, b.freeze_panes),
                              ("data validations", dvs(a), dvs(b)),
                              ("merged cells", str(a.merged_cells), str(b.merged_cells)),
                              ("sheet state", a.sheet_state, b.sheet_state),
                              ("conditional formatting", len(list(a.conditional_formatting)),
                               len(list(b.conditional_formatting)))):
            if fa != fb:
                D.append(dict(file=name, sheet=sn, cell=label, old=str(fa), new=str(fb), category=None))
    for k in ("creator", "lastModifiedBy", "title", "subject", "keywords", "description", "category", "language",
              "identifier", "revision", "version", "contentStatus", "lastPrinted", "created", "modified"):
        va, vb = getattr(wa.properties, k), getattr(wb.properties, k)
        if va != vb:
            D.append(dict(file=name, sheet="docProps", cell=k, old=str(va), new=str(vb),
                          category="METADATA" if k in ("created", "modified") else None))
    carrying = {README, INFO, ATT} | {s for s in wa.sheetnames if s.startswith("作答_")}
    with zipfile.ZipFile(pa) as za, zipfile.ZipFile(pb) as zb:
        if sorted(za.namelist()) != sorted(zb.namelist()):
            D.append(dict(file=name, sheet="zip", cell="part list", old=sorted(za.namelist()),
                          new=sorted(zb.namelist()), category=None))
        sheet_xml = {f"xl/worksheets/sheet{i + 1}.xml": s for i, s in enumerate(wa.sheetnames)}
        for p in za.namelist():
            if p in zb.namelist() and za.read(p) != zb.read(p):
                ok = p == "docProps/core.xml" or sheet_xml.get(p) in carrying
                D.append(dict(file=name, sheet="zip", cell=p + " (raw bytes)", old="", new="",
                              category="METADATA (raw part carrying the allowed changes)" if ok else None))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True)
    a = ap.parse_args()
    D = []
    for c_rel, d_rel in PAIRS:
        compare(os.path.join(a.repo, S1C, c_rel), os.path.join(a.repo, S1D, d_rel), D)
    bad = [d for d in D if d["category"] is None]
    verdict = "PASS" if not bad else "STOP"
    out = os.path.join(a.repo, S1D)
    with open(os.path.join(out, "tools", "s1c_to_s1d_diff_gate_result.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump(dict(verdict=verdict, n_differences=len(D), n_disallowed=len(bad), differences=D), f, indent=1,
                  ensure_ascii=False, default=str)
    esc = lambda v: "" if v is None else str(v).replace("|", "\\|").replace("\n", " ")  # noqa: E731
    L = ["# S1D_PATCH_DIFF — S1c → S1d 逐格差異（腳本產生）", "",
         "腳本：`tools/s1c_to_s1d_diff_gate.py`；比較 S1c（`claude/r14_merge_v1_8_2c_v1_s1c_zh_tw/`，唯讀）與 S1d 的四個活頁簿："
         "每個工作表的每個儲存格（值與樣式）、列高／隱藏、欄寬、凍結窗格、資料驗證、合併儲存格、條件式格式、工作表狀態、"
         "文件屬性，以及 zip 內每個部件的原始位元組。「填答說明」因插入新列而以內容區塊比對（difflib 對齊）："
         "對齊後未變的列文字與樣式必須完全相同；插入或取代的列必須是允許的 S1d 文字，且樣式與一般說明列相同。", "",
         f"**DIFF GATE：{verdict}** — 差異 {len(D)} 項，其中不允許的 {len(bad)} 項。", "",
         "| 檔案 | 工作表 | 儲存格 | 類別 | S1c（舊） | S1d（新） |", "|---|---|---|---|---|---|"]
    for d in D:
        L.append(f"| {d['file']} | {d['sheet']} | {d['cell']} | {d['category'] or '**NOT ALLOWED → STOP**'} | "
                 f"{esc(d['old'])} | {esc(d['new'])} |")
    with open(os.path.join(out, "S1D_PATCH_DIFF.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(L) + "\n")
    print(verdict, len(D), len(bad))
    for d in bad[:30]:
        print(d)
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main())
