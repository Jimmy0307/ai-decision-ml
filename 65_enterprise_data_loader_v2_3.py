"""65 — Enterprise empirical data loader and runner v2.3.

Turns the filled 62 data template (one folder of CSVs, or the 62 .xlsx workbook) into
  (1) a validation report (provenance, unit, interval and zero-fill rules of 43 / 62),
  (2) the parameter-identification coverage table (44 Table 4.5),
  (3) a 51 instance built ONLY from point-identified cells (every other cell stays UNIDENTIFIED = None),
  (4) solver runs through the 51 state machine (Models B, C, A+, C0, A per policy) and the derived quantities
      for Chapter 4 (architecture metrics, F_c thresholds, governance decomposition, price of requirement,
      design-response analysis, interface activation),
  (5) for interval-identified cells: a scenario-registry DRAFT + commitment, and a robust / VOI run that is allowed
      only after the registry commitment was written (pre-registration; 42 v2.2 section 13.1).

It does not change the model (42 v2.2) or the solver (51). It never fills a missing cell with zero.

Commands
  python 65_enterprise_data_loader_v2_3.py --make-template  OUTDIR      blank CSV schemas + .xlsx + dictionary
  python 65_enterprise_data_loader_v2_3.py --make-pilot     OUTDIR      63 minimum-pilot skeleton (all UNIDENTIFIED)
  python 65_enterprise_data_loader_v2_3.py --validate       DATA        validation + coverage only
  python 65_enterprise_data_loader_v2_3.py --run            DATA [--out DIR]   point solve + Chapter 4 tables
  python 65_enterprise_data_loader_v2_3.py --make-registry  DATA [--out DIR]   interval registry draft + commitment
  python 65_enterprise_data_loader_v2_3.py --robust         DATA --registry REG.json [--out DIR]   robust + VOI (Model B)
  python 65_enterprise_data_loader_v2_3.py --self-test
DATA is a folder of the 62 CSVs or the 62 .xlsx workbook.
"""
from __future__ import annotations

import copy
import csv
import hashlib
import importlib.util
import itertools
import json
import math
import os
import random
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
VERSION = "65_enterprise_data_loader_v2_3"


def _load51():
    spec = importlib.util.spec_from_file_location("s51", os.path.join(HERE, "51_enterprise_portfolio_solver_v2_2.py"))
    m = importlib.util.module_from_spec(spec)
    sys.modules["s51"] = m
    spec.loader.exec_module(m)
    return m


S = _load51()

# =============================================================================
# 1. Schema (single source of truth for 62)
# =============================================================================
VAL = ["value", "unit", "provenance_type", "source_id", "period", "uncertainty", "lower", "upper",
       "identification_status", "assumption_note"]
PROV = ["OBSERVED", "EXPERT_ELICITED_INTERNAL", "EXPERT_ELICITED_EXTERNAL", "PUBLIC_OBSERVED_BOUND",
        "SCENARIO_SENSITIVITY", "DEFINITIONAL", "UNIDENTIFIED"]
UNC = ["POINT", "INTERVAL", "NONE"]
IDS = ["IDENTIFIED_POINT", "IDENTIFIED_INTERVAL", "BOUND_ONLY", "SCENARIO_ONLY", "UNIDENTIFIED"]
KIND_UNIT = {"prob": "probability", "ord": "ordinal_level", "bin": "binary", "list": "list", "nominal": "nominal"}
CONFIG_MODES = ["HumanOnly", "AI->Human", "Human->AI", "Aggregation", "Delegated", "BoundedExecution"]
RESP = list(S.RESPONSES)
PATH_FIELDS = list(S.PATH_UNITS)

TABLES = {
    "S_sources": dict(block="S", cols=["source_id", "source_scope", "source_kind", "external_role", "description", "date",
                                       "custodian_role", "access_restriction"], values=False,
                      doc="Every source_id used anywhere must be registered here. source_scope=external sources can only "
                          "support EXPERT_ELICITED_EXTERNAL or PUBLIC_OBSERVED_BOUND cells; external roles are never optimizers."),
    "E_enterprise": dict(block="E", index=["parameter", "unit_id", "role"],
                         params={"A_bar_BUD": ("TWD/period", "num"), "A_bar_ENG": ("hours/period", "num"),
                                 "rev_cap": ("hours/period", "num"), "B0_BUD": ("TWD/period", "num"),
                                 "B0_ENG": ("hours/period", "num"), "lam": ("TWD/hour", "num"),
                                 "lam_u": ("TWD/hour", "num"), "lam_U": ("TWD/hour", "num")},
                         doc="Pooled budget and engineering capacity; unit review capacity rev_cap and status-quo envelopes B0 "
                             "(unit_id); time values lam (enterprise), lam_u (unit ledger), lam_U (upper bound) by role."),
    "E_capabilities": dict(block="E", index=["capability_id", "scope", "unit_id", "parameter"],
                           params={"F": ("TWD/period", "num"), "eng": ("hours/period", "num"), "y_cur": ("binary", "bin")},
                           doc="scope=SHARED (enterprise, unit_id blank) or LOCAL (unit-built alternative, unit_id required). "
                               "y_cur only for SHARED: 1 if already in operation."),
    "G_policy": dict(block="G", index=["policy_id", "policy_role", "parameter", "initiative_id", "config_id", "unit_id", "capability_id"],
                     params={"gov_cost": ("TWD/period", "num"), "allow": ("binary", "bin"), "allow_loc": ("binary", "bin"),
                             "G": ("ordinal_level", "ord"), "ebar": ("incidents/period", "num"), "must": ("binary", "bin"),
                             "forbid": ("binary", "bin")},
                     doc="policy_role in CURRENT / LEGAL_MINIMUM / ALTERNATIVE. allow per (initiative, config); allow_loc per "
                         "(unit, capability); G, ebar, must, forbid per initiative."),
    "G_requirements": dict(block="G", index=["policy_id", "parameter", "A_level", "kappa_level"],
                           params={"H_req": ("ordinal_level", "ord"), "WI_req": ("ordinal_level", "ord"),
                                   "EA_req": ("ordinal_level", "ord"), "G_req": ("ordinal_level", "ord"),
                                   "R_req": ("ordinal_level", "ord"), "D_req": ("ordinal_level", "ord")},
                           doc="Requirement tables keyed by (A level, kappa level), levels 0-3. R_req/D_req are enterprise-wide: "
                               "policy_id='*', kappa_level blank."),
    "X_initiatives": dict(block="X", index=["initiative_id", "unit_id", "parameter", "capability_id"],
                          params={"kappa": ("ordinal_level", "ord"), "relevant_capabilities": ("list", "list"),
                                  "tau": ("TWD/period", "num")},
                          doc="Owner unit, consequence class kappa, the capabilities whose state indexes the initiative's lookups "
                              "(semicolon list; NONE = empty list), chargeback tau per capability (DEFINITIONAL 0 only if the rule says none)."),
    "X_configurations": dict(block="X", index=["initiative_id", "config_id", "parameter"],
                             params={"is_status_quo": ("binary", "bin"), "A": ("ordinal_level", "ord"), "H": ("ordinal_level", "ord"),
                                     "WI": ("ordinal_level", "ord"), "EA": ("ordinal_level", "ord"), "EAX": ("ordinal_level", "ord"),
                                     "m": ("nominal", "nominal"), "prerequisites": ("list", "list"), "legal": ("binary", "bin"),
                                     "eff_H": ("binary", "bin"), "fallback": ("binary", "bin"), "cons": ("binary", "bin"),
                                     "allowed_responses": ("list", "list"), "impl_x": ("TWD/period", "num")},
                             doc="One status-quo configuration per initiative (is_status_quo=1). EAX is recorded as a rho grouping "
                                 "key only (conditional candidate; never a requirement)."),
    "X_config_costs": dict(block="X", index=["initiative_id", "config_id", "capability_state", "parameter"],
                           params={"impl": ("TWD/period", "num"), "eng": ("hours/period", "num")},
                           doc="Implementation cost and engineering hours by capability-source state, e.g. 'CAP1:S' (shared), "
                               "'CAP1:L' (local), 'NONE'; '*' = invariant across states (assumption_note required)."),
    "X_capability_state": dict(block="X", index=["initiative_id", "capability_state", "parameter"],
                               params={"R_level": ("ordinal_level", "ord"), "D_level": ("ordinal_level", "ord")},
                               doc="Data availability R_j(s) and interface level D_j(s) per capability-source state (F5e)."),
    "R_roles": dict(block="R", index=["initiative_id", "config_id", "role", "parameter"], params={"pi": ("probability", "prob")},
                    doc="Exposure weight of each role group; rows of one (initiative, config) sum to 1."),
    "R_signal_outcome": dict(block="R", index=["initiative_id", "config_id", "signal", "outcome", "parameter"],
                             params={"P": ("probability", "prob")},
                             doc="Joint probability of visible signal z and hidden outcome omega; rows sum to 1; z and omega labels disjoint."),
    "R_responses": dict(block="R", index=["initiative_id", "config_id", "signal", "role", "capability_state", "rho_kind", "response", "parameter"],
                        params={"rho": ("probability", "prob")},
                        doc="rho(r | k, z, role[, capability state]); rho_kind in design / observed / counterfactual. Never "
                            "conditioned on the hidden outcome. capability_state blank = all states."),
    "V_volume": dict(block="V", index=["initiative_id", "config_id", "parameter"], params={"N": ("events/period", "num")},
                     doc="Events per period by configuration."),
    "V_outcomes": dict(block="V", index=["initiative_id", "config_id", "role", "signal", "outcome", "response", "parameter"],
                       params={f: (S.PATH_UNITS[f].replace("TWD/event", "TWD/event"), "prob" if f == "I_sev" else "num") for f in PATH_FIELDS},
                       doc="Per-event path values: v_u/v_x value (unit ledger / outside it), l_u/l_x loss, c_u/c_x operating cost, "
                           "t role time, t_rev review time, I_sev severe-incident probability. '*' = invariant (note required)."),
    "U_registry": dict(block="U", cols=["group_id", "meaning", "eps_R", "tie_tol", "registered_before_results", "registration_date",
                                        "grouping_rule", "notes"], values=False,
                       doc="Scenario groups registered BEFORE any result is seen. grouping_rule documents how interval cells "
                           "are grouped into blocks (default: one block per elicited interval row)."),
}
for t in TABLES.values():
    if "params" in t:
        t["cols"] = t["index"] + VAL
WILDCARD_COLS = {"role", "signal", "outcome", "response", "capability_state"}
PER_PERIOD = ("/period",)


# =============================================================================
# 2. Reading
# =============================================================================
def read_data(path):
    tables = {}
    if os.path.isdir(path):
        for name, t in TABLES.items():
            p = os.path.join(path, f"{name}.csv")
            if not os.path.exists(p):
                tables[name] = []
                continue
            with open(p, encoding="utf-8-sig") as f:
                rd = csv.DictReader(f)
                hdr = [h.strip() for h in (rd.fieldnames or [])]
                if hdr != t["cols"]:
                    raise ValueError(f"{name}.csv: header must be {','.join(t['cols'])}")
                tables[name] = [{k.strip(): (v or "").strip() for k, v in r.items()} for r in rd]
    elif path.lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        for name, t in TABLES.items():
            if name not in wb.sheetnames:
                tables[name] = []
                continue
            ws = wb[name]
            rows = list(ws.iter_rows(values_only=True))
            hdr = [str(h).strip() if h is not None else "" for h in rows[0]][:len(t["cols"])]
            if hdr != t["cols"]:
                raise ValueError(f"sheet {name}: header must be {','.join(t['cols'])}")
            out = []
            for r in rows[1:]:
                vals = ["" if v is None else str(v).strip() for v in r[:len(t["cols"])]]
                if any(vals):
                    out.append(dict(zip(t["cols"], vals)))
            tables[name] = out
    else:
        raise ValueError("DATA must be a folder of 62 CSVs or the 62 .xlsx workbook")
    return tables


# =============================================================================
# 3. Validation (43 section 0 and 5; 62 rules)
# =============================================================================
def _float(s):
    try:
        x = float(s)
    except ValueError:
        return None
    return x if math.isfinite(x) else None


def validate(tables):
    errors, warnings = [], []
    srcs = {r["source_id"]: r for r in tables.get("S_sources", [])}
    ids = [r["source_id"] for r in tables.get("S_sources", [])]
    for sid in sorted({x for x in ids if ids.count(x) > 1}):
        errors.append(f"S_sources: duplicate source_id {sid} (each source has exactly one scope)")
    for sid, r in srcs.items():
        if r["source_scope"] not in ("internal", "external"):
            errors.append(f"S_sources {sid}: source_scope must be internal/external")
    for i, r in enumerate(tables.get("E_capabilities", [])):
        if r["scope"] not in ("SHARED", "LOCAL"):
            errors.append(f"E_capabilities row {i + 2}: scope must be SHARED or LOCAL")
        elif r["scope"] == "LOCAL" and (not r["unit_id"] or r["parameter"] == "y_cur"):
            errors.append(f"E_capabilities row {i + 2}: LOCAL rows need unit_id and only F/eng")
        elif r["scope"] == "SHARED" and r["unit_id"]:
            errors.append(f"E_capabilities row {i + 2}: SHARED rows have no unit_id")
    for i, r in enumerate(tables.get("R_responses", [])):
        if r["rho_kind"] not in ("design", "observed", "counterfactual"):
            errors.append(f"R_responses row {i + 2}: rho_kind must be design/observed/counterfactual")
        if r["capability_state"] == "*":
            errors.append(f"R_responses row {i + 2}: capability_state '*' not allowed (leave blank for all states)")
        if r["response"] not in RESP:
            errors.append(f"R_responses row {i + 2}: response must be one of {RESP}")
    for i, r in enumerate(tables.get("G_policy", [])):
        if r["policy_role"] not in ("CURRENT", "LEGAL_MINIMUM", "ALTERNATIVE"):
            errors.append(f"G_policy row {i + 2}: policy_role must be CURRENT/LEGAL_MINIMUM/ALTERNATIVE")
    period = set()
    for name, t in TABLES.items():
        if "params" not in t:
            continue
        seen = set()
        for i, r in enumerate(tables.get(name, [])):
            where = f"{name} row {i + 2}"
            p = r["parameter"]
            if p not in t["params"]:
                errors.append(f"{where}: unknown parameter {p!r}")
                continue
            unit, kind = t["params"][p]
            idx = tuple(r[c] for c in t["index"])
            if idx in seen:
                errors.append(f"{where}: duplicate index {idx}")
            seen.add(idx)
            if any(r.get(c) == "*" for c in WILDCARD_COLS & set(t["index"])) and not r["assumption_note"]:
                errors.append(f"{where}: '*' (invariance) requires an assumption_note stating the invariance")
            prov, unc, st = r["provenance_type"], r["uncertainty"], r["identification_status"]
            if prov not in PROV or unc not in UNC or st not in IDS:
                errors.append(f"{where}: provenance_type/uncertainty/identification_status not in the allowed lists")
                continue
            v, lo, hi = r["value"], r["lower"], r["upper"]
            if prov == "UNIDENTIFIED" or st == "UNIDENTIFIED":
                if not (prov == "UNIDENTIFIED" and st == "UNIDENTIFIED" and unc == "NONE"):
                    errors.append(f"{where}: UNIDENTIFIED requires provenance UNIDENTIFIED, status UNIDENTIFIED, uncertainty NONE")
                if v or lo or hi:
                    errors.append(f"{where}: UNIDENTIFIED cell must have blank value/lower/upper (no zero-fill)")
                continue
            if r["unit"] != unit:
                errors.append(f"{where}: unit must be {unit!r} (got {r['unit']!r})")
            if not r["source_id"]:
                errors.append(f"{where}: source_id required")
            elif r["source_id"] not in srcs:
                errors.append(f"{where}: source_id {r['source_id']} not registered in S_sources")
            else:
                if srcs[r["source_id"]]["source_scope"] == "external" and prov not in ("EXPERT_ELICITED_EXTERNAL", "PUBLIC_OBSERVED_BOUND"):
                    errors.append(f"{where}: external source can only give EXPERT_ELICITED_EXTERNAL or PUBLIC_OBSERVED_BOUND "
                                  "(never OBSERVED, never an optimizer)")
                if prov == "EXPERT_ELICITED_EXTERNAL" and srcs[r["source_id"]]["source_scope"] != "external":
                    errors.append(f"{where}: EXPERT_ELICITED_EXTERNAL needs an external source")
            if any(s in unit for s in PER_PERIOD) or unit.endswith("/event"):
                if not r["period"]:
                    errors.append(f"{where}: period required for {unit}")
                else:
                    period.add(r["period"])
            if unc == "POINT":
                if not v:
                    errors.append(f"{where}: POINT requires a value")
                if (lo or hi) and not (lo == v and hi == v):
                    errors.append(f"{where}: POINT must leave lower/upper blank (or equal to value)")
                if prov in ("PUBLIC_OBSERVED_BOUND", "SCENARIO_SENSITIVITY"):
                    errors.append(f"{where}: {prov} cannot be a point value (43 section 0)")
                if prov.startswith("EXPERT_ELICITED") and kind == "num" and not r["assumption_note"]:
                    errors.append(f"{where}: an expert point value needs a justification in assumption_note (43: interval by default)")
                if st != "IDENTIFIED_POINT":
                    errors.append(f"{where}: POINT cells have status IDENTIFIED_POINT")
            elif unc == "INTERVAL":
                if kind in ("ord", "bin", "list", "nominal"):
                    errors.append(f"{where}: {kind} parameters take a POINT value or UNIDENTIFIED")
                if v:
                    errors.append(f"{where}: INTERVAL leaves value blank (no invented midpoint)")
                a, b = _float(lo), _float(hi)
                if a is None or b is None or a > b:
                    errors.append(f"{where}: INTERVAL needs finite lower <= upper")
                need = {"PUBLIC_OBSERVED_BOUND": "BOUND_ONLY", "SCENARIO_SENSITIVITY": "SCENARIO_ONLY"}.get(prov, "IDENTIFIED_INTERVAL")
                if st != need:
                    errors.append(f"{where}: {prov} interval must have status {need}")
                if prov == "PUBLIC_OBSERVED_BOUND" and p in ("t", "t_rev", "v_u", "v_x") and "task similarity" not in r["assumption_note"].lower():
                    errors.append(f"{where}: public effect-size bounds on {p} need a 'task similarity' statement (43 section 0)")
            else:
                errors.append(f"{where}: identified cells use POINT or INTERVAL")
            for x in (v, lo, hi):
                if not x:
                    continue
                if kind in ("num", "prob"):
                    f = _float(x)
                    if f is None:
                        errors.append(f"{where}: non-finite or non-numeric {x!r}")
                    elif kind == "prob" and not (0 <= f <= 1):
                        errors.append(f"{where}: probability outside [0,1]")
                elif kind == "ord" and x not in ("0", "1", "2", "3"):
                    errors.append(f"{where}: ordinal level must be 0-3 (anchor level, not a cardinal score)")
                elif kind == "bin" and x not in ("0", "1"):
                    errors.append(f"{where}: binary must be 0/1")
                elif kind == "nominal" and p == "m" and x not in CONFIG_MODES:
                    errors.append(f"{where}: m must be one of {CONFIG_MODES}")
            if prov == "DEFINITIONAL" and kind == "num" and _float(v or "nan") == 0 and not r["assumption_note"]:
                warnings.append(f"{where}: DEFINITIONAL zero without a note — state the rule that makes it zero")
    if len(period) > 1:
        errors.append(f"mixed period definitions {sorted(period)}: rescale to one period first")
    return errors, warnings


# =============================================================================
# 4. Coverage (44 Table 4.5)
# =============================================================================
def coverage(tables):
    blocks = {}
    for name, t in TABLES.items():
        if "params" not in t:
            continue
        b = blocks.setdefault(t["block"], {k: 0 for k in ["cells"] + PROV})
        for r in tables.get(name, []):
            b["cells"] += 1
            b[r["provenance_type"] if r["provenance_type"] in PROV else "UNIDENTIFIED"] += 1
    reg = tables.get("U_registry", [])
    blocks["U"] = {"cells": len(reg), "registered_before_results": sum(r["registered_before_results"] == "yes" for r in reg)}
    return blocks


# =============================================================================
# 5. Instance building (point-identified cells only)
# =============================================================================
def _pattern(s):
    if s in ("", "NONE"):
        return frozenset()
    out = []
    for part in s.split(";"):
        c, src = part.split(":")
        if src not in ("S", "L"):
            raise ValueError(f"capability state {s!r}: source must be S or L")
        out.append((c.strip(), src.strip()))
    return frozenset(out)


def _pattern_str(p):
    return "NONE" if not p else ";".join(f"{c}:{s}" for c, s in sorted(p))


def _point(r, kind):
    if r["identification_status"] != "IDENTIFIED_POINT":
        return None
    v = r["value"]
    if kind in ("num", "prob"):
        return float(v)
    if kind in ("ord", "bin"):
        return int(v)
    if kind == "list":
        return [] if v == "NONE" else [x.strip() for x in v.split(";") if x.strip()]
    return v


def _interval(r):
    if r["uncertainty"] == "INTERVAL":
        return float(r["lower"]), float(r["upper"])
    return None


class Built:
    def __init__(self):
        self.inst, self.log, self.assumptions, self.kmap = None, [], [], {}
        self.winner = {}          # (target tuple, field) -> row that determined the cell (after override ordering)
        self.prob_rows = {}       # (target tuple, field) -> {entry: row} for probability rows (pi, P, rho)
        self.policy_role = {}


def _nwild(r):
    return sum(r.get(c) == "*" for c in WILDCARD_COLS if c in r)


def _ordered(rows):
    """Wildcard rows first (most '*' first), specific rows last: a specific row (even UNIDENTIFIED) always wins."""
    return sorted(rows, key=lambda r: -_nwild(r))


def build_instance(tables):
    """Returns Built with .inst (51 instance; None = UNIDENTIFIED), .log (loader exclusions), .assumptions (explicit
    DEF assumptions), .winner / .prob_rows (which template row determined each cell; used by the registry)."""
    out = Built()
    T = {n: tables.get(n, []) for n in TABLES}
    O = lambda c, lvl: S.O(c, lvl)
    pv = lambda r: _point(r, TABLES_KIND[(r["_t"], r["parameter"])])
    for n in T:
        for r in T[n]:
            r["_t"] = n

    def setc(container, target, field, r):
        container[field] = pv(r)
        out.winner[(tuple(target), field)] = r

    units = sorted({r["unit_id"] for r in T["X_initiatives"] if r["unit_id"]})
    ent = {(r["parameter"], r["unit_id"], r["role"]): r for r in T["E_enterprise"]}
    roles = sorted({r["role"] for r in T["R_roles"] if r["role"] not in ("", "*")})
    inst = dict(units=units, caps={}, loc_caps={}, initiatives={}, policies={}, y_cur=frozenset(),
                A_bar={"BUD": None, "ENG": None}, B0={}, rev_cap={u: None for u in units},
                lam={ro: None for ro in roles}, lam_u={ro: None for ro in roles}, lam_U={ro: None for ro in roles},
                R_req={}, D_req={}, codes=S.DEFAULT_CODES, provenance={})
    for (p, u, ro), r in ent.items():
        if p in ("A_bar_BUD", "A_bar_ENG"):
            setc(inst["A_bar"], ("A_bar",), p.split("_")[-1], r)
        elif p == "rev_cap":
            setc(inst["rev_cap"], ("rev_cap",), u, r)
        elif p in ("B0_BUD", "B0_ENG"):
            setc(inst["B0"].setdefault(u, {"BUD": None, "ENG": None}), ("B0", u), p.split("_")[-1], r)
        elif p in ("lam", "lam_u", "lam_U"):
            setc(inst[p], (p,), ro, r)
    for ro in roles:                        # lam interval: upper bound feeds lam_U (box [0, lam_U]; 42 v2.2 section 7.4)
        r = ent.get(("lam", "", ro))
        if r and _interval(r) and inst["lam_U"][ro] is None:
            inst["lam_U"][ro] = _interval(r)[1]
            out.assumptions.append(f"lam_U[{ro}] := upper bound of the elicited lam interval (scan box [0, lam_U])")
    ycur_unknown = set()
    for r in T["E_capabilities"]:
        c = r["capability_id"]
        if r["scope"] == "SHARED":
            d = inst["caps"].setdefault(c, {"F": None, "eng": None})
            if r["parameter"] == "y_cur":
                v = pv(r)
                if v is None:
                    ycur_unknown.add(c)
                elif v == 1:
                    inst["y_cur"] = inst["y_cur"] | {c}
            else:
                setc(d, ("caps", c), r["parameter"], r)
        elif r["scope"] == "LOCAL" and r["parameter"] in ("F", "eng"):
            setc(inst["loc_caps"].setdefault((r["unit_id"], c), {"F": None, "eng": None}), ("loc_caps", (r["unit_id"], c)), r["parameter"], r)
    for c in inst["caps"]:
        if not any(r["capability_id"] == c and r["scope"] == "SHARED" and r["parameter"] == "y_cur" for r in T["E_capabilities"]):
            ycur_unknown.add(c)
    for c in sorted(ycur_unknown):
        out.log.append(("UNIDENTIFIED_Y_CUR", c, "current operation status of the shared capability is UNIDENTIFIED: initiatives that "
                                                 "use it cannot fix their status-quo baseline and are excluded (never assumed 0)"))
    # requirement tables
    reqs = {}
    for r in T["G_requirements"]:
        reqs.setdefault(r["policy_id"], []).append(r)
    for a in range(4):
        for t in ("R_req", "D_req"):
            inst[t][O("A", a)] = None
    for r in reqs.get("*", []):
        if r["parameter"] in ("R_req", "D_req"):
            v = pv(r)
            inst[r["parameter"]][O("A", int(r["A_level"]))] = None if v is None else O(r["parameter"][0], v)
    # initiatives
    ini_rows, cfg_rows = {}, {}
    for r in T["X_initiatives"]:
        ini_rows.setdefault(r["initiative_id"], []).append(r)
    for r in T["X_configurations"]:
        cfg_rows.setdefault((r["initiative_id"], r["config_id"]), {})[r["parameter"]] = r
    for j in sorted(ini_rows):
        rows = ini_rows[j]
        unit = next((r["unit_id"] for r in rows if r["unit_id"]), None)
        kap = next((pv(r) for r in rows if r["parameter"] == "kappa"), None)
        rel_rows = [r for r in rows if r["parameter"] == "relevant_capabilities"]
        cfg_ids = sorted({c for (jj, c) in cfg_rows if jj == j})
        sq = [c for c in cfg_ids if cfg_rows[(j, c)].get("is_status_quo") and pv(cfg_rows[(j, c)]["is_status_quo"]) == 1]
        if unit is None or kap is None or len(sq) != 1:
            out.log.append(("UNEVALUABLE_INITIATIVE", j, "owner unit, kappa or exactly one status-quo config UNIDENTIFIED"))
            continue
        if rel_rows:
            rel = pv(rel_rows[0])
            if rel is None:
                out.log.append(("UNEVALUABLE_INITIATIVE", j, "relevant_capabilities explicitly UNIDENTIFIED: capability-state keys unknown"))
                continue
        else:
            rel = sorted({c for cid in cfg_ids for c in (pv(cfg_rows[(j, cid)]["prerequisites"]) or [])
                          if "prerequisites" in cfg_rows[(j, cid)]})
            out.assumptions.append(f"{j}: relevant_capabilities row absent; DEF := union of configuration prerequisites {rel}")
        if set(rel) & ycur_unknown:
            out.log.append(("UNEVALUABLE_INITIATIVE", j, f"y_cur UNIDENTIFIED for {sorted(set(rel) & ycur_unknown)}"))
            continue
        kmap = {sq[0]: 0}
        for n, c in enumerate([c for c in cfg_ids if c != sq[0]], start=1):
            kmap[c] = n
        out.kmap[j] = kmap
        caps, loc_caps = inst["caps"], inst["loc_caps"]
        pats = [frozenset(p) for p in itertools.chain.from_iterable(
            itertools.product(*[[(c, s) for s in ("S", "L") if (s == "S" and c in caps) or (s == "L" and (unit, c) in loc_caps)]
                                for c in sub]) for sub in S._subsets(sorted(rel)))]
        tau = {c: None for c in rel}
        for r in rows:
            if r["parameter"] == "tau" and r["capability_id"] in tau:
                setc(tau, ("initiatives", j, "tau"), r["capability_id"], r)
        Rst, Dst = {p: None for p in pats}, {p: None for p in pats}
        for r in _ordered([r for r in T["X_capability_state"] if r["initiative_id"] == j]):
            targets = pats if r["capability_state"] == "*" else [_pattern(r["capability_state"])]
            for p in targets:
                if p not in Rst:
                    raise ValueError(f"{j}: capability state {_pattern_str(p)} impossible for this initiative")
                v = pv(r)
                lvl = r["parameter"][0]
                (Rst if lvl == "R" else Dst)[p] = None if v is None else O(lvl, v)
        N = {}
        for r in T["V_volume"]:
            if r["initiative_id"] == j and r["config_id"] in kmap:
                setc(N, ("initiatives", j, "N"), kmap[r["config_id"]], r)
        configs = {}
        for cid in cfg_ids:
            k = kmap[cid]
            cr = cfg_rows[(j, cid)]
            g = lambda p: pv(cr[p]) if p in cr else None
            attrs = {p: g(p) for p in ("A", "H", "WI", "EA", "m", "legal", "eff_H", "fallback", "cons", "allowed_responses", "prerequisites")}
            missing = [p for p, v in attrs.items() if v is None]
            if missing:
                out.log.append(("UNEVALUABLE_INITIATIVE" if k == 0 else "EXCLUDED_CONFIG", f"{j}/{cid}",
                                f"configuration attributes UNIDENTIFIED: {missing}"))
                if k == 0:
                    configs = None
                    break
                continue
            allowed = set(attrs["allowed_responses"])
            if not allowed <= set(RESP):
                raise ValueError(f"{j}/{cid}: allowed_responses must be among {RESP}")
            base = ("initiatives", j, "configs", k)
            pi, P = {}, {}
            for r in T["R_roles"]:
                if r["initiative_id"] == j and r["config_id"] == cid:
                    pi[r["role"]] = pv(r)
                    out.prob_rows.setdefault((base, "pi"), {})[r["role"]] = r
            for r in T["R_signal_outcome"]:
                if r["initiative_id"] == j and r["config_id"] == cid:
                    P[(r["signal"], r["outcome"])] = pv(r)
                    out.prob_rows.setdefault((base, "P"), {})[(r["signal"], r["outcome"])] = r
            Z = sorted({z for z, _ in P}); W = sorted({w for _, w in P}); RL = sorted(pi)
            impl, eng = {p: None for p in pats}, {p: None for p in pats}
            for r in _ordered([r for r in T["X_config_costs"] if r["initiative_id"] == j and r["config_id"] == cid]):
                tg = pats if r["capability_state"] == "*" else [_pattern(r["capability_state"])]
                for p in tg:
                    if p not in impl:
                        raise ValueError(f"{j}/{cid}: capability state {_pattern_str(p)} impossible (capability or local version missing)")
                    setc(impl if r["parameter"] == "impl" else eng, base + (r["parameter"],), p, r)
            tabs = {}
            for kind in ("observed", "counterfactual", "design"):
                tab = {}
                for r in _ordered([r for r in T["R_responses"] if r["initiative_id"] == j and r["config_id"] == cid and r["rho_kind"] == kind]):
                    zs = Z if r["signal"] == "*" else [r["signal"]]
                    rs = RL if r["role"] == "*" else [r["role"]]
                    for z in zs:
                        for ro in rs:
                            key = (z, ro) if not r["capability_state"] else (z, ro, _pattern(r["capability_state"]))
                            tab.setdefault(key, {})[r["response"]] = r
                tabs[kind] = tab
            rho, rho_src = {}, {}
            for kind in ("observed", "counterfactual"):
                for key, rowmap in tabs[kind].items():
                    if key in rho:
                        continue
                    vals = {resp: pv(r) for resp, r in rowmap.items()}
                    if vals and all(v is not None for v in vals.values()):
                        rho[key], rho_src[key] = vals, kind
                        out.prob_rows[(base + ("rho",), key)] = rowmap
            for kind in ("observed", "counterfactual"):               # interval-only rows: remembered for the registry
                for key, rowmap in tabs[kind].items():
                    if key not in rho_src and all(r["identification_status"] != "UNIDENTIFIED" for r in rowmap.values()):
                        out.prob_rows.setdefault((base + ("rho",), key), rowmap)
            rho_design = {key: {resp: pv(r) for resp, r in rowmap.items()} for key, rowmap in tabs["design"].items()}
            if any(v is None for row in rho_design.values() for v in row.values()):
                rho_design = {}
                out.log.append(("RHO_DESIGN_INCOMPLETE", f"{j}/{cid}", "rho_design has UNIDENTIFIED entries: Q6 not computable"))
            for z in Z:
                for ro in RL:
                    if (z, ro) not in rho:
                        out.log.append(("EXCLUDED_CONFIG" if k else "UNEVALUABLE_INITIATIVE", f"{j}/{cid}",
                                        f"rho row ({z},{ro}) has no complete observed or counterfactual point row"))
            path = {}
            for ro in RL:
                for z in Z:
                    for w in W:
                        for r_ in sorted(allowed):
                            path[(ro, z, w, r_)] = {f: None for f in PATH_FIELDS}
            for r in _ordered([r for r in T["V_outcomes"] if r["initiative_id"] == j and r["config_id"] == cid]):
                for key in path:
                    if all(r[c] in ("*", key[i]) for i, c in enumerate(("role", "signal", "outcome", "response"))):
                        setc(path[key], base + ("path", key), r["parameter"], r)
            ix_row = cr.get("impl_x")
            cfgd = dict(A=O("A", attrs["A"]), H=O("H", attrs["H"]), WI=O("WI", attrs["WI"]), EA=O("EA", attrs["EA"]),
                        m=attrs["m"], pre=set(attrs["prerequisites"] or []), legal=attrs["legal"], effH=attrs["eff_H"],
                        FB=attrs["fallback"], cons=attrs["cons"], impl=impl, eng=eng, impl_x=None,
                        pi=pi, P=P, rho=rho, rho_design=rho_design, allowed_r=allowed, path=path)
            if ix_row is not None:
                setc(cfgd, base, "impl_x", ix_row)
            configs[k] = cfgd
        if configs is None or 0 not in configs:
            if configs is not None:
                out.log.append(("UNEVALUABLE_INITIATIVE", j, "status quo excluded"))
            continue
        inst["initiatives"][j] = dict(unit=unit, kappa=O("K", kap), N={k: N.get(k) for k in configs}, rel=sorted(rel), tau=tau,
                                      R_state=Rst, D_state=Dst, configs=configs)
    # policies
    pol_rows = {}
    for r in T["G_policy"]:
        pol_rows.setdefault(r["policy_id"], []).append(r)
    for g in sorted(pol_rows):
        rows = pol_rows[g]
        out.policy_role[g] = next((r["policy_role"] for r in rows if r["policy_role"]), "")
        pol = dict(gov_cost=None, allow={}, allow_loc={}, G={}, ebar={}, must=set(), forbid=set(),
                   H_req={}, WI_req={}, EA_req={}, G_req={})
        has = {"must": False, "forbid": False}
        for r in rows:
            p, v = r["parameter"], pv(r)
            j = r["initiative_id"]
            if p == "gov_cost":
                setc(pol, ("policies", g), "gov_cost", r)
            elif p == "allow":
                if j in out.kmap and r["config_id"] in out.kmap[j]:
                    pol["allow"][(j, out.kmap[j][r["config_id"]])] = v
            elif p == "allow_loc":
                pol["allow_loc"][(r["unit_id"], r["capability_id"])] = v
            elif p == "G":
                pol["G"][j] = None if v is None else O("G", v)
            elif p == "ebar":
                setc(pol["ebar"], ("policies", g, "ebar"), j, r)
            elif p in ("must", "forbid"):
                has[p] = True
                if v == 1:
                    pol[p].add(j)
                elif v is None:
                    out.log.append(("UNIDENTIFIED_POLICY_SET", g, f"{p}[{j}] UNIDENTIFIED: treated as not in the set — STATED; "
                                                                  "decisions involving {j} are conditional on it"))
        for p in ("must", "forbid"):
            if not has[p]:
                out.assumptions.append(f"policy {g}: J^{p} has no rows -> empty set (43 section 1.2: stated explicitly, not imputed)")
        for r in reqs.get(g, []):
            if r["parameter"] in ("H_req", "WI_req", "EA_req", "G_req"):
                v = pv(r)
                pol[r["parameter"]][(O("A", int(r["A_level"])), O("K", int(r["kappa_level"])))] = \
                    None if v is None else O(r["parameter"].split("_")[0], v)
        for j in inst["initiatives"]:
            pol["G"].setdefault(j, None)
            pol["ebar"].setdefault(j, None)
            for k in inst["initiatives"][j]["configs"]:
                pol["allow"].setdefault((j, k), None)
        for (u, c) in inst["loc_caps"]:
            pol["allow_loc"].setdefault((u, c), None)
        inst["policies"][g] = pol
    prov = {}
    for n in T:
        for r in T[n]:
            if "provenance_type" in r:
                prov[r["provenance_type"]] = prov.get(r["provenance_type"], 0) + 1
    inst["provenance"] = {f"cells:{k}": k for k in prov}
    out.provenance_counts = prov
    out.inst = inst
    return out


TABLES_KIND = {(n, p): k for n, t in TABLES.items() for p, (_, k) in t.get("params", {}).items()}


# =============================================================================
# 6. Runs and Chapter 4 quantities
# =============================================================================
def _f(x):
    return None if x is None else x


def run_point(built, out_dir):
    inst = built.inst
    res = dict(loader=VERSION, loader_log=built.log, assumptions=built.assumptions, runs={}, derived={})
    pols = sorted(inst["policies"])
    if not inst["initiatives"] or not pols:
        res["overall"] = "NO_EVALUABLE_INITIATIVE_OR_POLICY"
        return res
    cur = [g for g in pols if built.policy_role.get(g) == "CURRENT"] or pols[:1]
    legal = [g for g in pols if built.policy_role.get(g) == "LEGAL_MINIMUM"]
    for g in pols:
        for model in ("B", "C", "Aplus", "C0"):
            r = S.solve(inst, model, [g])
            res["runs"][f"{model}|{g}"] = r
    res["runs"][f"A|{cur[0]}"] = S.solve(inst, "A", [cur[0]])
    g = cur[0]
    get = lambda m: res["runs"].get(f"{m}|{g}", {})
    pt = lambda r: r.get("F_E") if r.get("status") in ("OK_POINT", "RESTRICTED_SOLVE", "STATUS_QUO_NONCOMPLIANT") else None
    A, C0, Ap, B, C = (pt(get(m)) for m in ("A", "C0", "Aplus", "B", "C"))
    d = {}
    if None not in (A, C0, Ap, B, C):
        d["architecture"] = dict(A=A, C0=C0, Aplus=Ap, B=B, C=C,
                                 VRA={m: C0[m] - A[m] for m in ("opt", "pess")}, DL0={m: Ap - C0[m] for m in ("opt", "pess")},
                                 VCP={m: Ap - A[m] for m in ("opt", "pess")}, VSC=B - Ap, DL={m: B - C[m] for m in ("opt", "pess")},
                                 NEV={m: C[m] - A[m] for m in ("opt", "pess")})
    else:
        d["architecture"] = dict(status="UNIDENTIFIED", available={m: v is not None for m, v in zip(("A", "C0", "Aplus", "B", "C"), (A, C0, Ap, B, C))},
                                 partial=dict(VSC=(B - Ap) if None not in (B, Ap) else None,
                                              DL={m: B - C[m] for m in ("opt", "pess")} if None not in (B, C) else None))
    # Q4 DL source decomposition (42 v2.2 section 8.2): close one alignment condition at a time, recompute DL
    if None not in (B, C):
        d["dl_sources"] = _dl_sources(inst, g, B)
    else:
        d["dl_sources"] = dict(status="UNIDENTIFIED", reason="needs B and C point solutions (C needs the ledger split)")
    # Q2 thresholds: F_c set UNIDENTIFIED in a copy -> THRESHOLD_MODE of 51
    th = {}
    for c in inst["caps"]:
        i2 = copy.deepcopy(inst); i2["caps"][c]["F"] = None
        for m in ("B", "C"):
            r = S.solve(i2, m, [g])
            th[f"{c}|{m}"] = dict(status=r["status"], threshold=r["details"].get("threshold", {}).get(c),
                                  F_c_identified=inst["caps"][c]["F"])
    d["capability_thresholds"] = th
    # Q3 governance decomposition (needs a legal-minimum policy)
    if legal and B is not None and pt(res["runs"].get(f"B|{legal[0]}", {})) is not None:
        try:
            d["governance"] = S.governance_decomposition(inst, g, legal[0])
        except Exception as e:                                   # noqa: BLE001  (reported, never hidden)
            d["governance"] = dict(status="UNIDENTIFIED", error=f"{type(e).__name__}: {e}")
    else:
        d["governance"] = dict(status="UNIDENTIFIED", reason="needs a solvable LEGAL_MINIMUM policy and current policy (43 section 2 Q3)")
    # Q5 price of requirement (relax one level, B)
    pr = {}
    if B is not None:
        for X in ("H_req", "WI_req", "EA_req", "G_req"):
            i2 = copy.deepcopy(inst)
            tab = i2["policies"][g][X]
            if any(v is None for v in tab.values()) or not tab:
                pr[X] = dict(status="UNIDENTIFIED", reason="requirement table has UNIDENTIFIED cells")
                continue
            for key, v in tab.items():
                tab[key] = S.O(v.construct, max(0, int(v.code) - 1))
            r = S.solve(i2, "B", [g])
            pr[X] = dict(PR=(r["F_E"] - B) if pt(r) is not None else None, status=r["status"],
                         relaxed_configs=r.get("configs"), base_configs=get("B").get("configs"))
        d["price_of_requirement"] = pr
        d["requirement_status"] = _at_above(inst, g, get("B").get("configs") or [])
    # Q6
    if B is not None and all(cfg.get("rho_design") for ini in inst["initiatives"].values() for kk, cfg in ini["configs"].items() if kk):
        try:
            d["design_response"] = S.design_response_analysis(inst, g)
        except Exception as e:                                   # noqa: BLE001
            d["design_response"] = dict(status="UNIDENTIFIED", error=f"{type(e).__name__}: {e}")
    else:
        d["design_response"] = dict(status="UNIDENTIFIED", reason="rho_design missing for some AI configuration or B not solvable")
    # Q8 interface activation (point; ROBUST variants need the registry run)
    d["interface_activation"] = _interfaces(pr, get("B").get("configs") or [], inst)
    res["derived"] = d
    res["policy_current"], res["policy_legal_minimum"] = g, legal[0] if legal else None
    return res


def _close(inst, which):
    i = copy.deepcopy(inst)
    for ini in i["initiatives"].values():
        if which in ("AL1", "ALL"):
            for k, cfg in ini["configs"].items():
                if k and cfg.get("impl_x") is not None:
                    cfg["impl_x"] = 0.0
                for cell in cfg["path"].values():
                    for u_, x_ in (("v_u", "v_x"), ("l_u", "l_x"), ("c_u", "c_x")):
                        if cell[u_] is not None and cell[x_] is not None:
                            cell[u_], cell[x_] = cell[u_] + cell[x_], 0.0      # unit internalises; enterprise sum unchanged
        if which in ("AL2", "ALL"):
            ini["tau"] = {c: (0.0 if v is not None else None) for c, v in ini["tau"].items()}
    if which in ("AL4", "ALL"):
        i["lam_u"] = {r: i["lam"].get(r) for r in i["lam_u"]}
    return i


def _dl_sources(inst, g, B):
    """DL after closing each (AL) condition separately and all together. (AL3) unit-only constraints are not part of
    the 51 baseline (43: added only with evidence), so it cannot be closed here. Single-closure DLs are not additive."""
    out = {}
    for which in ("AL1", "AL2", "AL4", "ALL"):
        r = S.solve(_close(inst, which), "C", [g])
        Fc = r.get("F_E") if r.get("status") in ("OK_POINT", "RESTRICTED_SOLVE", "STATUS_QUO_NONCOMPLIANT") else None
        out[which] = dict(status=r["status"], DL=None if Fc is None else {m: B - Fc[m] for m in ("opt", "pess")})
    out["note"] = "AL1 = external ledger items and central funding moved into the unit ledger; AL2 = chargeback removed; AL4 = lam_u := lam"
    return out


def _at_above(inst, g, configs):
    out = []
    pol = inst["policies"][g]
    for j, k in configs:
        cfg, kap = inst["initiatives"][j]["configs"][k], inst["initiatives"][j]["kappa"]
        row = {"initiative": j, "config": k}
        for X, attr in (("H_req", "H"), ("WI_req", "WI"), ("EA_req", "EA")):
            req = pol[X].get((cfg["A"], kap))
            row[attr] = None if req is None else ("AT_REQUIREMENT" if cfg[attr].code == req.code else "ABOVE_REQUIREMENT")
        out.append(row)
    return out


def _interfaces(pr, configs, inst):
    pr_pos = lambda X: (pr.get(X) or {}).get("PR")
    act = lambda v: "UNIDENTIFIED" if v is None else ("ACTIVE" if v > 1e-9 else "INACTIVE")
    changed = [jk for jk in configs if jk[1] != 0]
    return {"I1 Authority-to-act (H_req, G_req)": dict(H=act(pr_pos("H_req")), G=act(pr_pos("G_req"))),
            "I2 Routing & recovery (WI_req)": dict(WI=act(pr_pos("WI_req"))),
            "I3 Configuration & response (configs differ from status quo)": dict(changed=changed,
                                                                                   state="ACTIVE" if changed else ("INACTIVE" if configs else "UNIDENTIFIED")),
            "I4 Assurance (EA_req)": dict(EA=act(pr_pos("EA_req"))),
            "note": "point classification only; ACTIVE-ROBUST / ACTIVE-CONDITIONAL need the registered robust run (--robust)"}


# =============================================================================
# 7. Interval registry (pre-registration) and robust / VOI run
# =============================================================================
def simplex_box_vertices(lo, hi, tol=1e-12):
    """Vertices of {x : sum x = 1, lo <= x <= hi}. Each vertex has at least n-1 coordinates at a bound."""
    n, out = len(lo), set()
    for free in range(n):
        others = [i for i in range(n) if i != free]
        for bits in itertools.product((0, 1), repeat=n - 1):
            x = [0.0] * n
            for i, b in zip(others, bits):
                x[i] = hi[i] if b else lo[i]
            x[free] = 1 - sum(x[i] for i in others)
            if lo[free] - tol <= x[free] <= hi[free] + tol:
                out.add(tuple(round(v, 12) for v in x))
    return sorted(out)


NOT_IN_B_ROBUST = {"lam_u", "lam_U", "B0"}          # not used by Model B (robust module is Model B only)


def make_registry(tables, built):
    """Registry draft: one block per elicited interval row, applied ONLY to the cells that row actually determines
    (a more specific row overrides a wildcard row), plus simplex-box blocks for probability rows that are identified only
    by intervals. Returns (instance with nominal values, registry, vertex count, notes, unregistered)."""
    inst = copy.deepcopy(built.inst)
    notes, unregistered, groups = [], [], {}
    for (target, field), r in built.winner.items():
        if r["uncertainty"] != "INTERVAL":
            continue
        if target[0] in NOT_IN_B_ROBUST:
            notes.append(f"interval on {target[0]}[{field}] is not used by the Model B robust run")
            continue
        if target[0] == "initiatives" and target[1] not in inst["initiatives"]:
            continue
        key = (r["_t"], tuple(r[c] for c in TABLES[r["_t"]]["index"]))
        groups.setdefault(key, dict(lo=float(r["lower"]), hi=float(r["upper"]), cells=[], row=r))["cells"].append((target, field))
    blocks = []
    for (tname, idx), g in sorted(groups.items(), key=lambda kv: str(kv[0])):
        lo, hi = g["lo"], g["hi"]
        nominal = hi if abs(hi) >= abs(lo) else lo
        if nominal == 0:                                  # [0, 0]: a degenerate interval is a point
            L = U = 1.0
        else:
            L, U = sorted((lo / nominal, hi / nominal))
        for target, field in g["cells"]:
            S._get(inst, list(target))[field] = nominal
        blocks.append(dict(name=f"{tname}:{'|'.join(idx)}", kind="scale", targets=[(list(t), f) for t, f in g["cells"]],
                           L=L, U=U, provenance=g["row"]["provenance_type"], source_id=g["row"]["source_id"]))
    for (target, field), rowmap in sorted(built.prob_rows.items(), key=lambda kv: str(kv[0])):
        if not any(r["uncertainty"] == "INTERVAL" for r in rowmap.values()):
            continue
        if target[1] not in inst["initiatives"]:
            continue
        if any(r["identification_status"] == "UNIDENTIFIED" for r in rowmap.values()):
            unregistered.append(f"{target}/{field}: probability row partly UNIDENTIFIED")
            continue
        entries = sorted(rowmap)
        lo = [float(rowmap[e]["lower"]) if rowmap[e]["uncertainty"] == "INTERVAL" else float(rowmap[e]["value"]) for e in entries]
        hi = [float(rowmap[e]["upper"]) if rowmap[e]["uncertainty"] == "INTERVAL" else float(rowmap[e]["value"]) for e in entries]
        V = simplex_box_vertices(lo, hi)
        if not V:
            unregistered.append(f"{target}/{field}: interval bounds admit no probability vector")
            continue
        parent = S._get(inst, list(target))        # config dict (pi, P) or the rho table (field = (z, role[, state]))
        parent[field] = dict(zip(entries, V[0]))
        path = list(target) + [field]
        blocks.append(dict(name=f"prob:{'|'.join(map(str, target[1:]))}:{field}", kind="prob", path=path,
                           vertices=[dict(zip(entries, v)) for v in V]))
    regrows = [r for r in tables.get("U_registry", []) if r["group_id"] == "g1"]
    eps_R = float(regrows[0]["eps_R"]) if regrows and regrows[0]["eps_R"] else 0.0
    tie = float(regrows[0]["tie_tol"]) if regrows and regrows[0]["tie_tol"] else 1e-9
    if not regrows or not regrows[0]["eps_R"]:
        notes.append("eps_R not registered in U_registry: exact (eps_R = 0) classification only")
    nv = 1
    for b in blocks:
        nv *= 2 if b["kind"] == "scale" else len(b["vertices"])
    reg = {"groups": {"g1": {"meaning": regrows[0]["meaning"] if regrows else "elicited intervals (one block per elicited row)",
                             "blocks": blocks}}, "eps_R": eps_R, "tie_tol": tie, "seed": 20261011}
    return inst, reg, nv, notes, unregistered


def _jsonable_reg(reg):
    def conv(x):
        if isinstance(x, dict):
            return {str(k): conv(v) for k, v in x.items()}
        if isinstance(x, (list, tuple)):
            return [conv(v) for v in x]
        if isinstance(x, frozenset):
            return sorted(map(str, x))
        return x
    return conv(reg)


# =============================================================================
# 8. Rendering (Chapter 4 fragments)
# =============================================================================
def render_tables(cov, val, pres=None):
    L = ["# Chapter 4 data fragments (generated by 65; do not edit by hand)", "", "## Table 4.5 Parameter identification coverage", "",
         "| Block | Cells | OBSERVED | EXP internal | EXP external | PUBLIC bound | SCENARIO | DEFINITIONAL | UNIDENTIFIED |",
         "|---|---|---|---|---|---|---|---|---|"]
    for b in ("E", "G", "X", "R", "V"):
        c = cov.get(b, {})
        L.append(f"| {b} | {c.get('cells', 0)} | {c.get('OBSERVED', 0)} | {c.get('EXPERT_ELICITED_INTERNAL', 0)} | "
                 f"{c.get('EXPERT_ELICITED_EXTERNAL', 0)} | {c.get('PUBLIC_OBSERVED_BOUND', 0)} | {c.get('SCENARIO_SENSITIVITY', 0)} | "
                 f"{c.get('DEFINITIONAL', 0)} | {c.get('UNIDENTIFIED', 0)} |")
    L.append(f"| U | registry rows {cov.get('U', {}).get('cells', 0)} (registered before results: "
             f"{cov.get('U', {}).get('registered_before_results', 0)}) | | | | | | | |")
    L += ["", f"Validation: {len(val[0])} errors, {len(val[1])} warnings."]
    if pres:
        L += ["", "## Solver statuses (51 state machine)", "", "| Run | Status | Flags |", "|---|---|---|"]
        for k, r in pres["runs"].items():
            L.append(f"| {k} | {r['status']} | {', '.join(r['flags'])} |")
        a = pres["derived"].get("architecture", {})
        L += ["", "## Table 4.8 Architecture comparison (current policy)", "", "```", json.dumps(a, indent=1, default=str), "```"]
    return "\n".join(L) + "\n"


# =============================================================================
# 9. Template and pilot skeleton writers
# =============================================================================
def write_template(outdir, rows_by_table=None, xlsx=True):
    os.makedirs(outdir, exist_ok=True)
    for name, t in TABLES.items():
        with open(os.path.join(outdir, f"{name}.csv"), "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(t["cols"])
            for r in (rows_by_table or {}).get(name, []):
                w.writerow([r.get(c, "") for c in t["cols"]])
    with open(os.path.join(outdir, "DATA_DICTIONARY.md"), "w", encoding="utf-8") as f:
        f.write(dictionary_md())
    if xlsx:
        write_xlsx(os.path.join(outdir, "62_ENTERPRISE_EMPIRICAL_DATA_TEMPLATE_v2_3.xlsx"), rows_by_table)


def dictionary_md():
    L = ["# 62 data dictionary (generated from 65 TABLES; single source of truth)", "",
         "Common value columns for every parameter row: " + ", ".join(f"`{c}`" for c in VAL), "",
         f"- provenance_type: {', '.join(PROV)}", f"- uncertainty: {', '.join(UNC)}", f"- identification_status: {', '.join(IDS)}",
         "- UNIDENTIFIED rows keep value/lower/upper blank. A blank is never read as zero.", ""]
    for name, t in TABLES.items():
        L += [f"## {name} (Block {t['block']})", "", t["doc"], ""]
        if "params" in t:
            L += [f"Index columns: {', '.join(t['index'])}", "", "| parameter | unit | kind |", "|---|---|---|"]
            for p, (u, k) in t["params"].items():
                L.append(f"| {p} | {u} | {k} |")
        else:
            L.append(f"Columns: {', '.join(t['cols'])}")
        L.append("")
    return "\n".join(L)


def write_xlsx(path, rows_by_table=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    arial, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    lines = ["62 Enterprise empirical data template v2.3 (contract: 43 v2.1; model: 42 v2.2; loader: 65)",
             "Fill one sheet per block (S, E, G, X, R, V, U). One row = one parameter cell.",
             "Yellow header cells = columns you fill. Every identified cell needs provenance_type, source_id, period, unit, uncertainty, bounds and identification_status.",
             "Missing = provenance_type UNIDENTIFIED + identification_status UNIDENTIFIED + uncertainty NONE, with value/lower/upper blank. Never type 0 for missing.",
             "Expert values are intervals (lower/upper) unless a point is justified in assumption_note. Public bounds and scenarios are never point values.",
             "External sources (board-related, partner/supplier, investor) can only give EXPERT_ELICITED_EXTERNAL or PUBLIC_OBSERVED_BOUND. They are never optimizers.",
             "Ordinal constructs (A, H, WI, EA, EAX, G, R, D, kappa) are anchor levels 0-3, used only as within-construct thresholds or lookup keys.",
             "'*' in role/signal/outcome/response/capability_state = the value is claimed invariant across that index; the claim must be written in assumption_note.",
             "The EXAMPLES sheet shows the format with obviously fake IDs. It is never read by the loader.",
             "Run: python 65_enterprise_data_loader_v2_3.py --validate <this .xlsx>  then  --run <this .xlsx>"]
    for i, l in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=l).font = bold if i == 1 else arial
    ws.column_dimensions["A"].width = 160
    fill = PatternFill("solid", start_color="FFFF00")
    for name, t in TABLES.items():
        sh = wb.create_sheet(name)
        for c, col in enumerate(t["cols"], 1):
            cell = sh.cell(row=1, column=c, value=col)
            cell.font, cell.fill = bold, fill
            sh.column_dimensions[openpyxl.utils.get_column_letter(c)].width = max(12, len(col) + 2)
        for ri, r in enumerate((rows_by_table or {}).get(name, []), 2):
            for c, col in enumerate(t["cols"], 1):
                sh.cell(row=ri, column=c, value=r.get(col, "")).font = arial
        sh.freeze_panes = "A2"
        if "params" in t:
            n = len(t["cols"])
            colidx = {col: openpyxl.utils.get_column_letter(i + 1) for i, col in enumerate(t["cols"])}
            for col, opts in (("parameter", list(t["params"])), ("provenance_type", PROV), ("uncertainty", UNC),
                              ("identification_status", IDS)):
                dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
                sh.add_data_validation(dv)
                dv.add(f"{colidx[col]}2:{colidx[col]}5000")
    ex = wb.create_sheet("EXAMPLES")
    ex_rows = [["sheet", "row (format example only; fake IDs; never loaded)"],
               ["E_enterprise", "parameter=A_bar_BUD | value= | unit=TWD/period | provenance_type=EXPERT_ELICITED_INTERNAL | source_id=EX_SRC_01 | "
                                "period=EX_FY | uncertainty=INTERVAL | lower=1000000 | upper=1500000 | identification_status=IDENTIFIED_INTERVAL"],
               ["V_outcomes", "initiative_id=EX_INI | config_id=EX_AI | role=* | signal=* | outcome=EX_ERR | response=Use | parameter=l_u | "
                              "unit=TWD/event | provenance_type=OBSERVED | source_id=EX_SRC_02 | period=EX_FY | uncertainty=POINT | value=120 | "
                              "identification_status=IDENTIFIED_POINT | assumption_note=INVARIANT over role and signal: loss depends on outcome and response only"],
               ["G_policy", "policy_id=EX_CUR | policy_role=CURRENT | parameter=allow | initiative_id=EX_INI | config_id=EX_AI | value=1 | unit=binary | "
                            "provenance_type=OBSERVED | source_id=EX_SRC_03 | uncertainty=POINT | identification_status=IDENTIFIED_POINT"],
               ["R_responses", "initiative_id=EX_INI | config_id=EX_AI | signal=EX_FLAG | role=EX_ROLE | rho_kind=counterfactual | response=Verify | "
                               "parameter=rho | unit=probability | provenance_type=EXPERT_ELICITED_INTERNAL | source_id=EX_SRC_04 | uncertainty=INTERVAL | "
                               "lower=0.3 | upper=0.6 | identification_status=IDENTIFIED_INTERVAL (blind vignette; the respondent sees z, never omega)"]]
    for ri, r in enumerate(ex_rows, 1):
        for c, v in enumerate(r, 1):
            ex.cell(row=ri, column=c, value=v).font = bold if ri == 1 else arial
    ex.column_dimensions["A"].width = 16; ex.column_dimensions["B"].width = 220
    wb.save(path)


def _row(table, idx, param, unit, note=""):
    r = {c: "" for c in TABLES[table]["cols"]}
    r.update(idx); r.update(parameter=param, unit=unit, provenance_type="UNIDENTIFIED", uncertainty="NONE",
                            identification_status="UNIDENTIFIED", assumption_note=note)
    return r


PILOT = dict(units=["BU1", "BU2"], initiatives={"INI1": "BU1", "INI2": "BU1", "INI3": "BU2"}, cap="CAP1", local_unit="BU2",
             roles=["ROLE1", "ROLE2"], signals=["Z_FLAG", "Z_NOFLAG"], outcomes=["W_OK", "W_ERR"],
             configs={"SQ": True, "AI1": False}, ai_responses=["Use", "Verify", "Reject"], policies={"POL_CUR": "CURRENT", "POL_LEGALMIN": "LEGAL_MINIMUM"})


def pilot_rows(p=PILOT):
    """63 minimum identifiable pilot: every cell the 51 instance needs, at the recommended minimal granularity
    (invariance wildcards are pre-stated in assumption_note and must be confirmed or split)."""
    R = {n: [] for n in TABLES}
    cap, lu = p["cap"], p["local_unit"]
    R["E_enterprise"] += [_row("E_enterprise", {}, "A_bar_BUD", "TWD/period"), _row("E_enterprise", {}, "A_bar_ENG", "hours/period")]
    for u in p["units"]:
        R["E_enterprise"] += [_row("E_enterprise", {"unit_id": u}, x, un) for x, un in
                              (("rev_cap", "hours/period"), ("B0_BUD", "TWD/period"), ("B0_ENG", "hours/period"))]
    for ro in p["roles"]:
        R["E_enterprise"] += [_row("E_enterprise", {"role": ro}, x, "TWD/hour") for x in ("lam", "lam_u", "lam_U")]
    R["E_capabilities"] += [_row("E_capabilities", {"capability_id": cap, "scope": "SHARED"}, x, u)
                            for x, u in (("F", "TWD/period"), ("eng", "hours/period"), ("y_cur", "binary"))]
    R["E_capabilities"] += [_row("E_capabilities", {"capability_id": cap, "scope": "LOCAL", "unit_id": lu}, x, u)
                            for x, u in (("F", "TWD/period"), ("eng", "hours/period"))]
    states = {j: ["NONE", f"{cap}:S"] + ([f"{cap}:L"] if u == lu else []) for j, u in p["initiatives"].items()}
    for g, role in p["policies"].items():
        R["G_policy"].append(_row("G_policy", {"policy_id": g, "policy_role": role}, "gov_cost", "TWD/period"))
        for j in p["initiatives"]:
            for k in p["configs"]:
                R["G_policy"].append(_row("G_policy", {"policy_id": g, "policy_role": role, "initiative_id": j, "config_id": k}, "allow", "binary"))
            for x, un in (("G", "ordinal_level"), ("ebar", "incidents/period"), ("must", "binary"), ("forbid", "binary")):
                R["G_policy"].append(_row("G_policy", {"policy_id": g, "policy_role": role, "initiative_id": j}, x, un))
        R["G_policy"].append(_row("G_policy", {"policy_id": g, "policy_role": role, "unit_id": lu, "capability_id": cap}, "allow_loc", "binary"))
        for X in ("H_req", "WI_req", "EA_req", "G_req"):
            for a in range(4):
                for kk in range(3):
                    R["G_requirements"].append(_row("G_requirements", {"policy_id": g, "A_level": str(a), "kappa_level": str(kk)}, X, "ordinal_level"))
    for X in ("R_req", "D_req"):
        for a in range(4):
            R["G_requirements"].append(_row("G_requirements", {"policy_id": "*", "A_level": str(a)}, X, "ordinal_level"))
    for j, u in p["initiatives"].items():
        R["X_initiatives"] += [_row("X_initiatives", {"initiative_id": j, "unit_id": u}, "kappa", "ordinal_level"),
                               _row("X_initiatives", {"initiative_id": j, "unit_id": u}, "relevant_capabilities", "list"),
                               _row("X_initiatives", {"initiative_id": j, "unit_id": u, "capability_id": cap}, "tau", "TWD/period")]
        for s in states[j]:
            for x in ("R_level", "D_level"):
                R["X_capability_state"].append(_row("X_capability_state", {"initiative_id": j, "capability_state": s}, x, "ordinal_level"))
        for k, sq in p["configs"].items():
            base = {"initiative_id": j, "config_id": k}
            for x in ("is_status_quo", "A", "H", "WI", "EA", "EAX", "m", "prerequisites", "legal", "eff_H", "fallback", "cons",
                      "allowed_responses", "impl_x"):
                unit, _ = TABLES["X_configurations"]["params"][x]
                R["X_configurations"].append(_row("X_configurations", base, x, unit))
            for s in (states[j] if not sq else ["*"]):
                for x, un in (("impl", "TWD/period"), ("eng", "hours/period")):
                    R["X_config_costs"].append(_row("X_config_costs", dict(base, capability_state=s), x, un,
                                                    "INVARIANT across capability states: status quo needs no implementation" if s == "*" else ""))
            for ro in p["roles"]:
                R["R_roles"].append(_row("R_roles", dict(base, role=ro), "pi", "probability"))
            for z in p["signals"]:
                for w in p["outcomes"]:
                    R["R_signal_outcome"].append(_row("R_signal_outcome", dict(base, signal=z, outcome=w), "P", "probability"))
            resps = ["NA"] if sq else p["ai_responses"]
            kinds = ["observed"] if sq else ["observed", "counterfactual", "design"]
            for z in p["signals"]:
                for ro in p["roles"]:
                    for kind in kinds:
                        for r_ in resps:
                            R["R_responses"].append(_row("R_responses", dict(base, signal=z, role=ro, rho_kind=kind, response=r_), "rho", "probability"))
            R["V_volume"].append(_row("V_volume", base, "N", "events/period"))
            for w in p["outcomes"]:
                for r_ in resps:
                    for x in ("v_u", "v_x", "l_u", "l_x", "I_sev"):
                        R["V_outcomes"].append(_row("V_outcomes", dict(base, role="*", signal="*", outcome=w, response=r_), x,
                                                    TABLES["V_outcomes"]["params"][x][0],
                                                    "INVARIANT over role and signal: value/loss/incident depend on outcome and response only — confirm or split"))
            for x in ("c_u", "c_x"):
                R["V_outcomes"].append(_row("V_outcomes", dict(base, role="*", signal="*", outcome="*", response="*"), x, "TWD/event",
                                            "INVARIANT over role, signal, outcome, response: per-event operating cost of the configuration — confirm or split"))
            for ro in p["roles"]:
                for r_ in resps:
                    R["V_outcomes"].append(_row("V_outcomes", dict(base, role=ro, signal="*", outcome="*", response=r_), "t", "hours/event",
                                                "INVARIANT over signal and outcome: handling time depends on role and response — confirm or split"))
            for r_ in resps:
                R["V_outcomes"].append(_row("V_outcomes", dict(base, role="*", signal="*", outcome="*", response=r_), "t_rev", "hours/event",
                                            "INVARIANT over role, signal and outcome: review time depends on response — confirm or split"))
    R["U_registry"].append({c: "" for c in TABLES["U_registry"]["cols"]} | dict(group_id="g1", grouping_rule="one block per elicited interval row",
                                                                             registered_before_results="no"))
    return R


# =============================================================================
# 10. CLI and self-test
# =============================================================================
def cmd_validate(data, out=None):
    tables = read_data(data)
    errs, warns = validate(tables)
    cov = coverage(tables)
    rep = dict(loader=VERSION, errors=errs, warnings=warns, coverage=cov)
    if out:
        os.makedirs(out, exist_ok=True)
        json.dump(rep, open(os.path.join(out, "VALIDATION.json"), "w"), indent=1)
        open(os.path.join(out, "CH4_FRAGMENTS.md"), "w").write(render_tables(cov, (errs, warns)))
    return tables, rep


def cmd_run(data, out):
    tables, rep = cmd_validate(data, out)
    if rep["errors"]:
        return dict(status="VALIDATION_FAILED", errors=rep["errors"][:50], n_errors=len(rep["errors"]))
    built = build_instance(tables)
    res = run_point(built, out)
    res["coverage"] = rep["coverage"]
    os.makedirs(out, exist_ok=True)
    json.dump(res, open(os.path.join(out, "ENTERPRISE_RESULTS.json"), "w"), indent=1, default=str)
    open(os.path.join(out, "CH4_FRAGMENTS.md"), "w").write(render_tables(rep["coverage"], (rep["errors"], rep["warnings"]), res))
    return res


def _registration(tables):
    rows = [r for r in tables.get("U_registry", []) if r["group_id"] == "g1"]
    if not rows or rows[0]["registered_before_results"] != "yes" or not rows[0]["registration_date"]:
        return False
    return True


def cmd_make_registry(data, out):
    tables, rep = cmd_validate(data)
    if rep["errors"]:
        return dict(status="VALIDATION_FAILED", errors=rep["errors"][:50])
    built = build_instance(tables)
    inst, reg, nv, notes, unreg = make_registry(tables, built)
    if unreg:
        return dict(status="REGISTRY_INCOMPLETE", unregistered=unreg, detail="fix these interval rows before registering")
    S.validate_registry(reg, inst)
    os.makedirs(out, exist_ok=True)
    json.dump(_jsonable_reg(reg), open(os.path.join(out, "REGISTRY_DRAFT.json"), "w"), indent=1, default=str)
    commit = S.registry_commitment(reg, inst)
    open(os.path.join(out, "REGISTRY_COMMITMENT.txt"), "w").write(f"{commit}  registry+nominal (51 registry_commitment)\n")
    return dict(status="REGISTRY_DRAFT_WRITTEN", vertices=nv, blocks=len(reg["groups"]["g1"]["blocks"]), commitment=commit, notes=notes,
                next_step="set U_registry g1 registered_before_results=yes with a date, commit REGISTRY_DRAFT.json, "
                          "REGISTRY_COMMITMENT.txt and the data to git BEFORE running --robust (the commitment proves the data did not "
                          "change afterwards; the git history proves the order)"
                          + ("; vertex count exceeds 4096: register a coarser grouping first" if nv > 4096 else ""))


def cmd_robust(data, out, regdir):
    tables, rep = cmd_validate(data)
    if rep["errors"]:
        return dict(status="VALIDATION_FAILED")
    if not _registration(tables):
        return dict(status="NOT_PRE_REGISTERED", detail="U_registry g1 needs registered_before_results=yes and registration_date")
    built = build_instance(tables)
    inst, reg, nv, notes, unreg = make_registry(tables, built)
    if unreg:
        return dict(status="REGISTRY_INCOMPLETE", unregistered=unreg)
    commit = S.registry_commitment(reg, inst)
    rec = open(os.path.join(regdir, "REGISTRY_COMMITMENT.txt")).read().split()[0]
    if rec != commit:
        return dict(status="REGISTRY_COMMITMENT_MISMATCH", detail="data or registry changed after registration; re-register (and disclose)")
    if nv > 4096:
        return dict(status="REGISTRY_TOO_LARGE", vertices=nv)
    g = [p for p in inst["policies"] if built.policy_role.get(p) == "CURRENT"] or sorted(inst["policies"])[:1]
    pre = S.solve(inst, "B", [g[0]])
    if "F_E" not in pre:
        return dict(status="ROBUST_NOT_RUN", detail=f"the nominal instance has no Model B point solution ({pre['status']}); "
                                                     "identify the missing inputs first", flags=pre["flags"])
    try:
        rr = S.solve_robust(inst, g[0], reg, "g1")
        res = dict(robust=rr, notes=notes)
        if rr.get("status") == "OK_INTERVAL":
            blocks = reg["groups"]["g1"]["blocks"]
            tinst = S.theta_instances(inst, blocks)
            cands = [d for d in S.enumerate_portfolios(inst, g[0], static_only=True)
                     if S.robust_feasibility(inst, g[0], d, blocks, tinst, spts=[]) == "ROBUSTLY_FEASIBLE"]
            res["voi"] = S.voi_table(inst, g[0], cands, blocks, grid=4)
            res["element_y"] = S.classify_element(inst, g[0], reg, lambda sig: tuple(sig[0]) if sig else None)
    except (S.MissingInput, S.InvalidInput, S.ScopeError) as e:
        return dict(status="ROBUST_NOT_RUN", detail=f"{type(e).__name__}: {e}")
    os.makedirs(out, exist_ok=True)
    json.dump(res, open(os.path.join(out, "ROBUST_RESULTS.json"), "w"), indent=1, default=str)
    return dict(status=rr.get("status"), max_regret=rr.get("max_regret"))


def _fill_test(R, rng):
    """TEST FIXTURE ONLY: fill the pilot skeleton with arbitrary numbers to exercise the pipeline. Never data."""
    src = {c: "" for c in TABLES["S_sources"]["cols"]}
    src.update(source_id="TEST_ONLY", source_scope="internal", source_kind="other", description="TEST FIXTURE - not data")
    R["S_sources"] = [src]

    def pt(r, v):
        r.update(value=str(v), provenance_type="DEFINITIONAL", source_id="TEST_ONLY", period="TEST", uncertainty="POINT",
                 identification_status="IDENTIFIED_POINT", assumption_note=r["assumption_note"] or "TEST")
    probs = {}
    for name, rows in R.items():
        for r in rows:
            if "parameter" not in r:
                continue
            p = r["parameter"]
            kind = TABLES_KIND[(name, p)]
            if name == "E_enterprise":
                v = {"A_bar_BUD": 400, "A_bar_ENG": 200, "rev_cap": 1e6, "B0_BUD": 150, "B0_ENG": 80, "lam": 1.0, "lam_u": 1.0, "lam_U": 3.0}[p]
            elif name == "E_capabilities":
                v = {"F": 40 if r["scope"] == "SHARED" else 30, "eng": 10, "y_cur": 0}[p]
            elif name == "G_policy":
                v = {"gov_cost": 5, "allow": 1, "allow_loc": 1, "G": 2, "ebar": 10, "must": 0, "forbid": 0}[p]
            elif name == "G_requirements":
                a = int(r["A_level"])
                v = {"H_req": 2 if a >= 2 else 0, "WI_req": 1 if a >= 2 else 0, "EA_req": 1 if a >= 2 else 0, "G_req": 1,
                     "R_req": min(a, 2), "D_req": min(a, 2)}[p]
            elif name == "X_initiatives":
                v = {"kappa": 1, "relevant_capabilities": "CAP1", "tau": 0}[p]
            elif name == "X_capability_state":
                v = 3
            elif name == "X_configurations":
                sq = r["config_id"] == "SQ"
                v = dict(is_status_quo=int(sq), A=0 if sq else 2, H=3 if sq else 2, WI=0 if sq else 2, EA=0 if sq else 1, EAX=1,
                         m="HumanOnly" if sq else "AI->Human", prerequisites="" if sq else "CAP1", legal=1, eff_H=1, fallback=1, cons=1,
                         allowed_responses="NA" if sq else "Use;Verify;Reject", impl_x=0)[p]
                if p == "prerequisites" and sq:
                    v = "NONE"
            elif name == "X_config_costs":
                v = 0 if r["config_id"] == "SQ" else (round(rng.uniform(10, 30), 2) if r["capability_state"] != "NONE" else 999)
            elif name == "R_roles":
                v = 0.5
            elif name == "R_signal_outcome":
                v = {("Z_FLAG", "W_ERR"): 0.07, ("Z_NOFLAG", "W_ERR"): 0.03, ("Z_FLAG", "W_OK"): 0.43, ("Z_NOFLAG", "W_OK"): 0.47}[(r["signal"], r["outcome"])]
            elif name == "R_responses":
                key = (r["initiative_id"], r["config_id"], r["signal"], r["role"], r["rho_kind"])
                if r["response"] == "NA":
                    v = 1.0
                else:
                    if key not in probs:
                        a = rng.uniform(0.3, 0.6); b = rng.uniform(0, 1 - a)
                        probs[key] = {"Use": round(a, 4), "Verify": round(b, 4)}
                        probs[key]["Reject"] = round(1 - probs[key]["Use"] - probs[key]["Verify"], 4)
                    v = probs[key][r["response"]]
            elif name == "V_volume":
                v = 100
            elif name == "V_outcomes":
                bad = r["outcome"] == "W_ERR" and r["response"] in ("Use", "NA")
                v = {"v_u": 0 if r["config_id"] == "SQ" else round(rng.uniform(0, 2), 2), "v_x": 0,
                     "l_u": round(rng.uniform(5, 15), 2) if bad else 0, "l_x": 0, "c_u": 0.1, "c_x": 0,
                     "t": {"NA": 1.0, "Use": 0.2, "Verify": 0.6, "Reject": 1.1}.get(r["response"], 0.5),
                     "t_rev": 0.2, "I_sev": 0.02 if bad else 0}[p]
            else:
                continue
            pt(r, v)
    return R


def self_test():
    tmp = tempfile.mkdtemp(prefix="loader_selftest_")
    try:
        write_template(os.path.join(tmp, "blank"))
        t, rep = cmd_validate(os.path.join(tmp, "blank"))
        assert not rep["errors"], rep["errors"][:3]
        sk = os.path.join(tmp, "pilot")
        write_template(sk, pilot_rows())
        t, rep = cmd_validate(sk)
        assert not rep["errors"], rep["errors"][:3]
        assert rep["coverage"]["V"]["UNIDENTIFIED"] == rep["coverage"]["V"]["cells"] > 0
        r = cmd_run(sk, os.path.join(tmp, "out0"))
        assert r.get("overall") == "NO_EVALUABLE_INITIATIVE_OR_POLICY" or all(x["status"] != "OK_POINT" for x in r["runs"].values())
        # xlsx path reads identically
        t2 = read_data(os.path.join(sk, "62_ENTERPRISE_EMPIRICAL_DATA_TEMPLATE_v2_3.xlsx"))
        assert sum(len(v) for v in t2.values()) == sum(len(v) for v in t.values())
        # filled TEST fixture: full pipeline
        R = _fill_test(pilot_rows(), random.Random(3))
        fx = os.path.join(tmp, "fixture")
        write_template(fx, R, xlsx=False)
        t, rep = cmd_validate(fx)
        assert not rep["errors"], rep["errors"][:5]
        res = cmd_run(fx, os.path.join(tmp, "out1"))
        sts = {k: v["status"] for k, v in res["runs"].items()}
        assert sts["B|POL_CUR"] in ("OK_POINT", "RESTRICTED_SOLVE"), sts
        assert "architecture" in res["derived"]
        # zero-fill and external-OBSERVED guards
        bad = copy.deepcopy(R)
        bad["E_enterprise"][0].update(value="0", provenance_type="UNIDENTIFIED", identification_status="UNIDENTIFIED", uncertainty="NONE")
        write_template(os.path.join(tmp, "bad"), bad, xlsx=False)
        assert any("no zero-fill" in e for e in cmd_validate(os.path.join(tmp, "bad"))[1]["errors"])
        bad = copy.deepcopy(R)
        bad["S_sources"].append(dict(bad["S_sources"][0], source_id="EXT1", source_scope="external", external_role="investor"))
        bad["E_enterprise"][0].update(source_id="EXT1", provenance_type="OBSERVED")
        write_template(os.path.join(tmp, "bad2"), bad, xlsx=False)
        assert any("external source" in e for e in cmd_validate(os.path.join(tmp, "bad2"))[1]["errors"])
        bad = copy.deepcopy(R)
        bad["E_enterprise"][0].update(value="", uncertainty="INTERVAL", lower="300", upper="500", provenance_type="PUBLIC_OBSERVED_BOUND",
                                      identification_status="IDENTIFIED_INTERVAL")
        write_template(os.path.join(tmp, "bad3"), bad, xlsx=False)
        assert any("BOUND_ONLY" in e for e in cmd_validate(os.path.join(tmp, "bad3"))[1]["errors"])
        # missing lookup stays UNIDENTIFIED (not zero): blank an AI loss cell -> RESTRICTED or different from zero
        miss = copy.deepcopy(R)
        for r in miss["V_outcomes"]:
            if r["parameter"] == "l_u" and r["config_id"] == "AI1" and r["outcome"] == "W_ERR" and r["response"] == "Use":
                r.update(value="", provenance_type="UNIDENTIFIED", identification_status="UNIDENTIFIED", uncertainty="NONE", source_id="",
                         period="", unit="TWD/event")
        write_template(os.path.join(tmp, "miss"), miss, xlsx=False)
        rm = cmd_run(os.path.join(tmp, "miss"), os.path.join(tmp, "out2"))
        assert rm["runs"]["B|POL_CUR"]["status"] != "OK_POINT", rm["runs"]["B|POL_CUR"]["status"]
        # override order: a specific UNIDENTIFIED row beats a wildcard value row (never filled from '*')
        ov = copy.deepcopy(R)
        spec = next(r for r in ov["X_config_costs"] if r["initiative_id"] == "INI1" and r["config_id"] == "AI1"
                    and r["capability_state"] == "CAP1:S" and r["parameter"] == "impl")
        spec.update(value="", provenance_type="UNIDENTIFIED", identification_status="UNIDENTIFIED", uncertainty="NONE",
                    source_id="", period="")
        ov["X_config_costs"].append(dict(spec, capability_state="*", value="5", provenance_type="DEFINITIONAL", source_id="TEST_ONLY",
                                         period="TEST", uncertainty="POINT", identification_status="IDENTIFIED_POINT",
                                         assumption_note="TEST invariance"))
        ov["X_config_costs"] = [r for r in ov["X_config_costs"] if not (r["initiative_id"] == "INI1" and r["config_id"] == "AI1"
                                and r["parameter"] == "impl" and r["capability_state"] not in ("CAP1:S", "*"))] 
        bo = build_instance(ov)
        impl = bo.inst["initiatives"]["INI1"]["configs"][1]["impl"]
        assert impl[frozenset({("CAP1", "S")})] is None and impl[frozenset()] == 5.0, impl
        # duplicate source id with different scope is rejected
        dup = copy.deepcopy(R)
        dup["S_sources"].append(dict(dup["S_sources"][0], source_scope="external"))
        write_template(os.path.join(tmp, "dup"), dup, xlsx=False)
        assert any("duplicate source_id" in e for e in cmd_validate(os.path.join(tmp, "dup"))[1]["errors"])
        # missing y_cur is never read as 0
        yc = copy.deepcopy(R)
        yc["E_capabilities"] = [r for r in yc["E_capabilities"] if r["parameter"] != "y_cur"]
        by = build_instance(yc)
        assert not by.inst["initiatives"] and any(e[0] == "UNIDENTIFIED_Y_CUR" for e in by.log)
        # simplex-box vertices
        V = simplex_box_vertices([0.2, 0.1, 0.0], [0.6, 0.5, 0.4])
        assert all(abs(sum(v) - 1) < 1e-9 for v in V) and len(V) >= 3
        # interval registry: one interval cell -> one scale block, commitment, robust run
        iv = copy.deepcopy(R)
        for r in iv["V_outcomes"]:
            if r["parameter"] == "l_u" and r["config_id"] == "AI1" and r["outcome"] == "W_ERR" and r["response"] == "Use" and r["initiative_id"] == "INI1":
                r.update(value="", uncertainty="INTERVAL", lower="4", upper="20", provenance_type="EXPERT_ELICITED_INTERNAL",
                         identification_status="IDENTIFIED_INTERVAL")
        write_template(os.path.join(tmp, "iv"), iv, xlsx=False)
        assert cmd_robust(os.path.join(tmp, "iv"), os.path.join(tmp, "out3"), os.path.join(tmp, "reg"))["status"] == "NOT_PRE_REGISTERED"
        iv["U_registry"][0].update(registered_before_results="yes", registration_date="TEST", eps_R="0")
        write_template(os.path.join(tmp, "iv"), iv, xlsx=False)
        mr = cmd_make_registry(os.path.join(tmp, "iv"), os.path.join(tmp, "reg"))
        assert mr["status"] == "REGISTRY_DRAFT_WRITTEN" and mr["blocks"] == 1 and mr["vertices"] == 2, mr
        rr = cmd_robust(os.path.join(tmp, "iv"), os.path.join(tmp, "out3"), os.path.join(tmp, "reg"))
        assert rr["status"] in ("OK_INTERVAL", "INFEASIBLE"), rr
        open(os.path.join(tmp, "reg", "REGISTRY_COMMITMENT.txt"), "w").write("0" * 64 + "\n")
        assert cmd_robust(os.path.join(tmp, "iv"), os.path.join(tmp, "out3"), os.path.join(tmp, "reg"))["status"] == "REGISTRY_COMMITMENT_MISMATCH"
    finally:
        shutil.rmtree(tmp)
    print("SELF-TEST PASS: blank template and pilot skeleton validate; all-UNIDENTIFIED skeleton yields no point decision; xlsx == csv; "
          "TEST fixture runs end-to-end through 51 (B, C, A+, C0, A, thresholds, PR, interfaces); zero-fill, external-OBSERVED and "
          "public-bound-as-identified rejected; missing loss cell not zero-filled; simplex-box vertices; registry commitment enforced "
          "(TEST fixtures only; nothing written to the repository)")


def main(argv):
    if not argv or argv[0] == "--help":
        print(__doc__); return 0
    cmd = argv[0]
    opt = {argv[i]: argv[i + 1] for i in range(2, len(argv) - 1, 2)} if len(argv) > 2 else {}
    out = opt.get("--out", os.path.join(HERE, "enterprise_results"))
    if cmd == "--self-test":
        self_test(); return 0
    if cmd == "--make-template":
        write_template(argv[1]); print(f"template written to {argv[1]}"); return 0
    if cmd == "--make-pilot":
        write_template(argv[1], pilot_rows()); print(f"63 pilot skeleton written to {argv[1]}"); return 0
    if cmd == "--validate":
        _, rep = cmd_validate(argv[1], out)
        print(json.dumps(dict(errors=len(rep["errors"]), first_errors=rep["errors"][:10], warnings=len(rep["warnings"]), coverage=rep["coverage"]), indent=1))
        return 1 if rep["errors"] else 0
    if cmd == "--run":
        r = cmd_run(argv[1], out)
        print(json.dumps({k: r[k] for k in r if k in ("status", "errors", "n_errors", "overall")} or
                         {k: v["status"] for k, v in r["runs"].items()}, indent=1, default=str))
        return 0
    if cmd == "--make-registry":
        print(json.dumps(cmd_make_registry(argv[1], out), indent=1)); return 0
    if cmd == "--robust":
        print(json.dumps(cmd_robust(argv[1], out, opt["--registry"]), indent=1, default=str)); return 0
    print(__doc__); return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
